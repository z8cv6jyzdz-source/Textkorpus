# -*- coding: utf-8 -*-
"""
Textvorschlag_2.4_2026-09-27.py — Task 7 des Plans, Fassung 4 (2.4 knapp nach Lloyd et al., 2016)
Bachelorarbeit U15-Plyometrie · DSHS Köln

Zweck: (1) Messung von Master, Fassung 3 (26.09.) und Fassung 4 in beiden Varianten (Wörter als Leerraum-Token,
Sätze nach der Satzteilung von Manuskriptstand_2026-09-25.py, Median, längster Satz, Sätze über 32 und 40 Wörter,
Semikola außerhalb der Zitierklammern, Doppelpunkte, nummerierte Abschnittsverweise, Belegklammern),
(2) Kürzungsleiter vom Master über Fassung 3 zu Fassung 4: Stufen 0 bis 5 aus der Teilstückzuordnung des Masters
(Kuerzungsleiter_2.4_2026-09-26.tsv, Wort für Wort geprüft), Stufe 6 = Fassung 3, Stufe 7 = Kern nach Lloyd aus dem
Verbleib jedes Satzes der Fassung 3 und der Herkunft jedes Satzes der Fassung 4,
(3) optionale Stufen unterhalb der Empfehlung,
(4) Probeexemplare NUR IM ARBEITSORDNER: Kopie des Masters mit Variante A (2.4 ohne Unterabschnitte, die Überschriften
2.4.1 bis 2.4.3 entfallen) und mit Variante B (Unterabschnitte bleiben), darauf Manuskriptstand_2026-09-25.py,
auf Variante A zusätzlich Endabgleich_Manuskript_2026-09-25.py.
Der Master wird nicht verändert (Prozessregel F16 § 1.2, Verfasser 26.09.: keine Änderung am Master durch Claude).

Aufruf:
python Textvorschlag_2.4_2026-09-27.py <Master.docx> <Variante A .txt> <Variante B .txt> <Fassung 3 .txt>
       <Kuerzungsleiter_2.4_2026-09-26.tsv> <Verbleib_Fassung3 .tsv> <Herkunft_Fassung4 .tsv> <Optionale_Stufen .tsv>
       <Ordner 03_Skripte> <Ordner 02_Befunde> <Kennzahlen_2026-09-22.md aus dem Archiv> <Arbeitsordner> <Laufprotokoll.txt>
Übertragungsvorlage: Zeilen mit # sind Kommentare, eine Zeile „@@ 2.4“ eröffnet den Abschnitt, Absätze durch Leerzeilen getrennt.
Ohne Semikolon im Skript (chr(59)).
Fassung: 2026-09-27, erste Fassung (ersetzt für Fassung 4 das Skript vom 26.09., das nur Vorlagen mit allen vier Abschnitten kennt).
"""
import sys
import os
import re
import copy
import hashlib
import statistics
import subprocess
import zipfile
from collections import OrderedDict
from docx import Document
from docx.text.paragraph import Paragraph

(MASTER, TXT_A, TXT_B, TXT_F3, LEITER, VERBLEIB, HERKUNFT, OPT,
 SKR, BEF, KALT, ARB, LOG) = sys.argv[1:14]
SEMI = chr(59)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
ALLE = ['2.4', '2.4.1', '2.4.2', '2.4.3']
BUDGET = 700
os.makedirs(ARB, exist_ok=True)
AUS = []


def log(*t):
    AUS.append(' '.join(str(x) for x in t))


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()


def saetze(text):
    """Satzteilung wörtlich wie Manuskriptstand_2026-09-25.py (Fassung 2)"""
    t = re.sub(r'(\d)\.(\d)', r'\1<P>\2', text)
    t = re.sub(r'\b(et al|Abschn|Tab|Abb|vgl|bzw|ca|Nr|Aufl|Hrsg|Jg)\.', r'\1<P>', t)
    t = re.sub(r'\b([A-Z])\.\s', r'\1<P> ', t)
    t = re.sub(r'\bS\.\s', 'S<P> ', t)
    t = re.sub(r'\b(u|z|d)\.\s?(a|B|h)\.', r'\1<P>\2<P>', t)
    t = re.sub(r'(\d{2})\.(\d{2})\.(\d{4})', r'\1<P>\2<P>\3', t)
    t = re.sub(r'(\d{2})\.(\d{2})\.', r'\1<P>\2<P>', t)
    t = re.sub(r'(\d)\.\s', r'\1<P> ', t)
    teile = [s.strip() for s in re.split(r'(?<=[.!?])\s+(?=[A-ZÄÖÜ„(⟨])', t) if s.strip()]
    return [s.replace('<P>', '.') for s in teile]


def ohne_zitierklammern(text):
    return re.sub(r'\([^)]*\d{4}[^)]*\)', '', text)


RE_REF = re.compile(r'Abschn(?:itt)?\.?\s*\d|Kapitel\s*\d|siehe (?:oben|unten)')


def kennwerte(absaetze):
    text = ' '.join(absaetze)
    ss = [s for a in absaetze for s in saetze(a)]
    lang = [len(s.split()) for s in ss]
    return OrderedDict([
        ('Wörter', sum(len(a.split()) for a in absaetze)),
        ('Absätze', len(absaetze)),
        ('Sätze', len(ss)),
        ('Median', statistics.median(lang) if lang else 0),
        ('längster', max(lang) if lang else 0),
        ('über 32', sum(1 for x in lang if x > 32)),
        ('über 40', sum(1 for x in lang if x > 40)),
        ('Semikola', ohne_zitierklammern(text).count(SEMI)),
        ('Doppelpunkte', text.count(':')),
        ('Verweise', len(RE_REF.findall(text))),
        ('Belegklammern', len(re.findall(r'\((?:[A-ZÄÖÜ][^()]*?\d{4}[^()]*)\)', text))),
    ])


def zeile(k):
    return ' · '.join('%s %s' % (a, b) for a, b in k.items())


def abschnitt(doc, nr):
    ps = doc.paragraphs
    kopf = [k for k, p in enumerate(ps) if p.style.name.startswith('Heading') and p.text.strip().startswith(nr + ' ')]
    if len(kopf) != 1:
        raise SystemExit('Überschrift ' + nr + ' nicht genau einmal gefunden')
    j = kopf[0] + 1
    body = []
    while j < len(ps) and not ps[j].style.name.startswith('Heading'):
        if ps[j].text.strip():
            body.append(ps[j])
        j += 1
    return ps[kopf[0]], body


def vorlage_lesen(pfad):
    sek = OrderedDict()
    akt = None
    for block in re.split(r'\n\s*\n', open(pfad, encoding='utf-8').read()):
        zeilen = []
        for z in block.split('\n'):
            if z.startswith('@@'):
                akt = z[2:].strip()
                sek.setdefault(akt, [])
                continue
            if z.startswith('#') or not z.strip():
                continue
            zeilen.append(z.strip())
        if zeilen:
            sek[akt].append(' '.join(zeilen))
    return sek


def tsv(pfad):
    return [z.split('\t') for z in open(pfad, encoding='utf-8').read().split('\n') if z.strip() and not z.startswith('#')]


# ------------------------------------------------------------------ Eingänge
log('Textvorschlag_2.4_2026-09-27.py, Laufprotokoll')
for p in (MASTER, TXT_A, TXT_B, TXT_F3, LEITER, VERBLEIB, HERKUNFT, OPT, KALT):
    log('Eingang', os.path.basename(p), os.path.getsize(p), 'Byte, SHA-256', sha(p))
log('Master ohne comments.xml' if 'word/comments.xml' not in zipfile.ZipFile(MASTER).namelist() else 'WARNUNG: Master trägt Kommentare')
d = Document(MASTER)
alt = OrderedDict((nr, [p.text.strip() for p in abschnitt(d, nr)[1]]) for nr in ALLE)
va = vorlage_lesen(TXT_A)
vb = vorlage_lesen(TXT_B)
f3 = vorlage_lesen(TXT_F3)
if list(va.keys()) != ['2.4']:
    raise SystemExit('Variante A trägt nicht genau den Abschnitt 2.4')
if list(vb.keys()) != ALLE or list(f3.keys()) != ALLE:
    raise SystemExit('Variante B oder Fassung 3 trägt nicht genau die Abschnitte ' + ', '.join(ALLE))

# ------------------------------------------------------------------ Messung
log('')
log('Messung (Master, Fassung 3, Fassung 4 Variante A und B)')
km = kennwerte([a for nr in ALLE for a in alt[nr]])
k3 = kennwerte([a for nr in ALLE for a in f3[nr]])
ka = kennwerte(va['2.4'])
kb = kennwerte([a for nr in ALLE for a in vb[nr]])
log('  Master:        ' + zeile(km))
log('  Fassung 3:     ' + zeile(k3))
log('  Fassung 4 A:   ' + zeile(ka))
log('  Fassung 4 B:   ' + zeile(kb))
for nr in ALLE:
    log('  Variante B %-6s %s' % (nr, zeile(kennwerte(vb[nr]))))
log('  Budget %d · Variante A %d (Abstand %+d) · Variante B %d (Abstand %+d) · gegenüber Fassung 3 %+d und %+d · gegenüber dem Master %+d und %+d' % (
    BUDGET, ka['Wörter'], ka['Wörter'] - BUDGET, kb['Wörter'], kb['Wörter'] - BUDGET,
    ka['Wörter'] - k3['Wörter'], kb['Wörter'] - k3['Wörter'], ka['Wörter'] - km['Wörter'], kb['Wörter'] - km['Wörter']))
log('')
log('Sätze der Variante A mit Wortzahl')
sa_ids = []
for i, a in enumerate(va['2.4'], 1):
    for j, s in enumerate(saetze(a), 1):
        sa_ids.append(('2.4', i, j, s))
        log('  2.4 Abs. %d S%d (%d): %s' % (i, j, len(s.split()), s))
log('')
log('Sätze der Variante B mit Wortzahl')
for nr in ALLE:
    for i, a in enumerate(vb[nr], 1):
        for j, s in enumerate(saetze(a), 1):
            log('  %s Abs. %d S%d (%d): %s' % (nr, i, j, len(s.split()), s))

# ------------------------------------------------------------------ Kürzungsleiter Stufen 0 bis 5 (Master)
log('')
log('Kürzungsleiter')
zuordnung = [(a, int(b), int(s), k, g, t) for a, b, s, k, g, t in tsv(LEITER)]
bestand = ' '.join(a for nr in ALLE for a in alt[nr]).split()
stuecke = ' '.join(z[5] for z in zuordnung).split()
if bestand != stuecke:
    raise SystemExit('Die Zuordnung gibt den Masterbestand nicht Wort für Wort wieder')
log('  Zuordnung gibt den Masterbestand Wort für Wort wieder: %d Wörter in %d Teilstücken' % (len(bestand), len(zuordnung)))
kat = OrderedDict([('K1', 'Selbstkommentar, Überleitung und Abschnittsverweise streichen'),
                   ('K2', 'Fehler und Zitierfallen streichen (W14, W16, T4)'),
                   ('K3', 'Methodendiskussion und Bezüge auf die eigene Studie nach 6.1 und 6.2 verlagern'),
                   ('K4', 'Rollentrennung: Mechanismus nach 2.1, Testaufbau steht in 4.4'),
                   ('K5', 'Doppelungen und entbehrliche Einzelbefunde streichen')])
summe = OrderedDict((k, sum(len(z[5].split()) for z in zuordnung if z[3] == k)) for k in list(kat) + ['U'])
stand = len(bestand)
log('  Stufe 0: Master %d' % stand)
for i, (k, text) in enumerate(kat.items(), 1):
    stand -= summe[k]
    log('  Stufe %d: %s, −%d → %d' % (i, text, summe[k], stand))
if stand != summe['U']:
    raise SystemExit('Leiterrechnung der Stufen 1 bis 5 stimmt nicht')
log('  Stufe 6: Fassung 3 vom 26.09. (Übernommenes verdichtet, korrigiert, ergänzt), %+d → %d' % (k3['Wörter'] - stand, k3['Wörter']))

# ------------------------------------------------------------------ Stufe 7 (Fassung 3 → Fassung 4 A)
f3_ids = []
for nr in ALLE:
    for i, a in enumerate(f3[nr], 1):
        for j, s in enumerate(saetze(a), 1):
            f3_ids.append((nr, i, j, s))
verbleib = OrderedDict(((a, int(b), int(s)), (k, z, g)) for a, b, s, k, z, g in tsv(VERBLEIB))
if [x[:3] for x in f3_ids] != list(verbleib.keys()):
    raise SystemExit('Verbleibtabelle deckt die Sätze der Fassung 3 nicht genau ab')
herkunft = OrderedDict(((a, int(b), int(s)), k) for a, b, s, k in tsv(HERKUNFT))
if [x[:3] for x in sa_ids] != list(herkunft.keys()):
    raise SystemExit('Herkunftstabelle deckt die Sätze der Variante A nicht genau ab')
w3 = OrderedDict()
for nr, i, j, s in f3_ids:
    k = verbleib[(nr, i, j)][0]
    w3[k] = w3.get(k, 0) + len(s.split())
w4 = OrderedDict()
for nr, i, j, s in sa_ids:
    k = herkunft[(nr, i, j)]
    w4[k] = w4.get(k, 0) + len(s.split())
if sum(w3.values()) != k3['Wörter'] or sum(w4.values()) != ka['Wörter']:
    raise SystemExit('Satzsummen stimmen nicht mit den Wortzahlen überein')
namen = OrderedDict([('N', 'entbehrliche Einzelbefunde und Kennwerte streichen'),
                     ('D', 'Doppelung streichen, steht an anderer Stelle (4.4.2)'),
                     ('L', 'Validitätsdebatte und Gegenbefunde zur eigenen Messung nach Kapitel 6 verlagern'),
                     ('M', 'Rollentrennung: Mechanismus und Komponentenbegriff nach 2.1')])
stand = k3['Wörter']
log('  Stufe 7: Kern nach Lloyd et al. (2016), Fassung 3 → Fassung 4, Variante A')
for k, text in namen.items():
    stand -= w3.get(k, 0)
    log('    7%s %s, −%d → %d' % (k, text, w3.get(k, 0), stand))
d_v = w4.get('V', 0) - w3.get('V', 0)
stand += d_v
log('    7V Verdichten (%d Wörter der Fassung 3 in %d Wörtern der Fassung 4), %+d → %d' % (w3.get('V', 0), w4.get('V', 0), d_v, stand))
d_b = w4.get('B', 0) - w3.get('B', 0)
stand += d_b
log('    7B Übernommenes angepasst (Doppelpunkt, Überdehnung, Beleg, Gruppenschluss), %+d → %d' % (d_b, stand))
stand += w4.get('neu', 0)
log('    7neu Leitgedanke Messfehler und Stichprobenumfang (Hopkins, 2000) mit Definition, %+d → %d' % (w4.get('neu', 0), stand))
if stand != ka['Wörter']:
    raise SystemExit('Leiterrechnung der Stufe 7 stimmt nicht')
log('  Verbleib der Sätze der Fassung 3:')
for nr, i, j, s in f3_ids:
    k, z, g = verbleib[(nr, i, j)]
    log('    %s %s Abs. %d S%d (%d) → %s: %s … · %s' % (k, nr, i, j, len(s.split()), z, ' '.join(s.split()[:8]), g))

# ------------------------------------------------------------------ optionale Stufen
log('')
log('Optionale Stufen unterhalb der Empfehlung (Variante A, kumuliert)')
opt = list(va['2.4'])
stand = ka['Wörter']
for stufe, nr, a_alt, a_neu, preis in tsv(OPT):
    treffer = [i for i, a in enumerate(opt) if a_alt in a]
    if len(treffer) != 1 or opt[treffer[0]].count(a_alt) != 1:
        raise SystemExit('Optionale Stufe ' + stufe + ': Wortlaut nicht genau einmal gefunden')
    opt[treffer[0]] = opt[treffer[0]].replace(a_alt, a_neu, 1).strip()
    neu_stand = kennwerte(opt)['Wörter']
    log('  %s: %+d → %d · Preis: %s' % (stufe, neu_stand - stand, neu_stand, preis))
    stand = neu_stand

# ------------------------------------------------------------------ Satzvergleich
s3 = [x[3] for x in f3_ids]
log('')
log('Satzvergleich: Variante A %d Sätze, davon wortgleich aus Fassung 3 %d · Fassung 3 %d Sätze' % (
    len(sa_ids), sum(1 for x in sa_ids if x[3] in s3), len(s3)))


# ------------------------------------------------------------------ Probeexemplare (nur Arbeitsordner)
def probe_bauen(ziel, vorlage, abschnitte):
    doc = Document(MASTER)
    kopf24, body24 = abschnitt(doc, '2.4')
    if not all(p.style.name == 'Normal' for p in body24):
        raise SystemExit('Abschnitt 2.4 enthält Absätze außerhalb der Formatvorlage Standard')
    muster = copy.deepcopy(body24[0]._p)
    # Unterabschnitte, die die Vorlage nicht trägt, samt Überschrift entfernen
    for nr in ALLE[1:]:
        if nr not in abschnitte:
            k, b = abschnitt(doc, nr)
            for p in b:
                p._p.getparent().remove(p._p)
            k._p.getparent().remove(k._p)
    for nr in abschnitte:
        kopf, body = abschnitt(doc, nr)
        for p in body:
            p._p.getparent().remove(p._p)
        letzte = kopf._p
        for txt in vorlage[nr]:
            el = copy.deepcopy(muster)
            for r in el.findall(W + 'r'):
                el.remove(r)
            letzte.addnext(el)
            Paragraph(el, kopf._parent).add_run(txt)
            letzte = el
    doc.save(ziel)
    kontrolle = Document(ziel)
    for nr in abschnitte:
        if [p.text.strip() for p in abschnitt(kontrolle, nr)[1]] != vorlage[nr]:
            raise SystemExit('Probeexemplar gibt ' + nr + ' nicht wortgleich wieder')
    for nr in ALLE[1:]:
        if nr not in abschnitte and any(p.style.name.startswith('Heading') and p.text.strip().startswith(nr + ' ') for p in kontrolle.paragraphs):
            raise SystemExit('Überschrift ' + nr + ' im Probeexemplar nicht entfernt')
    return ziel


def manuskriptstand(probe, name):
    ms_txt = os.path.join(ARB, name)
    subprocess.run([sys.executable, os.path.join(SKR, 'Manuskriptstand_2026-09-25.py'), probe, ms_txt], check=True, capture_output=True)
    ms = open(ms_txt, encoding='utf-8').read()
    for z in ms.split('\n'):
        zs = z.strip()
        if re.match(r'^2\.4(\.\d)?\s', zs) or zs.startswith('Kapitel 2') or zs.startswith('Absatztext') or zs.startswith('Semikola') or zs.startswith('Nummerierte'):
            log('    ' + zs)


log('')
pa = probe_bauen(os.path.join(ARB, 'Probeexemplar_2.4_A_nicht_zurueckschreiben.docx'), va, ['2.4'])
log('Probeexemplar Variante A', os.path.basename(pa), os.path.getsize(pa), 'Byte (nur Arbeitsordner), 2.4 wortgleich, Überschriften 2.4.1 bis 2.4.3 entfernt')
log('  Manuskriptstand am Probeexemplar A:')
manuskriptstand(pa, 'Manuskriptstand_Probe_2.4_A.txt')
pb = probe_bauen(os.path.join(ARB, 'Probeexemplar_2.4_B_nicht_zurueckschreiben.docx'), vb, ALLE)
log('Probeexemplar Variante B', os.path.basename(pb), os.path.getsize(pb), 'Byte (nur Arbeitsordner), 2.4 bis 2.4.3 wortgleich')
log('  Manuskriptstand am Probeexemplar B:')
manuskriptstand(pb, 'Manuskriptstand_Probe_2.4_B.txt')

ea = os.path.join(ARB, 'Endabgleich_Probe_2.4_A')
subprocess.run([sys.executable, os.path.join(SKR, 'Endabgleich_Manuskript_2026-09-25.py'), pa,
                os.path.join(BEF, 'Kennzahlen_2026-09-25_Werte.csv'), os.path.join(BEF, 'Kennzahlen_2026-09-25.md'),
                KALT, os.path.join(SKR, 'Programmkennzahlen_2026-09-23.txt'), os.path.join(SKR, 'Objekte_2026-09-25'), ea],
               check=True, capture_output=True)
lauf = open(os.path.join(ea, 'Endabgleich_Manuskript_2026-09-25.txt'), encoding='utf-8').read()
log('Endabgleich am Probeexemplar A:')
for z in lauf.split('\n'):
    if z.startswith('Zahlen im Master') or z.startswith('Satzprüfungen:') or 'ABWEICHUNG' in z or z.strip().startswith('Wortprüfung') or z.startswith('Vorschläge'):
        log('  ' + z.strip()[:300])

with open(LOG, 'w', encoding='utf-8', newline='\n') as f:
    f.write('\n'.join(AUS) + '\n')
print('\n'.join(AUS))
