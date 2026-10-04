# -*- coding: utf-8 -*-
"""
Werkzeug_Objekte_Docx_2026-10-01.py — Word-Vorlage der Objekte nach dvs (2020) für das Einsetzen in Task 18
Bachelorarbeit U15-Plyometrie · Setzwerkzeug, kein Manuskripttext. Bildet keine Zahl, liest nur.

Eingang: Ausgabeordner von Objekte_2026-10-01.R (Objekte_2026-10-01_Liste.csv mit Kennzeichnung, Titel und
Anmerkung je Objekt, je Tabelle eine CSV, Abbildungen als PNG) und von Anhang_G_Diagramme_2026-10-01.R
(Anhang_G_2026-10-01_Liste.csv, PNG). Seitenformat, Ränder und Standard-Formatvorlage kommen aus einer
gestagten Kopie des Masters, der Master selbst wird nicht berührt.

Form nach dvs (2020) S. 2 bis 4 und Musterseite S. 10 mit der Klickfreigabe vom 01.10.2026
(05_Protokolle\\Pruefprotokoll_Objekte_dvs_2026-10-01, § 12):
  Formatvorlagen (KF10, gleichlautend für den Master in Task 18):
    „Tabellenüberschrift“: 10 pt kursiv, genau 12 pt, Blocksatz, Tabstopp und hängender Einzug 1,25 cm,
      vor 12, nach 6 pt, mit dem nächsten Absatz zusammen. „Tab. X.“ aufrecht (KF1), kein Schlusspunkt (KF2).
    „Tabellenüberschrift Anhang“: wie oben mit Tabstopp und hängendem Einzug 2,5 cm (KF12).
    „Abbildungsunterschrift“ und „Abbildungsunterschrift Anhang“: 10 pt aufrecht, genau 12 pt, Blocksatz,
      Tabstopp 1,25 bzw. 2,5 cm, vor 6, nach 12 pt, „Abb. X.“ kursiv (KF1).
    „Tabellentext“: 10 pt, genau 12 pt, Einzug links und rechts 0,1 cm.
    „Anmerkung“: „Anmerkung.“ kursiv, Text aufrecht, 10 pt, genau 12 pt, Blocksatz, vor 3 pt (KF3).
    „Abbildung“: zentriert (KF22), Einzug 0,1 cm, Rahmen ¾ pt, mit dem nächsten Absatz zusammen.
    „Abbildungsverzeichnis“: Tabstopp und hängender Einzug 2,5 cm, rechter Tabstopp mit Füllpunkten.
    Tabellenformatvorlage „Tabelle dvs“: Gitternetz, außen, unter dem Kopf und rechts der Vorspalte 1,5 pt,
      innen ¾ pt, Kopf 15 % grau und fett, Vorspalte fett ohne Schattierung (KF11), Zellinnenabstand 0 (KF19).
  Tabellen: Kopfzeilen wiederholt, Zeilen nicht über Seiten getrennt, Kopf zweistufig bei „Ober::Unter“,
    gleiche Werte der ersten Vorspalte senkrecht verbunden, Zellabsatz bei „<br>“, Index „~x~“ um 2 pt
    tiefgestellt (dvs S. 3), Zahlen rechtsbündig (Dezimal-Tabstopp abgeschaltet, siehe DEZIMAL_TABSTOPP), Zahl mit
    Intervall, Klammer oder Einheit geschützt (keine Trennung in der Zahl). Spaltenbreiten aus den
    Arial-Metriken, Tabelle über die Satzspiegelbreite 14,5 cm. Schmale Tabellen (natürliche Breite höchstens
    drei Viertel des Satzspiegels) behalten ihre Breite und stehen zentriert (KF22).
  Beschriftungen mit SEQ-Feldern (KF13): Tabelle, TabelleH (Unterbuchstaben mit \\c), Abbildung,
    AbbildungG, AbbildungH. Verzeichnisse über die Formatvorlagen (TOC \\t) am Anfang der Vorlage zur Probe.
  Sprache de-DE (Silbentrennung, Rechtschreibung).

Aufruf: python Werkzeug_Objekte_Docx_2026-10-01.py <Master-Kopie.docx> <Objekte-Ordner> <Anhang-G-Ordner> <Ziel.docx>
Ausgabe: Ziel.docx und <Ziel>_Spaltenbreiten.csv (je Tabelle Breiten und Status).
"""
import sys, os, re, csv
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageFont

MASTER, OBJ, GORD, ZIEL = sys.argv[1:5]
FASSUNG_OBJ = 'Objekte_2026-10-01'
FASSUNG_G = 'Anhang_G_2026-10-01'
NBSP = '\u00a0'
LEER = '\u2013'

d = Document(MASTER)
body = d.element.body
sectPr = body.find(qn('w:sectPr'))
for el in list(body):
    if el is not sectPr:
        body.remove(el)

# ------------------------------------------------------------------ Sprache de-DE
styles_el = d.styles.element
dd = styles_el.find(qn('w:docDefaults'))
rprd = dd.find(qn('w:rPrDefault')) if dd is not None else None
if rprd is not None:
    rpr = rprd.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr'); rprd.append(rpr)
    lang = rpr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang'); rpr.append(lang)
    lang.set(qn('w:val'), 'de-DE'); lang.set(qn('w:eastAsia'), 'de-DE')

# ------------------------------------------------------------------ Formatvorlagen
def stil(name, basis='Normal', groesse=10, kursiv=None, fett=None, vor=0, nach=0, genau=12, jc=WD_ALIGN_PARAGRAPH.JUSTIFY,
         tab_cm=None, haengend_cm=None, einzug_lr_cm=None, zusammen=False):
    vorhanden = [s for s in d.styles if s.name == name]
    s = vorhanden[0] if vorhanden else d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = d.styles[basis]
    s.font.name = 'Arial'
    s.font.size = Pt(groesse)
    if kursiv is not None:
        s.font.italic = kursiv
    if fett is not None:
        s.font.bold = fett
    pf = s.paragraph_format
    pf.space_before, pf.space_after = Pt(vor), Pt(nach)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(genau)
    pf.alignment = jc
    pf.keep_with_next = zusammen
    if tab_cm is not None:
        pf.tab_stops.add_tab_stop(Cm(tab_cm), WD_TAB_ALIGNMENT.LEFT)
    if haengend_cm is not None:
        pf.left_indent = Cm(haengend_cm)
        pf.first_line_indent = Cm(-haengend_cm)
    if einzug_lr_cm is not None:
        pf.left_indent = Cm(einzug_lr_cm)
        pf.right_indent = Cm(einzug_lr_cm)
    return s


stil('Tabellenüberschrift', kursiv=True, vor=12, nach=6, tab_cm=1.25, haengend_cm=1.25, zusammen=True)
stil('Tabellenüberschrift Anhang', kursiv=True, vor=12, nach=6, tab_cm=2.5, haengend_cm=2.5, zusammen=True)
stil('Abbildungsunterschrift', kursiv=False, vor=6, nach=12, tab_cm=1.25, haengend_cm=1.25)
stil('Abbildungsunterschrift Anhang', kursiv=False, vor=6, nach=12, tab_cm=2.5, haengend_cm=2.5)
st_txt = stil('Tabellentext', einzug_lr_cm=0.1, jc=WD_ALIGN_PARAGRAPH.LEFT)
_sah = OxmlElement('w:suppressAutoHyphens'); st_txt.element.get_or_add_pPr().insert(0, _sah)
stil('Anmerkung', vor=3, nach=0)
st_bild = stil('Abbildung', jc=WD_ALIGN_PARAGRAPH.CENTER, einzug_lr_cm=0.1, zusammen=True)
st_bild.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
ppr = st_bild.element.get_or_add_pPr()
bdr = OxmlElement('w:pBdr')
for seite in ('top', 'left', 'bottom', 'right'):
    b = OxmlElement('w:' + seite)
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '6'); b.set(qn('w:space'), '1'); b.set(qn('w:color'), '000000')
    bdr.append(b)
ppr.append(bdr)
# Verzeichniseinträge (TOC \t schreibt die Einträge in „Verzeichnis 1“, \c in „Abbildungsverzeichnis“, beide mit Tabstopp)
# Eingebaute Formatvorlage „Abbildungsverzeichnis“ (intern „table of figures“), die Word für TOC \t … \c verwendet
s = stil('table of figures', jc=WD_ALIGN_PARAGRAPH.LEFT, tab_cm=2.5, haengend_cm=2.5, genau=12)
s.paragraph_format.tab_stops.add_tab_stop(Cm(14.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
# als eingebaute Vorlage kennzeichnen (Kennung wie in deutschem Word, nicht benutzerdefiniert), sonst legt Word eine eigene an
s.element.set(qn('w:styleId'), 'Abbildungsverzeichnis')
if s.element.get(qn('w:customStyle')) is not None:
    del s.element.attrib[qn('w:customStyle')]

# Tabellenformatvorlage „Tabelle dvs“ (für den Master, die Tabellen tragen dieselbe Form zusätzlich direkt)
def tabellenstil():
    if any(s.name == 'Tabelle dvs' for s in d.styles):
        return
    st = OxmlElement('w:style'); st.set(qn('w:type'), 'table'); st.set(qn('w:customStyle'), '1'); st.set(qn('w:styleId'), 'Tabelledvs')
    nm = OxmlElement('w:name'); nm.set(qn('w:val'), 'Tabelle dvs'); st.append(nm)
    pp = OxmlElement('w:pPr')
    sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), '0'); sp.set(qn('w:after'), '0'); sp.set(qn('w:line'), '240'); sp.set(qn('w:lineRule'), 'exact'); pp.append(sp)
    ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '57'); ind.set(qn('w:right'), '57'); pp.append(ind)
    st.append(pp)
    rp = OxmlElement('w:rPr')
    f = OxmlElement('w:rFonts'); f.set(qn('w:ascii'), 'Arial'); f.set(qn('w:hAnsi'), 'Arial'); f.set(qn('w:cs'), 'Arial'); rp.append(f)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '20'); rp.append(sz)
    st.append(rp)
    tp = OxmlElement('w:tblPr')
    tb = OxmlElement('w:tblBorders')
    for s_, w_ in (('top', 12), ('left', 12), ('bottom', 12), ('right', 12), ('insideH', 6), ('insideV', 6)):
        b = OxmlElement('w:' + s_); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), str(w_)); b.set(qn('w:space'), '0'); b.set(qn('w:color'), '000000'); tb.append(b)
    tp.append(tb)
    mar = OxmlElement('w:tblCellMar')
    for s_ in ('left', 'right'):
        m = OxmlElement('w:' + s_); m.set(qn('w:w'), '0'); m.set(qn('w:type'), 'dxa'); mar.append(m)
    tp.append(mar)
    st.append(tp)
    for typ, fett, grau, linie in (('firstRow', True, True, 'bottom'), ('firstCol', True, False, 'right')):
        tsp = OxmlElement('w:tblStylePr'); tsp.set(qn('w:type'), typ)
        r_ = OxmlElement('w:rPr'); bb = OxmlElement('w:b'); r_.append(bb); tsp.append(r_)
        tcp = OxmlElement('w:tcPr')
        tcb = OxmlElement('w:tcBorders'); b = OxmlElement('w:' + linie); b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '12'); b.set(qn('w:space'), '0'); b.set(qn('w:color'), '000000'); tcb.append(b); tcp.append(tcb)
        if grau:
            shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'D9D9D9'); tcp.append(shd)
        tsp.append(tcp)
        st.append(tsp)
    styles_el.append(st)


tabellenstil()

# ------------------------------------------------------------------ Läufe, Felder, Schutz
def feld(p, instr, ergebnis='1', kursiv=None):
    def r(kind=None, text=None, it=False):
        run = OxmlElement('w:r')
        if kursiv is not None:
            rpr = OxmlElement('w:rPr'); i = OxmlElement('w:i'); i.set(qn('w:val'), '1' if kursiv else '0'); rpr.append(i); run.append(rpr)
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); run.append(fc)
        if it:
            t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = instr; run.append(t)
        if text is not None:
            t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text; run.append(t)
        p._p.append(run)
    r('begin'); r(it=True); r('separate'); r(text=ergebnis); r('end')


def tief(run):
    """Index nach dvs S. 3: 10 pt, um 2 pt tiefgestellt (w:position in halben Punkten)."""
    rpr = run._r.get_or_add_rPr()
    pos = OxmlElement('w:position'); pos.set(qn('w:val'), '-4'); rpr.append(pos)


AUFRECHT = re.compile(r'(k̄|Δ)')


def lauf_mit_trennstellen(p, text):
    """Lauf, in dem U+00AD als <w:softHyphen/> steht (Word zeigt das Zeichen sonst als Bindestrich)."""
    r = p.add_run('')
    teile = text.split(SHY)
    r.text = teile[0]
    for t in teile[1:]:
        r._r.append(OxmlElement('w:softHyphen'))
        el = OxmlElement('w:t'); el.set(qn('xml:space'), 'preserve'); el.text = t; r._r.append(el)
    return r


def text_laeufe(p, text, kursiv=None, fett=None, symbole_aufrecht=False):
    """Text mit „~x~“ als tiefgestellte Läufe, in kursiven Titeln Symbole aufrecht (KF4)."""
    for i, teil in enumerate(re.split(r'~([^~]+)~', text)):
        if not teil:
            continue
        if i % 2 == 1:
            r = p.add_run(teil); r.italic = kursiv; r.bold = fett; tief(r)
            continue
        stuecke = AUFRECHT.split(teil) if symbole_aufrecht else [teil]
        for s in stuecke:
            if not s:
                continue
            r = lauf_mit_trennstellen(p, s)
            r.italic = False if (symbole_aufrecht and AUFRECHT.fullmatch(s)) else kursiv
            if fett is not None:
                r.bold = fett


def schuetze_text(t):
    """Geschützte Leerzeichen zwischen Zahl und Einheit und in „n = 16“ (Befund nbsp)."""
    t = re.sub(r'(\d) (s|cm|%|AU|Wochen|Einheiten|Ziehungen)\b', lambda m: m.group(1) + NBSP + m.group(2), t)
    t = re.sub(r'\b([nNBWFtpdgbJ]|df|c~3~|R²) = (?=[\u2212+<]?\d)', lambda m: m.group(1) + NBSP + '=' + NBSP, t)
    t = re.sub(r'< (?=\d)', '<' + NBSP, t)
    t = t.replace(' = ', NBSP + '= ')
    return t


# ------------------------------------------------------------------ Zahlzellen und Breiten
F_R = ImageFont.truetype(r'C:\Windows\Fonts\arial.ttf', 1000)
F_B = ImageFont.truetype(r'C:\Windows\Fonts\arialbd.ttf', 1000)
PT = 10.0
EINZUG = 0.1 * 72 / 2.54                    # 0,1 cm in pt
POLSTER = 2 * EINZUG + 3.0                  # Einzug beidseitig, 3 pt für Linien und Rundung
VERFUEGBAR = 14.5 * 72 / 2.54
SCHMAL = 0.75
DEZ_VERSATZ = 0.2 * 72 / 2.54             # Versatz des Dezimal-Tabstopps in Word-Zellen (Probe p44_dezimalprobe.py)
# Dezimal-Tabstopp abgeschaltet: Word 2019 setzt das Komma in den Zellen nicht reproduzierbar an den Tabstopp
# (Probe p44 ohne Formatvorlage +0,1 cm, Vorlage mit „Tabellentext“ +0,3 cm, Folge: Umbruch in der Zahl). Zahlen stehen
# rechtsbündig, bei gleicher Stellenzahl stehen die Kommas untereinander (Ziffern in Arial gleich breit). Offen für Task 18.
DEZIMAL_TABSTOPP = False
RE_DEZ = re.compile(r'^[\[(]?(<\s?|[\u2212+])?\d')        # Absatz einer Datenzelle beginnt mit einer Zahl (auch Intervall, Klammer, „< 0,001“)
# Vorspalten: führende Spalten mit diesen Köpfen (die Zielgrößen beginnen teils mit Ziffern und sind trotzdem Text)
VORSPALTEN_KOPF = {'Zielgröße', 'Grund', 'Kennwert', 'Prüfung', 'Menge', 'Gruppe', 'Variante', 'Spieler', 'Programmwoche', 'Vollständige Einheiten'}
RE_DEZKOMMA = re.compile(r'\d,\d')
RE_EINZELZAHL = re.compile(r'^[\u2212+]?\d+(,\d+)?$')
breiten_protokoll = []


SHY = '­'
# Weiche Trennstellen in langen Wörtern von Kopf und Vorspalte (Word trennt nur dort, wenn die Spalte es verlangt)
TRENNSTELLEN = {'Standweitsprung': 'Standweit' + SHY + 'sprung', 'Eingangstestung': 'Eingangs' + SHY + 'testung',
                'Abschlusstestung': 'Abschluss' + SHY + 'testung', 'Ausgangswert': 'Ausgangs' + SHY + 'wert',
                'Abschlusswert': 'Abschluss' + SHY + 'wert', 'Seitenmittel': 'Seiten' + SHY + 'mittel',
                'Überlappung': 'Über' + SHY + 'lappung', 'Programmwoche': 'Programm' + SHY + 'woche',
                'Beanspruchung': 'Bean' + SHY + 'spruchung', 'Erhebungsanteil': 'Erhebungs' + SHY + 'anteil',
                'Einheitennummern': 'Einheiten' + SHY + 'nummern', 'Familiarisierungsangabe': 'Familiarisierungs' + SHY + 'angabe',
                'Familiarisierungstermine': 'Familiarisierungs' + SHY + 'termine', 'Lichtschranke': 'Licht' + SHY + 'schranke',
                'Fehlversuch': 'Fehl' + SHY + 'versuch', 'wiederholbarer': 'wieder' + SHY + 'holbarer',
                'Änderungswertmodell': 'Änderungswert' + SHY + 'modell', 'Sensitivitätsanalyse': 'Sensitivitäts' + SHY + 'analyse',
                'Interaktionsterm': 'Interaktions' + SHY + 'term', 'Ausgangswerte': 'Ausgangs' + SHY + 'werte'}


def trennbar(text):
    for w_, t_ in TRENNSTELLEN.items():
        text = text.replace(w_, t_)
    return text


def laenge(text, fett=False):
    return (F_B if fett else F_R).getlength(text.replace('~', '').replace(SHY, '')) * PT / 1000


def kleinste(text, fett):
    """Breite des längsten unteilbaren Stücks: Umbruch an Leerzeichen, nach Bindestrich und an weichen Trennstellen (mit Trennstrich)."""
    best = 0.0
    for wort in re.split(r'[ ]+', text):
        for teil in re.split(r'(?<=-)(?=\S)', wort):
            stuecke = teil.split(SHY)
            for k, st in enumerate(stuecke):
                best = max(best, laenge(st + ('-' if k < len(stuecke) - 1 else ''), fett))
    return best


def dezimal_absatz(a):
    return bool(RE_DEZ.match(a))


def teile_am_komma(a):
    m = RE_DEZKOMMA.search(a)
    if not m:
        return a, ''
    k = m.start() + 1
    return a[:k], a[k:]


def absaetze(wert):
    return wert.split('<br>')


def spaltenbreiten(kopf_unten, kopf_oben, zeilen, n_label, name):
    n = len(kopf_unten)
    nat, mind, dez = [], [], []
    for j in range(n):
        # Kopf: natürlich ganz, kleinste Breite = längstes unteilbares Stück (fett). Die Oberzeile zählt je Gruppe (unten).
        kt = kopf_unten[j] if kopf_unten[j] else kopf_oben[j]
        kn = max((laenge(a, True) for a in absaetze(kt)), default=0)
        km = max((kleinste(a, True) for a in absaetze(kt)), default=0)
        if j < n_label:
            zn = max((laenge(a, True) for z in zeilen for a in absaetze(z[j])), default=0)
            zm = max((kleinste(a, True) for z in zeilen for a in absaetze(z[j])), default=0)
            nat.append(max(kn, zn) + POLSTER); mind.append(max(km, zm) + POLSTER); dez.append(None)
            continue
        zahl_abs = [a for z in zeilen for a in absaetze(z[j]) if dezimal_absatz(a)]
        # Dezimal-Tabstopp nur in Spalten mit reinen Einzelzahlen mit Dezimalkomma (Probe im Prüfexemplar: in zusammengesetzten
        # Zellen bricht Word die Zahl um). Zusammengesetzte Zahlzellen stehen rechtsbündig, bei gleicher Stellenzahl stehen die
        # Kommas damit ebenfalls untereinander (Ziffern in Arial gleich breit).
        hat_dez = DEZIMAL_TABSTOPP and bool(zahl_abs) and all(RE_EINZELZAHL.match(a) for a in zahl_abs) and any(RE_DEZKOMMA.search(a) for a in zahl_abs)
        links = max((laenge(teile_am_komma(a)[0]) for a in zahl_abs), default=0) if hat_dez else 0.0
        rechts = max((laenge(teile_am_komma(a)[1]) for a in zahl_abs), default=0) if hat_dez else 0.0
        andere = max((laenge(a) for z in zeilen for a in absaetze(z[j]) if not (hat_dez and dezimal_absatz(a))), default=0)
        bedarf = max(links + rechts + EINZUG + 2.0 if hat_dez else 0.0, andere)
        nat.append(max(kn, bedarf) + POLSTER); mind.append(max(km, bedarf) + POLSTER)
        dez.append((links, rechts) if hat_dez else None)
    # Oberzeile: die Gruppe muss die Oberzeile tragen (natürlich ganz, kleinste Breite längstes Stück)
    j = 0
    while j < n:
        if not kopf_oben[j]:
            j += 1
            continue
        k = j
        while k + 1 < n and kopf_oben[k + 1] == kopf_oben[j]:
            k += 1
        if k > j:
            for liste_, wert in ((nat, laenge(kopf_oben[j], True) + POLSTER), (mind, kleinste(kopf_oben[j], True) + POLSTER)):
                fehlt = wert - sum(liste_[j:k + 1])
                if fehlt > 0:
                    for jj in range(j, k + 1):
                        liste_[jj] += fehlt / (k - j + 1)
        j = k + 1
    s_nat, s_min = sum(nat), sum(mind)
    if s_nat <= SCHMAL * VERFUEGBAR:
        status, breiten, zentriert = 'schmal, natürliche Breite, zentriert', nat, True
    elif s_nat <= VERFUEGBAR:
        rest = VERFUEGBAR - s_nat
        zahl = [j for j in range(n) if j >= n_label] or list(range(n))
        breiten = [b + (rest / len(zahl) if j in zahl else 0) for j, b in enumerate(nat)]
        status, zentriert = 'passt ohne Umbruch, Satzspiegelbreite', False
    elif s_min <= VERFUEGBAR:
        rest = VERFUEGBAR - s_min
        spiel = [a - b for a, b in zip(nat, mind)]
        bedarf_v = sum(spiel[:n_label])
        breiten = list(mind)
        if rest >= bedarf_v:
            for jj in range(n_label):
                breiten[jj] = nat[jj]
            rest2 = rest - bedarf_v
            sp2 = sum(spiel[n_label:])
            for jj in range(n_label, n):
                breiten[jj] += rest2 * (spiel[jj] / sp2 if sp2 > 0 else 1 / (n - n_label))
        else:
            for jj in range(n_label):
                breiten[jj] += rest * spiel[jj] / bedarf_v
        status, zentriert = 'passt mit Umbruch in Kopf und Textzellen, Satzspiegelbreite', False
    else:
        breiten, zentriert = mind, False
        status = 'PASST NICHT (kleinste Breite über dem Satzspiegel)'
    breiten_protokoll.append({'objekt': name, 'spalten': n, 'natuerlich_cm': round(s_nat / 72 * 2.54, 2), 'kleinste_cm': round(s_min / 72 * 2.54, 2),
                              'gesetzt_cm': round(sum(breiten) / 72 * 2.54, 2), 'status': status, 'breiten_cm': [round(b / 72 * 2.54, 2) for b in breiten]})
    return breiten, dez, zentriert


def rand(el, seite, sz):
    b = OxmlElement('w:' + seite)
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), str(sz)); b.set(qn('w:space'), '0'); b.set(qn('w:color'), '000000')
    el.append(b)


def zelle_setzen(c, wert, fett, rechts, dez_pos, kopf):
    c.text = ''
    for k, a in enumerate(absaetze(wert)):
        p = c.paragraphs[0] if k == 0 else c.add_paragraph()
        p.style = d.styles['Tabellentext']
        if kopf:
            text_laeufe(p, a, fett=True)
            if rechts:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            continue
        if dez_pos is not None and RE_EINZELZAHL.match(a):
            p.paragraph_format.tab_stops.add_tab_stop(Pt(dez_pos), WD_TAB_ALIGNMENT.DECIMAL)
            text_laeufe(p, '\t' + a.replace(' ', NBSP), fett=fett)
        elif rechts:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            text_laeufe(p, a.replace(' ', NBSP) if dezimal_absatz(a) else a, fett=fett)
        else:
            text_laeufe(p, a, fett=fett)


def tabelle(kopf, zeilen, name):
    kopf = [trennbar(k) for k in kopf]
    kopf_oben = [k.split('::')[0] if '::' in k else '' for k in kopf]
    kopf_unten = [k.split('::')[1] if '::' in k else k for k in kopf]
    zweistufig = any(kopf_oben)
    n = len(kopf)
    n_label = 0
    while n_label < n and kopf_unten[n_label].replace(SHY, '') in VORSPALTEN_KOPF and not kopf_oben[n_label]:
        n_label += 1
    n_label = max(n_label, 1)
    for z in zeilen:
        for j in range(n):
            if j < n_label:
                z[j] = trennbar(z[j])
            else:
                z[j] = '<br>'.join(a if dezimal_absatz(a) else trennbar(a) for a in absaetze(z[j]))
    breiten, dez, zentriert = spaltenbreiten(kopf_unten, kopf_oben, zeilen, n_label, name)
    nk = 2 if zweistufig else 1
    t = d.add_table(rows=nk + len(zeilen), cols=n)
    t.style = d.styles['Tabelle dvs']
    t.alignment = WD_TABLE_ALIGNMENT.CENTER if zentriert else WD_TABLE_ALIGNMENT.LEFT
    tbl = t._tbl
    tpr = tbl.tblPr
    tw = tpr.find(qn('w:tblW'))
    if tw is None:
        tw = OxmlElement('w:tblW'); tpr.append(tw)
    twips = [int(round(b * 20)) for b in breiten]
    tw.set(qn('w:w'), str(sum(twips))); tw.set(qn('w:type'), 'dxa')
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tpr.append(lay)
    mar = OxmlElement('w:tblCellMar')
    for seite in ('left', 'right'):
        m = OxmlElement('w:' + seite); m.set(qn('w:w'), '0'); m.set(qn('w:type'), 'dxa'); mar.append(m)
    tpr.append(mar)
    look = tpr.find(qn('w:tblLook'))
    if look is not None:
        look.set(qn('w:firstRow'), '1'); look.set(qn('w:firstColumn'), '1'); look.set(qn('w:noVBand'), '1'); look.set(qn('w:noHBand'), '1')
    for gc, b in zip(tbl.tblGrid.findall(qn('w:gridCol')), twips):
        gc.set(qn('w:w'), str(b))
    for tr in t.rows:
        for c, b in zip(tr.cells, twips):
            tcw = c._tc.get_or_add_tcPr().get_or_add_tcW()
            tcw.set(qn('w:w'), str(b)); tcw.set(qn('w:type'), 'dxa')
    tb = OxmlElement('w:tblBorders')
    for s_ in ('top', 'left', 'bottom', 'right'):
        rand(tb, s_, 12)
    rand(tb, 'insideH', 6); rand(tb, 'insideV', 6)
    tpr.append(tb)
    # Dezimal-Tabstopp: rechtsbündig am Komma, Position vom Zellrand
    dez_pos = []
    for j in range(n):
        if dez[j] is None:
            dez_pos.append(None)
        else:
            # Probe p44 (Word 2019): das Komma steht rund 0,2 cm rechts der eingestellten Position, gemessen vom Zellrand
            dez_pos.append(breiten[j] - EINZUG - dez[j][1] - DEZ_VERSATZ - 3.0)
    # Inhalte
    for i, tr in enumerate(t.rows):
        trpr = tr._tr.get_or_add_trPr()
        cs = OxmlElement('w:cantSplit'); trpr.append(cs)
        if i < nk:
            h = OxmlElement('w:tblHeader'); trpr.append(h)
    for j in range(n):
        if zweistufig:
            zelle_setzen(t.cell(0, j), kopf_oben[j] if kopf_oben[j] else kopf_unten[j], True, j >= n_label, None, True)
            if kopf_oben[j]:
                for pp in t.cell(0, j).paragraphs:
                    pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            zelle_setzen(t.cell(1, j), kopf_unten[j] if kopf_oben[j] else '', True, j >= n_label, None, True)
        else:
            zelle_setzen(t.cell(0, j), kopf_unten[j], True, j >= n_label, None, True)
    for i, z in enumerate(zeilen):
        for j in range(n):
            zelle_setzen(t.cell(nk + i, j), z[j], j < n_label, j >= n_label, dez_pos[j], False)
    # Schattierung, Kopflinie, Vorspaltenlinie
    for i, tr in enumerate(t.rows):
        for j, c in enumerate(tr.cells):
            tcpr = c._tc.get_or_add_tcPr()
            if i < nk:
                shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'D9D9D9'); tcpr.append(shd)
            tcb = OxmlElement('w:tcBorders')
            if i == nk - 1:
                rand(tcb, 'bottom', 12)
            if j == n_label - 1:
                rand(tcb, 'right', 12)
            if len(tcb):
                tcpr.append(tcb)
    # Verbinden: zweistufiger Kopf, gleiche Werte der ersten Vorspalte
    if zweistufig:
        j = 0
        while j < n:
            if not kopf_oben[j]:
                t.cell(0, j).merge(t.cell(1, j))
                j += 1
                continue
            k = j
            while k + 1 < n and kopf_oben[k + 1] == kopf_oben[j]:
                k += 1
            if k > j:
                oben = t.cell(0, j).merge(t.cell(0, k))
                for extra in oben.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
            j = k + 1
    if n_label >= 2:
        i = 0
        while i < len(zeilen):
            k = i
            while k + 1 < len(zeilen) and zeilen[k + 1][0] == zeilen[i][0]:
                k += 1
            if k > i:
                c = t.cell(nk + i, 0).merge(t.cell(nk + k, 0))
                for extra in c.paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
            i = k + 1
    return t


# ------------------------------------------------------------------ Beschriftungen, Anmerkung, Abbildung
def kennung_teile(k):
    m = re.match(r'^(Tab|Abb)\. ([A-Z]?)(\d+)([a-e]?)\.$', k)
    if not m:
        raise SystemExit('Kennzeichnung nicht lesbar: ' + k)
    return m.groups()


SEQ_NAME = {('Tab', ''): 'Tabelle', ('Tab', 'H'): 'TabelleH', ('Abb', ''): 'Abbildung', ('Abb', 'H'): 'AbbildungH', ('Abb', 'G'): 'AbbildungG'}
zaehler = {}
letzte_nummer = {}


def beschriftung(kennzeichnung, titel, nach_anmerkung=False):
    art, anhang, nr, buchst = kennung_teile(kennzeichnung)
    seq = SEQ_NAME[(art, anhang)]
    # SEQ-Feld: gleiche Nummer mit \c für die Unterbuchstaben, sonst fortlaufend, Probe gegen die Kennzeichnung
    gleich = letzte_nummer.get(seq) == nr
    if not gleich:
        zaehler[seq] = zaehler.get(seq, 0) + 1
    letzte_nummer[seq] = nr
    if str(zaehler[seq]) != nr:
        raise SystemExit(f'SEQ-Folge {seq} ergibt {zaehler[seq]}, Kennzeichnung {kennzeichnung}')
    stilname = ('Tabellenüberschrift' if art == 'Tab' else 'Abbildungsunterschrift') + (' Anhang' if anhang else '')
    p = d.add_paragraph(style=stilname)
    kursiv_kenn = art == 'Abb'
    r = p.add_run(f'{art}. {anhang}'); r.italic = kursiv_kenn
    feld(p, f' SEQ {seq} \\* ARABIC ' + ('\\c ' if gleich else ''), ergebnis=nr, kursiv=kursiv_kenn)
    r = p.add_run(f'{buchst}.'); r.italic = kursiv_kenn
    p.add_run('\t')
    text_laeufe(p, titel, kursiv=(art == 'Tab'), symbole_aufrecht=(art == 'Tab'))
    if nach_anmerkung:
        p.paragraph_format.space_after = Pt(0)
    return p


def anmerkung(text, nach_pt=None):
    p = d.add_paragraph(style='Anmerkung')
    r = p.add_run('Anmerkung.'); r.italic = True
    text_laeufe(p, ' ' + schuetze_text(text))
    if nach_pt is not None:
        p.paragraph_format.space_after = Pt(nach_pt)
    return p


def leerzeile():
    d.add_paragraph(style='Normal')


def abbildung(pfad):
    im = Image.open(pfad)
    dpi = im.info.get('dpi', (300, 300))[0]
    breite_cm = im.size[0] / dpi * 2.54
    if breite_cm > 14.26:
        raise SystemExit(f'Abbildung breiter als 14,25 cm: {pfad}')
    leerzeile()
    p = d.add_paragraph(style='Abbildung')
    p.add_run().add_picture(pfad, width=Cm(breite_cm))
    return p


def seite():
    p = d.add_paragraph(style='Normal')
    p.add_run().add_break(WD_BREAK.PAGE)


def hinweis(text):
    p = d.add_paragraph(style='Normal')
    p.add_run(text).bold = True


# ------------------------------------------------------------------ Objekte lesen
def liste(pfad):
    with open(pfad, encoding='utf-8') as f:
        return {r['kennzeichnung']: r for r in csv.DictReader(f)}


L = liste(os.path.join(OBJ, FASSUNG_OBJ + '_Liste.csv'))
LG = liste(os.path.join(GORD, FASSUNG_G + '_Liste.csv'))


def tabelle_lesen(datei):
    with open(os.path.join(OBJ, datei), encoding='utf-8') as f:
        r = list(csv.reader(f))
    return r[0], r[1:]


def setze_tabelle(kenn):
    o = L[kenn]
    beschriftung(kenn, o['titel'])
    kopf, zeilen = tabelle_lesen(o['datei'])
    tabelle(kopf, zeilen, kenn)
    anmerkung(o['anmerkung'])
    leerzeile()


def setze_abbildung(kenn, quelle, ordner):
    o = quelle[kenn]
    abbildung(os.path.join(ordner, o['datei']))
    beschriftung(kenn, o['titel'], nach_anmerkung=bool(o['anmerkung']))
    if o['anmerkung']:
        anmerkung(o['anmerkung'], nach_pt=12)


# ------------------------------------------------------------------ Dokument
hinweis('Word-Vorlage der Objekte nach dvs (2020), erzeugt mit Werkzeug_Objekte_Docx_2026-10-01.py aus Objekte_2026-10-01 und Anhang_G_2026-10-01. Kein Manuskripttext. Zum Einsetzen in Task 18: Formatvorlagen und Objekte übernehmen, Verzeichnisse im Master über die Formatvorlagen führen, erst SEQ-Felder, dann Verzeichnisse aktualisieren.')
# \t mit leerem \c: Word führt das Feld als Abbildungsverzeichnis, Einträge in der Formatvorlage „Abbildungsverzeichnis“
for instr, kopf in (('TOC \\h \\z \\t "Tabellenüberschrift;1;Tabellenüberschrift Anhang;1" \\c', 'Tabellenverzeichnis (Probe, TOC \\t mit \\c)'),
                    ('TOC \\h \\z \\t "Abbildungsunterschrift;1;Abbildungsunterschrift Anhang;1" \\c', 'Abbildungsverzeichnis (Probe, TOC \\t mit \\c)')):
    hinweis(kopf)
    p = d.add_paragraph(style='Normal')
    feld(p, ' ' + instr + ' ', ergebnis='(Verzeichnis wird von Word aktualisiert)')
seite()
hinweis('Textteil (Reihenfolge der Erstverweise: Tab. 1 in 4.4, Abb. 1 und Tab. 2 in Kapitel 5 A1, Tab. 3 und Abb. 2 in A5)')
setze_tabelle('Tab. 1.')
setze_abbildung('Abb. 1.', L, OBJ)
setze_tabelle('Tab. 2.')
setze_tabelle('Tab. 3.')
setze_abbildung('Abb. 2.', L, OBJ)
seite()
hinweis('Anhang G (Diagramme der Voraussetzungsprüfung)')
for k in sorted(LG, key=lambda x: int(re.search(r'\d+', x).group())):
    setze_abbildung(k, LG, GORD)
seite()
hinweis('Anhang H (Ergänzende Tabellen und Abbildung)')
for k in [k for k in L if k.startswith('Tab. H')]:
    setze_tabelle(k)
for k in [k for k in L if k.startswith('Abb. H')]:
    setze_abbildung(k, L, OBJ)

fehlend = set(L) - {k for k in L if k.startswith(('Tab. H', 'Abb. H'))} - {'Tab. 1.', 'Tab. 2.', 'Tab. 3.', 'Abb. 1.', 'Abb. 2.'}
if fehlend:
    raise SystemExit('Objekte nicht gesetzt: ' + ', '.join(sorted(fehlend)))
d.save(ZIEL)
with open(os.path.splitext(ZIEL)[0] + '_Spaltenbreiten.csv', 'w', encoding='utf-8-sig', newline='') as f:
    wr = csv.writer(f, delimiter=';')
    kk = ['objekt', 'spalten', 'natuerlich_cm', 'kleinste_cm', 'gesetzt_cm', 'status']
    wr.writerow(kk + ['breiten_cm'])
    for e in breiten_protokoll:
        wr.writerow([e[k] for k in kk] + [' '.join(str(x).replace('.', ',') for x in e['breiten_cm'])])
nicht = [e['objekt'] for e in breiten_protokoll if e['status'].startswith('PASST NICHT')]
print('geschrieben:', ZIEL, '· Tabellen:', len(breiten_protokoll), '· Abbildungen:', sum(1 for k in L if k.startswith('Abb')) + len(LG),
      '· passt nicht:', ', '.join(nicht) if nicht else 'keine')
for e in breiten_protokoll:
    print(f"  {e['objekt']:<11} {e['spalten']:>2} Sp. natürlich {e['natuerlich_cm']:>6} cm, kleinste {e['kleinste_cm']:>6} cm, gesetzt {e['gesetzt_cm']:>6} cm · {e['status']}")
