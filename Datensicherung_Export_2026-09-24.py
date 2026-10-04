# -*- coding: utf-8 -*-
"""
Datensicherung_Export_2026-09-24.py

Zweck:    Friert den Datenstand für die blinde R-Rechnung ein: nur die Analysevariablen nach
          Spezifikation_2026-09-24 Abschnitt 0.4 als CSV, dazu Datenwörterbuch und SHA-256-
          Prüfsummen. Vor dem Einfrieren zwei Proben: (1) Rückleseprobe, die CSV-Dateien
          zurückgelesen ergeben Wert für Wert den Inhalt von Workbook und JSON, (2) Abbruch-
          probe, alle Abbruchregeln der Spezifikation (0.5 Nr. 10, S01 Regeln 2 bis 4, S03 mit
          K5, S06 Regeln 3 und 7) laufen am Datenstand ohne Verstoß. Erst dann die Prüfsummen.
Schritt:  Auswertungsverfahren_2026-09-24, Phase 1, Schritt 1.5
Eingang:  Statistik/Studiendaten_U15_gesamt.xlsx, Blätter 01_Personen und 02_Rohdaten, Stand
              nach den Korrekturen aus Schritt 1.4 (Prüfsumme im Laufprotokoll unten)
          SoSci_Fragebogen/data_test546007_2026-08-31_11-37.json (Originalexport, maßgeblich)
              SHA-256 7faf00a66b1bfcdfe78e100c8259b81b426bd797c006d10d756f243c1b5b40ab
          SoSci_Fragebogen/Fragebogen Datensatz.txt (nur zum Gegenlesen der elf Spalten)
              SHA-256 8e827b4eee9e637a783df819f4beba737b4f842fcfe0c2e27e86f4f8c264465c
          Claude/02_Befunde/Spezifikation_2026-09-24_Vokabular.csv und _Fragebogen.csv
              (nur für die Abbruchprobe)
Fassung:  2026-09-24, Fassung 1 (Claude, Task Datensicherung)
Aufruf:   python Datensicherung_Export_2026-09-24.py [Basisordner] [--datum JJJJ-MM-TT]
          Ziel: Statistik/Datenstand_JJJJ-MM-TT/ (Datum des Einfrierens, Voreinstellung heute).
          Ein vorhandener Zielordner wird nicht überschrieben (Abbruch).
Regeln:   Keine Namen, kein Geburtsdatum, kein Testdatum, keine Freitexte in Datenstand,
          Protokoll oder Konsole. Das Alter wird aus den beiden Datumsfeldern gerechnet.
          Keine Kennzahlen im Datenwörterbuch.
Format:   UTF-8 ohne BOM, Komma als Trennzeichen, Dezimalpunkt, Zeilenende LF, fehlender Wert
          leer, ganze Zahlen ohne Nachkommastellen, Dezimalzahlen ungerundet in der kürzesten
          Darstellung, die denselben Gleitkommawert ergibt (repr).
"""

import sys, os, io, re, csv, json, hashlib, datetime, collections

FASSUNG = "2026-09-24, Fassung 1"
args = [a for a in sys.argv[1:] if not a.startswith("--")]
BASIS = os.path.abspath(args[0]) if args and os.path.isdir(args[0]) else os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
DATUM = sys.argv[sys.argv.index("--datum") + 1] if "--datum" in sys.argv else datetime.date.today().strftime("%Y-%m-%d")
P_WB = os.path.join(BASIS, "Statistik", "Studiendaten_U15_gesamt.xlsx")
P_JSON = os.path.join(BASIS, "SoSci_Fragebogen", "data_test546007_2026-08-31_11-37.json")
P_TXT = os.path.join(BASIS, "SoSci_Fragebogen", "Fragebogen Datensatz.txt")
P_VOK = os.path.join(BASIS, "Claude", "02_Befunde", "Spezifikation_2026-09-24_Vokabular.csv")
P_FB = os.path.join(BASIS, "Claude", "02_Befunde", "Spezifikation_2026-09-24_Fragebogen.csv")
ZIEL = os.path.join(BASIS, "Statistik", "Datenstand_" + DATUM)
P_OUT = os.path.splitext(os.path.abspath(__file__))[0] + ".txt"

out = io.StringIO()
def w(s=""): out.write(s + "\n")
def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""): h.update(chunk)
    return h.hexdigest()
def ende(code):
    text = out.getvalue()
    with open(P_OUT, "w", encoding="utf-8", newline="\n") as f: f.write(text)
    print(text); sys.exit(code)
def fmt(v):
    """Zahl in der kürzesten Darstellung, ganze Zahlen ohne Nachkommastellen, leer bleibt leer."""
    if v is None or v == "": return ""
    if isinstance(v, bool): raise ValueError("bool")
    if isinstance(v, int): return str(v)
    if isinstance(v, float): return str(int(v)) if v.is_integer() else repr(v)
    return str(v)

w("Datensicherung, Schritt 1.5 Datenstand einfrieren")
w("Skriptfassung " + FASSUNG + " · Lauf " + datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S") + " · Datenstand " + DATUM)
w("Basisordner: " + BASIS)
w("Eingang mit SHA-256:")
for p in (P_WB, P_JSON, P_TXT, P_VOK, P_FB):
    w("  " + os.path.relpath(p, BASIS) + "  " + sha256(p) + "  " + str(os.path.getsize(p)) + " B")
w()
if os.path.exists(ZIEL):
    w("ABBRUCH: Zielordner existiert schon: " + ZIEL + " (ein späterer Stand bekommt ein neues Datum)"); ende(2)

# ----------------------------------------------------------------------------
# Workbook lesen
# ----------------------------------------------------------------------------
import openpyxl
wb = openpyxl.load_workbook(P_WB, data_only=False)
CODE_RE = re.compile(r"^(HL|BW|VS)-\d{2}$")
VEREIN_MAP = {"Hohenlind": "A", "Blau-Weiß Köln": "B", "Vorwärts Spoho": "C"}
GRUPPE_MAP = {"Intervention": "IG", "Kontrolle": "KG"}
ws = wb["01_Personen"]
hdr = [c.value for c in ws[1]]
col = {h: i + 1 for i, h in enumerate(hdr) if h is not None}
personen = []
for r in range(2, ws.max_row + 1):
    code = ws.cell(r, col["Code"]).value
    if not (isinstance(code, str) and CODE_RE.match(code.strip())):
        continue
    g = lambda h: ws.cell(r, col[h]).value
    ver, gr, st = g("Verein"), g("Gruppe"), g("Status")
    if ver not in VEREIN_MAP: w("ABBRUCH: Verein %r bei %s" % (ver, code)); ende(2)
    if gr not in GRUPPE_MAP: w("ABBRUCH: Gruppe %r bei %s" % (gr, code)); ende(2)
    if st == "ausgewertet": st_neu = "ausgewertet"
    elif isinstance(st, str) and st.startswith("Post-Testung nicht angetreten"): st_neu = "nicht angetreten"
    else: w("ABBRUCH: Status %r bei %s außerhalb der festen Umsetzung" % (st, code)); ende(2)
    gd, tp = g("Geburtsdatum"), g("Testdatum_prä")
    if not (isinstance(gd, (datetime.date, datetime.datetime)) and isinstance(tp, (datetime.date, datetime.datetime))):
        w("ABBRUCH: Geburtsdatum oder Testdatum prä fehlt bei %s" % code); ende(2)
    gd = gd.date() if isinstance(gd, datetime.datetime) else gd
    tp = tp.date() if isinstance(tp, datetime.datetime) else tp
    alter = (tp - gd).days / 365.25
    for h in ("Größe_prä_cm", "Gewicht_prä_kg", "Größe_Mutter_cm", "Größe_Vater_cm", "Familiarisierung"):
        v = g(h)
        if v not in (None, "") and (isinstance(v, str) or isinstance(v, bool)):
            w("ABBRUCH: %s bei %s ist Text: %r" % (h, code, v)); ende(2)
    fam = g("Familiarisierung")
    if fam not in (None, "", 1, 2): w("ABBRUCH: Familiarisierung %r bei %s" % (fam, code)); ende(2)
    personen.append({"Code": code.strip(), "Verein": VEREIN_MAP[ver], "Gruppe": GRUPPE_MAP[gr], "Alter_prae": alter,
                     "Koerperhoehe_prae": g("Größe_prä_cm"), "Koerpermasse_prae": g("Gewicht_prä_kg"),
                     "Groesse_Mutter": g("Größe_Mutter_cm"), "Groesse_Vater": g("Größe_Vater_cm"),
                     "Familiarisierung": fam, "Status": st_neu, "_zeile": r})
ws2 = wb["02_Rohdaten"]
SOLL2 = ["Code", "Verein", "Zeitpunkt", "Test", "Seite", "Versuch", "Wert", "Einheit", "ungültig", "Bemerkung", "gilt"]
if [c.value for c in ws2[1]][:11] != SOLL2: w("ABBRUCH: Kopfzeile 02_Rohdaten"); ende(2)
versuche = []
for r in range(2, ws2.max_row + 1):
    vals = [ws2.cell(r, c).value for c in range(1, 11)]
    if all(v in (None, "") for v in vals): continue
    d = dict(zip(SOLL2[:10], vals)); d["_zeile"] = r
    if isinstance(d["Wert"], (str, bool)): w("ABBRUCH: Wert ist Text in Zeile %d" % r); ende(2)
    versuche.append(d)

# ----------------------------------------------------------------------------
# Fragebogen lesen
# ----------------------------------------------------------------------------
with open(P_JSON, encoding="utf-8") as f:
    js = json.load(f)
COLS = ["CASE", "STARTED", "H010", "H002", "H003", "H004", "H005", "H006", "H007", "H008", "H009"]
meldungen = []
for key, rec in js["data"].items():
    row = {}
    for c in COLS:
        if c not in rec: w("ABBRUCH: Spalte %s fehlt bei CASE %s" % (c, rec.get("CASE"))); ende(2)
        row[c] = rec[c]
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", str(row["STARTED"])):
        w("ABBRUCH: STARTED %r nicht im Format JJJJ-MM-TT hh:mm:ss" % row["STARTED"]); ende(2)
    meldungen.append(row)
meldungen.sort(key=lambda r: int(r["CASE"]))

# Gegenlesen gegen die TXT (Excel-gespeicherter Export, STARTED auf Minutenebene)
with open(P_TXT, "rb") as f:
    raw = f.read()
txt = raw.decode("utf-16")
zeilen = [z for z in txt.splitlines() if z.strip()]
kopf = zeilen[0].split("\t")
idx = {h: i for i, h in enumerate(kopf)}
txt_rows = {}
for z in zeilen[1:]:
    t = z.split("\t")
    txt_rows[int(t[idx["CASE"]])] = t
abw = 0
for row in meldungen:
    t = txt_rows.get(int(row["CASE"]))
    if t is None: w("TXT: CASE %s fehlt" % row["CASE"]); abw += 1; continue
    for c in COLS:
        a = str(row[c]); b = t[idx[c]]
        if c == "STARTED":
            a = datetime.datetime.strptime(a, "%Y-%m-%d %H:%M:%S").strftime("%d.%m.%Y %H:%M")
        if a != b:
            w("TXT-Abweichung CASE %s %s: JSON %r, TXT %r" % (row["CASE"], c, a, b)); abw += 1
if set(txt_rows) != {int(r["CASE"]) for r in meldungen}:
    w("TXT: CASE-Mengen weichen ab"); abw += 1
w("Gegenlesen JSON gegen TXT, elf Spalten (STARTED auf Minutenebene): %d Abweichungen" % abw)

# ----------------------------------------------------------------------------
# Zuordnungstabelle und Listenplatz (Verfasser per Klick am 24.09.2026 bestätigt)
# ----------------------------------------------------------------------------
ZUORDNUNG = [("HL-01", "HL-01"), ("HL-02", "HL-02"), ("HL-03", "HL-03"), ("HL-04", "HL-05"), ("HL-05", ""),
             ("HL-06", "HL-06"), ("HL-07", "HL-07"), ("HL-08", "HL-08"),
             ("BW-01", "BW-01"), ("BW-02", "BW-02"), ("BW-03", ""), ("BW-04", "BW-04"), ("BW-05", "BW-05"),
             ("BW-06", "BW-06"), ("BW-07", "BW-07"), ("BW-08", "BW-08"), ("BW-09", "BW-09"), ("BW-10", "BW-10"),
             ("BW-11", "BW-11"), ("BW-12", ""), ("BW-13", ""), ("BW-14", "")]
KEIN_LISTENPLATZ = {"BW-21"}
codes = {p["Code"] for p in personen}
for lab, ac in ZUORDNUNG:
    if ac and ac not in codes: w("ABBRUCH: Analysecode %s der Zuordnung fehlt in den Personendaten" % ac); ende(2)
ig_codes = [p["Code"] for p in personen if p["Gruppe"] == "IG"]
for c in KEIN_LISTENPLATZ:
    if c not in ig_codes: w("ABBRUCH: %s ohne Listenplatz ist kein IG-Spieler" % c); ende(2)

# ----------------------------------------------------------------------------
# Schreiben
# ----------------------------------------------------------------------------
os.makedirs(ZIEL)
def schreibe_csv(name, kopf, zeilen):
    p = os.path.join(ZIEL, name)
    with open(p, "w", encoding="utf-8", newline="") as f:
        wr = csv.writer(f, lineterminator="\n")
        wr.writerow(kopf)
        for z in zeilen: wr.writerow(z)
    return p

p1 = schreibe_csv("Versuchsdaten.csv", ["Code", "Zeitpunkt", "Test", "Seite", "Versuch", "Wert", "ungültig", "Bemerkung"],
                  [[d["Code"].strip(), d["Zeitpunkt"], d["Test"], d["Seite"], fmt(d["Versuch"]), fmt(d["Wert"]),
                    "" if d["ungültig"] in (None, "") else d["ungültig"], "" if d["Bemerkung"] in (None, "") else d["Bemerkung"]] for d in versuche])
p2 = schreibe_csv("Personendaten.csv", ["Code", "Verein", "Gruppe", "Alter_prae", "Koerperhoehe_prae", "Koerpermasse_prae", "Groesse_Mutter", "Groesse_Vater", "Familiarisierung", "Status"],
                  [[p["Code"], p["Verein"], p["Gruppe"], fmt(p["Alter_prae"]), fmt(p["Koerperhoehe_prae"]), fmt(p["Koerpermasse_prae"]),
                    fmt(p["Groesse_Mutter"]), fmt(p["Groesse_Vater"]), fmt(p["Familiarisierung"]), p["Status"]] for p in personen])
p3 = schreibe_csv("Fragebogen_A.csv", COLS, [[fmt(r[c]) if c != "STARTED" else r[c] for c in COLS] for r in meldungen])
p4 = schreibe_csv("Zuordnung_Fragebogen.csv", ["Listenlabel", "Analysecode"], [[lab, ac] for lab, ac in ZUORDNUNG])
p5 = schreibe_csv("Listenplatz_IG.csv", ["Code", "kein_Listenplatz"], [[c, "ja" if c in KEIN_LISTENPLATZ else "nein"] for c in ig_codes])
w("Geschrieben: " + ZIEL)
for p in (p1, p2, p3, p4, p5):
    with open(p, "rb") as f: b = f.read()
    w("  %-26s %7d B  BOM %s  CRLF %s" % (os.path.basename(p), len(b), "nein" if not b.startswith(b"\xef\xbb\xbf") else "JA", "nein" if b"\r\n" not in b else "JA"))

# ----------------------------------------------------------------------------
# Datenwörterbuch (ohne Kennzahlen)
# ----------------------------------------------------------------------------
DW = """# Datenwörterbuch des Datenstands %(datum)s

Eingefrorener Datenstand der Studie „U15-Plyometrie“ (DSHS Köln) für die Rechnung nach der Spezifikation der Rechenschritte vom 24.09.2026 (Abschnitt 0.4). Erzeugt am %(datum)s aus dem Workbook `Studiendaten_U15_gesamt.xlsx` (Blätter `01_Personen` und `02_Rohdaten`, Stand nach der Belegprüfung der Phase 1) und dem Originalexport des Fragebogens A (`data_test546007_2026-08-31_11-37.json`, SoSci Survey, Projekt test546007). Die Prüfsummen aller Dateien stehen in `Pruefsummen.txt`. Der Ordner wird nach dem Einfrieren nicht mehr verändert, ein späterer Stand bekommt einen neuen Ordner mit neuem Datum.

**Format aller CSV-Dateien:** UTF-8 ohne BOM, Komma als Trennzeichen, Dezimalpunkt, Zeilenende LF, erste Zeile Spaltennamen, fehlender Wert als leeres Feld, ganze Zahlen ohne Nachkommastellen, Dezimalzahlen ungerundet in der kürzesten Darstellung, die denselben Gleitkommawert ergibt. Textfelder ohne Anführungszeichen, außer sie enthalten ein Komma.

**Pseudonymisierung:** Spielercodes `HL-nn`, `BW-nn`, `VS-nn` sind die Analysecodes. Sie enthalten keinen Personenbezug. Der Datenstand enthält keine Namen, kein Geburtsdatum, kein Testdatum und keine Freitexte.

## Versuchsdaten.csv

Eine Zeile je Spieler × Zeitpunkt × Test × Seite × Versuch, alle Zeilen des Rasters, auch solche ohne Wert. Herkunft: Blatt `02_Rohdaten`, Zeilen in der Reihenfolge des Blatts. Die Spalten `Verein`, `Einheit` und `gilt` des Blatts sind nicht übernommen (Verein steht in den Personendaten, die Einheit folgt aus dem Test, `gilt` ist eine berechnete Spalte und wird nach Spezifikation 0.4 nicht verwendet).

| Spalte | Typ | Einheit | Zulässige Werte | Fehlend | Herkunft | Umsetzung |
|---|---|---|---|---|---|---|
| Code | Text | – | Analysecode nach dem Muster `XX-nn`, jeder Code steht in Personendaten.csv | nie | `02_Rohdaten`, Spalte A `Code` | unverändert |
| Zeitpunkt | Text | – | `prä`, `post` | nie | Spalte C `Zeitpunkt` | unverändert |
| Test | Text | – | `Sprint_5m`, `Sprint_10m`, `Sprint_30m`, `COD_505`, `Standweitsprung` | nie | Spalte D `Test` | unverändert |
| Seite | Text | – | `L` oder `R` genau dann, wenn Test = `COD_505`, sonst `–` (Halbgeviertstrich, wie im Workbook) | nie | Spalte E `Seite` | unverändert |
| Versuch | ganze Zahl | – | 1, 2, 3 | nie | Spalte F `Versuch` | unverändert |
| Wert | Dezimalzahl | s bei Sprint_5m, Sprint_10m, Sprint_30m und COD_505, cm bei Standweitsprung | Zahl größer 0 | leer, wenn kein Wert erfasst wurde (Versuch nicht durchgeführt, Lichtschranke ohne Aufzeichnung oder Spieler zur Post-Testung nicht angetreten) | Spalte G `Wert` | unverändert, ohne Rundung |
| ungültig | Text | – | leer oder `x` | leer = kein Eintrag | Spalte I `ungültig` | unverändert. Nach Spezifikation S02 zählt jeder Eintrag als ungültig |
| Bemerkung | Text | – | festes Vokabular nach Anlage `Spezifikation_2026-09-24_Vokabular.csv` (`System nicht aufgenommen`, `Nur 2 Versuche aus Zeitmangel`, `Linie nicht getroffen`, `Kein fester Stand`, `Fehlversuch`, `System falsch aufgenommen`) oder leer | leer = keine Bemerkung | Spalte J `Bemerkung` | unverändert nach Vereinheitlichung der Schreibweise in Phase 1 |

**Sprintlauf (Spezifikation S02, Hinweis, vom Verfasser am 24.09.2026 bestätigt):** Die drei Teilzeiten eines Sprintlaufs bei 5 m, 10 m und 30 m tragen dieselbe Versuchsnummer. Ein Lauf ist die Menge der Zeilen mit gleichem Code, Zeitpunkt und Versuch über `Sprint_5m`, `Sprint_10m` und `Sprint_30m`.

**Zeilen ohne Wert und ohne Bemerkung** kommen nur bei Spielern mit Status `nicht angetreten` zum Zeitpunkt `post` vor (Spezifikation S03 Regel 1). Eine Zeile mit Wert und Eintrag in `ungültig` ist ein gemessener, aber ungültiger Versuch (Spezifikation S02 Regel 1).

## Personendaten.csv

Eine Zeile je Spieler, nur echte Spielerzeilen des Blatts `01_Personen` (Codes nach dem Muster `XX-nn`), in der Reihenfolge des Blatts. Nicht übernommen: Geburtsdatum, Testdaten, alle berechneten Spalten (Alter dezimal, Elterngröße mittel, PAH, PAH_Prozent, Wachstum, Reifeband, Hilfsspalten), Post-Anthropometrie und die Bemerkung.

| Spalte | Typ | Einheit | Zulässige Werte | Fehlend | Herkunft | Umsetzung |
|---|---|---|---|---|---|---|
| Code | Text | – | Analysecode `XX-nn`, eindeutig | nie | `01_Personen`, Spalte A `Code` | unverändert |
| Verein | Text | – | `A`, `B`, `C` | nie | Spalte B `Verein` | `Hohenlind` → `A`, `Blau-Weiß Köln` → `B`, `Vorwärts Spoho` → `C` |
| Gruppe | Text | – | `IG`, `KG` | nie | Spalte C `Gruppe` | `Intervention` → `IG`, `Kontrolle` → `KG`. IG genau dann, wenn Verein `A` oder `B` |
| Alter_prae | Dezimalzahl | Jahre | Zahl größer 0 | nie | Spalten D `Geburtsdatum` und E `Testdatum_prä` | (Testdatum der Prätestung − Geburtsdatum) in Tagen / 365,25, im Skript aus den beiden Datumsfeldern gerechnet, nicht aus der Formelspalte, ungerundet (Spezifikation S05 Regel 1, K1) |
| Koerperhoehe_prae | Zahl | cm | Zahl größer 0 | leer, wenn bei der Prätestung nicht erfasst | Spalte I `Größe_prä_cm` | unverändert |
| Koerpermasse_prae | Zahl | kg | Zahl größer 0 | leer, wenn bei der Prätestung nicht erfasst | Spalte J `Gewicht_prä_kg` | unverändert |
| Groesse_Mutter | Zahl | cm | Zahl größer 0 | leer, wenn nicht angegeben | Spalte K `Größe_Mutter_cm` | unverändert |
| Groesse_Vater | Zahl | cm | Zahl größer 0 | leer, wenn nicht angegeben | Spalte L `Größe_Vater_cm` | unverändert |
| Familiarisierung | ganze Zahl | Termine | 1, 2 | leer, wenn nicht dokumentiert | Spalte P `Familiarisierung` | unverändert (Zahl der besuchten Familiarisierungstermine laut Anwesenheitsliste) |
| Status | Text | – | `ausgewertet`, `nicht angetreten` | nie | Spalte Q `Status` | feste Umsetzung: `ausgewertet` → `ausgewertet`, Text beginnend mit `Post-Testung nicht angetreten` → `nicht angetreten`, jeder andere Text bricht den Export ab |

## Fragebogen_A.csv

Eine Zeile je Meldung des Fragebogens A (Einheitsprotokoll der Interventionsgruppe), alle Meldungen des Originalexports, sortiert nach CASE. Herkunft: `data_test546007_2026-08-31_11-37.json`, Objekt `data`, ein Datensatz je Meldung. Werte unverändert, keine weiteren Spalten, keine Freitexte, keine Bearbeitungsdauer. Die Fallkorrekturen der CASE-Nummern (Anlage `_Fragebogen.csv`, Zeilen `CASE_KORREKTUR`) sind nicht angewandt, die Spezifikation wendet sie in S06 Regel 2 an. Die Bedeutung der Codes steht im Codebuch und in der Anlage `Spezifikation_2026-09-24_Fragebogen.csv`.

| Spalte | Typ | Zulässige Werte | Fehlend | Herkunft | Umsetzung |
|---|---|---|---|---|---|
| CASE | ganze Zahl | eindeutige Fallnummer des Exports | nie | Feld `CASE` | unverändert |
| STARTED | Text | Zeitstempel `JJJJ-MM-TT hh:mm:ss`, Ortszeit des Exports | nie | Feld `STARTED` | unverändert, das Exportformat entspricht bereits der Vorgabe |
| H010 | ganze Zahl | 1 bis 22 (Listenplatz, Codebuch H010) | nie | Feld `H010` | unverändert |
| H002 | ganze Zahl | 1 (Datum „Heute“, einzige Option) | nie | Feld `H002` | unverändert |
| H003 | ganze Zahl | 1 bis 12 (Einheit, Codebuch H003) | nie | Feld `H003` | unverändert |
| H004 | ganze Zahl | 1 ganz, 2 teilweise, 3 gar nicht | nie | Feld `H004` | unverändert |
| H005 | ganze Zahl | 1 bis 11 (CR-10-Stufe 0 bis 10 = Code − 1) | nie | Feld `H005` | unverändert |
| H006 | ganze Zahl | 1 bis 5 (Untergrund) | nie | Feld `H006` | unverändert |
| H007 | ganze Zahl | 1 ja, 2 nein (Schmerzen oder Probleme) | nie | Feld `H007` | unverändert |
| H008 | ganze Zahl | 1 ja, 2 nein (Video angehalten oder Übungen wiederholt) | nie | Feld `H008` | unverändert |
| H009 | ganze Zahl | 1 ja, 2 nein (anderes Sporttraining seit der letzten Einheit) | nie | Feld `H009` | unverändert |

## Zuordnung_Fragebogen.csv

Zuordnung der Listenlabels des Fragebogens (Codebuch H010, altes Codeschema) zu den Analysecodes, vom Verfasser am 24.09.2026 bestätigt. 22 Zeilen, eine je Listenplatz.

| Spalte | Typ | Zulässige Werte | Fehlend | Herkunft | Umsetzung |
|---|---|---|---|---|---|
| Listenlabel | Text | `HL-01` bis `HL-08`, `BW-01` bis `BW-14` | nie | Codebuch H010 | Listenplatz 1 bis 8 → `HL-01` bis `HL-08`, 9 bis 22 → `BW-01` bis `BW-14` |
| Analysecode | Text | Code aus Personendaten.csv oder leer | leer = kein realer Spieler hinter dem Listenplatz | Codezuordnung Phase 1 (Schritt 1.2 b) | Listenlabel = Workbook-Code, mit einer Ausnahme: Listenlabel `HL-04` → Analysecode `HL-05` (der Spieler mit dem Workbook-Code `HL-05` trug in der Auswahlliste den Platz `HL-04`) |

## Listenplatz_IG.csv

Eine Zeile je Spieler der Interventionsgruppe (Gruppe `IG` in Personendaten.csv), in der Reihenfolge der Personendaten.

| Spalte | Typ | Zulässige Werte | Fehlend | Herkunft | Umsetzung |
|---|---|---|---|---|---|
| Code | Text | Analysecode der IG | nie | Personendaten.csv | – |
| kein_Listenplatz | Text | `ja`, `nein` | nie | Codezuordnung Phase 1 (Schritt 1.2 b) | `ja`, wenn der Spieler in der Auswahlliste des Fragebogens nicht zur Verfügung stand (Instrumentenfehler, Spezifikation S07 Regel 5 und K9), sonst `nein` |

## Pruefsummen.txt

SHA-256 aller Dateien dieses Ordners außer der Prüfsummendatei selbst, eine Zeile je Datei (`Prüfsumme  Dateiname`).
""" % {"datum": DATUM}
p6 = os.path.join(ZIEL, "Datenwoerterbuch.md")
with open(p6, "w", encoding="utf-8", newline="\n") as f: f.write(DW)
w("  Datenwoerterbuch.md geschrieben (%d B)" % os.path.getsize(p6))

# ----------------------------------------------------------------------------
# Probe 1: Rücklesen
# ----------------------------------------------------------------------------
w()
w("Probe 1, Rückleseprobe:")
def lies(name):
    with open(os.path.join(ZIEL, name), encoding="utf-8", newline="") as f:
        b = f.read(1)
        if b == "﻿": raise RuntimeError("BOM")
        f.seek(0)
        return list(csv.DictReader(f))
fehler = 0
V = lies("Versuchsdaten.csv")
if len(V) != len(versuche): w("  Versuchsdaten: Zeilenzahl weicht ab"); fehler += 1
for d, z in zip(versuche, V):
    soll = {"Code": d["Code"].strip(), "Zeitpunkt": d["Zeitpunkt"], "Test": d["Test"], "Seite": d["Seite"], "Versuch": fmt(d["Versuch"]),
            "Wert": fmt(d["Wert"]), "ungültig": "" if d["ungültig"] in (None, "") else d["ungültig"], "Bemerkung": "" if d["Bemerkung"] in (None, "") else d["Bemerkung"]}
    if z != soll: fehler += 1; w("  Versuchsdaten Zeile %d weicht ab" % d["_zeile"])
    # Zahlwert exakt gleich
    if d["Wert"] not in (None, "") and float(z["Wert"]) != float(d["Wert"]): fehler += 1; w("  Wert nicht identisch Zeile %d" % d["_zeile"])
P = lies("Personendaten.csv")
if len(P) != len(personen): w("  Personendaten: Zeilenzahl weicht ab"); fehler += 1
for p, z in zip(personen, P):
    soll = {k: (fmt(p[k]) if k not in ("Code", "Verein", "Gruppe", "Status") else p[k]) for k in ["Code", "Verein", "Gruppe", "Alter_prae", "Koerperhoehe_prae", "Koerpermasse_prae", "Groesse_Mutter", "Groesse_Vater", "Familiarisierung", "Status"]}
    if z != soll: fehler += 1; w("  Personendaten %s weicht ab" % p["Code"])
    if float(z["Alter_prae"]) != p["Alter_prae"]: fehler += 1; w("  Alter nicht identisch %s" % p["Code"])
F = lies("Fragebogen_A.csv")
if len(F) != len(meldungen): w("  Fragebogen: Zeilenzahl weicht ab"); fehler += 1
for r, z in zip(meldungen, F):
    soll = {c: (fmt(r[c]) if c != "STARTED" else r[c]) for c in COLS}
    if z != soll: fehler += 1; w("  Fragebogen CASE %s weicht ab" % r["CASE"])
Z = lies("Zuordnung_Fragebogen.csv")
if [(z["Listenlabel"], z["Analysecode"]) for z in Z] != ZUORDNUNG: fehler += 1; w("  Zuordnung weicht ab")
L = lies("Listenplatz_IG.csv")
if [(z["Code"], z["kein_Listenplatz"]) for z in L] != [(c, "ja" if c in KEIN_LISTENPLATZ else "nein") for c in ig_codes]: fehler += 1; w("  Listenplatz weicht ab")
w("  Abweichungen: %d" % fehler)
if fehler: w("ABBRUCH: Rückleseprobe nicht bestanden"); ende(2)
w("  Rückleseprobe bestanden: ja")

# ----------------------------------------------------------------------------
# Probe 2: Abbruchregeln der Spezifikation am Datenstand
# ----------------------------------------------------------------------------
w()
w("Probe 2, Abbruchprobe (0.5 Nr. 10, S01 Regeln 2 bis 4, S03 mit K5, S06 Regeln 3 und 7):")
with open(P_VOK, encoding="utf-8", newline="") as f: VOK = {r["bemerkung_exakt"]: r["kategorie"] for r in csv.DictReader(f)}
with open(P_FB, encoding="utf-8", newline="") as f: FBR = list(csv.DictReader(f))
FB_CODES = collections.defaultdict(set); KORR = {}
for r in FBR:
    if r["variable"] == "CASE_KORREKTUR": KORR[r["code"]] = r["bedeutung"].replace("Listenlabel ", "")
    else: FB_CODES[r["variable"]].add(r["code"])
verst = []
# S01 Regel 3, Personendaten
pcodes = [z["Code"] for z in P]
if len(set(pcodes)) != len(pcodes): verst.append("S01.3 Code nicht eindeutig")
gruppe_von = {}
status_von = {}
for z in P:
    if z["Verein"] not in ("A", "B", "C"): verst.append("S01.3 Verein " + z["Verein"])
    if z["Gruppe"] != ("IG" if z["Verein"] in ("A", "B") else "KG"): verst.append("S01.3 Gruppe passt nicht zum Verein bei " + z["Code"])
    if z["Familiarisierung"] not in ("", "1", "2"): verst.append("S01.3 Familiarisierung bei " + z["Code"])
    if z["Status"] not in ("ausgewertet", "nicht angetreten"): verst.append("S01.3 Status bei " + z["Code"])
    for k in ("Alter_prae", "Koerperhoehe_prae", "Koerpermasse_prae", "Groesse_Mutter", "Groesse_Vater"):
        if z[k] != "":
            try:
                if float(z[k]) <= 0: verst.append("Wert nicht größer 0: %s %s" % (z["Code"], k))
            except ValueError: verst.append("0.5.10 Format %s %s" % (z["Code"], k))
    gruppe_von[z["Code"]] = z["Gruppe"]; status_von[z["Code"]] = z["Status"]
# S01 Regel 2, Versuchsdaten
keys = collections.Counter()
for z in V:
    if z["Zeitpunkt"] not in ("prä", "post"): verst.append("S01.2 Zeitpunkt " + z["Zeitpunkt"])
    if z["Test"] not in ("Sprint_5m", "Sprint_10m", "Sprint_30m", "COD_505", "Standweitsprung"): verst.append("S01.2 Test " + z["Test"])
    if (z["Test"] == "COD_505") != (z["Seite"] in ("L", "R")): verst.append("S01.2 Seite bei %s %s" % (z["Code"], z["Test"]))
    if z["Versuch"] not in ("1", "2", "3"): verst.append("S01.2 Versuch " + z["Versuch"])
    if z["Wert"] != "":
        try:
            if float(z["Wert"]) <= 0: verst.append("S01.2 Wert nicht größer 0")
        except ValueError: verst.append("0.5.10 Wert kein Zahlformat: " + z["Wert"])
    if z["Code"] not in gruppe_von: verst.append("S01.2 Code %s nicht in den Personendaten" % z["Code"])
    keys[(z["Code"], z["Zeitpunkt"], z["Test"], z["Seite"], z["Versuch"])] += 1
if any(n > 1 for n in keys.values()): verst.append("S01.2 Schlüssel doppelt")
# S03 mit K5
for z in V:
    gueltig_roh = (z["Wert"] != "") and (z["ungültig"] == "")
    if gueltig_roh: continue
    if z["Zeitpunkt"] == "post" and status_von.get(z["Code"]) == "nicht angetreten":
        if z["Wert"] != "": verst.append("S03.1 Wert trotz nicht angetreten bei " + z["Code"])
        continue
    if z["Bemerkung"] == "": verst.append("K5 ungültige Zeile ohne Bemerkung: %s %s %s %s V%s" % (z["Code"], z["Zeitpunkt"], z["Test"], z["Seite"], z["Versuch"]))
    elif z["Bemerkung"] not in VOK: verst.append("K5 Bemerkung nicht im Vokabular: %r" % z["Bemerkung"])
# S01 Regel 4 und S06 Regeln 3 und 7, Fragebogen
cases = collections.Counter(z["CASE"] for z in F)
if any(n > 1 for n in cases.values()): verst.append("S01.4 CASE doppelt")
W1 = datetime.datetime(2026, 7, 20); W6 = datetime.datetime(2026, 8, 30, 23, 59, 59)
zuo = dict(ZUORDNUNG)
lab_von_code = {r["code"]: r["bedeutung"] for r in FBR if r["variable"] == "H010"}
for z in F:
    for c in ("H010", "H002", "H003", "H004", "H005", "H006", "H007", "H008", "H009"):
        if z[c] not in FB_CODES[c]: verst.append("S01.4 %s = %r bei CASE %s" % (c, z[c], z["CASE"]))
    try:
        ts = datetime.datetime.strptime(z["STARTED"], "%Y-%m-%d %H:%M:%S")
        if not (W1 <= ts <= W6): verst.append("S06.7 Meldung außerhalb W1 bis W6, CASE " + z["CASE"])
    except ValueError: verst.append("S01.4 STARTED Format CASE " + z["CASE"])
    lab = KORR.get(z["CASE"], lab_von_code.get(z["H010"]))
    ac = zuo.get(lab, "")
    if ac == "": verst.append("S06.3 Meldung an Listenlabel ohne Analysecode: %s (CASE %s)" % (lab, z["CASE"]))
    elif gruppe_von.get(ac) == "KG": verst.append("S06.3 Analysecode der KG: " + ac)
for z in L:
    if gruppe_von.get(z["Code"]) != "IG": verst.append("Listenplatz_IG: %s ist kein IG-Spieler" % z["Code"])
if sorted(z["Code"] for z in L) != sorted(c for c, g in gruppe_von.items() if g == "IG"): verst.append("Listenplatz_IG: Menge weicht von der IG ab")
w("  Verstöße: %d" % len(verst))
for v in verst: w("   " + v)
if verst: w("ABBRUCH: Abbruchprobe nicht bestanden. Ordner bleibt zur Ansicht, Prüfsummen nicht geschrieben."); ende(2)
w("  Abbruchprobe bestanden: ja")

# ----------------------------------------------------------------------------
# Prüfsummen
# ----------------------------------------------------------------------------
dateien = sorted(f for f in os.listdir(ZIEL) if f != "Pruefsummen.txt")
with open(os.path.join(ZIEL, "Pruefsummen.txt"), "w", encoding="utf-8", newline="\n") as f:
    for name in dateien:
        f.write("%s  %s\n" % (sha256(os.path.join(ZIEL, name)), name))
w()
w("Pruefsummen.txt:")
with open(os.path.join(ZIEL, "Pruefsummen.txt"), encoding="utf-8") as f:
    for zeile in f: w("  " + zeile.rstrip())
w()
w("Datenstand eingefroren: " + ZIEL + ". Danach wird nichts mehr verändert.")
ende(0)
