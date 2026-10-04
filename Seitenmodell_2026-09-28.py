# -*- coding: utf-8 -*-
"""
Seitenmodell_2026-09-28.py — Seitenrechnung der Bachelorarbeit als Modellrechnung, keine Messung
Bachelorarbeit U15-Plyometrie · DSHS Köln · Task Steuerdokumente 28.09. (Übergabe § 3.3)

Zweck: Quelle der Seitenrechnung in Projektanweisungen Fassung 17 § 1.1 und der Parameter für den Block
„Seitenschätzung“ des Messskripts `Manuskriptstand_2026-09-25.py` (Fassung 3). Grenze des Verfassers vom 28.09.:
höchstens 33 Seiten einschließlich Literaturverzeichnis, Zählweise nach Klick K1 (Einleitung bis Ende des
Literaturverzeichnisses, ohne Vorspann und Anhang).

Vier Teile (Übergabe § 3.3):
(a) Dichte und Umbruchverlust am heutigen Master: Master per LibreOffice als PDF gerendert, je Kapitel mit Text die
    Seiten vom Kapitelanfang bis zur letzten Textzeile gegen die Wörter des Absatztexts (Zählregel wie im Messskript).
    Word und LibreOffice brechen leicht verschieden um, Kontrolle bleibt die Word-Messung am fertigen Dokument.
(b) Literaturverzeichnis an einer Probe: 20 echte Einträge aus T1 (APA 7) im Layout des Masters (Formatvorlage
    Standard, Arial 11, Zeilenabstand 1,5 nach SMK-Leitfaden 4.4.2, hängender Einzug 1,25 cm nach SMK 4.6),
    Einträge je Seite gemessen.
(c) Zahl der Einträge geschätzt: Belege im Master in Kapitel 4, Belege im Wortlaut des Textvorschlags Einleitung (§ 1,
    seit Fassung 2, vorher Quellen der Übergabe Einleitung § 3), vorgemerkte Vorstudien für Kapitel 6 (F17 § 6.5), ohne
    Doppelzählung. Obere Variante mit den für Kapitel 6 vorgemerkten weiteren Quellen (Maßnahmenliste G33 d und f, G34 c,
    F17 § 12, H9, K18).
(d) Vorspann nach K1: zählt nicht, die gerenderte Seitenzahl steht nur zur Information da.
Jede Zeile der Seitenrechnung ist als Modellrechnung gekennzeichnet (F17 § 1.1: Modellrechnung und Messung trennen).

Aufruf: python Seitenmodell_2026-09-28.py <Master.docx> <T1_steckbriefe.csv> <Textvorschlag_Einleitung.md>
        <Ausgabe.txt> <Parameter.csv>
Ausgabe: Laufprotokoll mit Seitenrechnung (.txt), Parameter für das Messskript (.csv, Trennzeichen chr(59)).
Braucht: python-docx, pdfplumber, LibreOffice (soffice). Ohne Semikolon im Skript (chr(59)).
Fassung: 2026-09-28, Fassung 2 (Einträge der Einleitung aus dem Wortlaut des Textvorschlags, Zweitprüfung 28.09. Befund 28).
Fassung 1: Dichte je Kapiteltyp, Einleitung nach der Quellenliste der Übergabe Einleitung § 3.
"""
import sys
import os
import re
import csv
import shutil
import tempfile
import subprocess
import datetime
import hashlib
from docx import Document
from docx.shared import Cm, Pt
import pdfplumber

MASTER, T1, UEB_EINL, OUT_TXT, OUT_CSV = sys.argv[1:6]
SEMI = chr(59)
LOG = []


def log(*t):
    s = ' '.join(str(x) for x in t)
    LOG.append(s)
    print(s)


def de(x, nk=1):
    """Deutsche Zahlendarstellung mit Tausenderpunkt und Dezimalkomma."""
    s = ('%.' + str(nk) + 'f') % x
    ganz, _, dez = s.partition('.')
    neg = ganz.startswith('-')
    ganz = ganz.lstrip('-')
    gruppen = []
    while len(ganz) > 3:
        gruppen.insert(0, ganz[-3:])
        ganz = ganz[:-3]
    gruppen.insert(0, ganz)
    r = '.'.join(gruppen)
    if dez:
        r += ',' + dez
    return ('−' if neg else '') + r


def md5(p):
    return hashlib.md5(open(p, 'rb').read()).hexdigest()


# ---------------------------------------------------------------- Zählregel wie im Messskript (Fassung 2 und 3)
def absatztext(doc):
    """Liefert je Überschrift die Absätze in Formatvorlage Standard ohne Beschriftungen, Anmerkungen, Platzhalter."""
    abschnitte = []
    akt = None
    for p in doc.paragraphs:
        st = p.style.name
        txt = p.text.strip()
        if st.startswith('Heading'):
            m = re.match(r'^(\d+(?:\.\d+)*)\s', txt)
            akt = {'nr': m.group(1) if m else txt, 'titel': txt, 'ebene': int(st[-1]) if st[-1].isdigit() else 0,
                   'absaetze': []}
            abschnitte.append(akt)
            continue
        if akt is None or not txt:
            continue
        if txt.startswith('⟨') or st == 'Caption' or re.match(r'^(Tab|Abb)\.\s[A-H]?\d+\.(\s|$)', txt):
            continue
        if txt.startswith('Anmerkung.') or st != 'Normal':
            continue
        akt['absaetze'].append(txt)
    return abschnitte


def woerter(abschnitt):
    return sum(len(x.split()) for x in abschnitt['absaetze'])


# ---------------------------------------------------------------- Belege (Autor-Jahr) wie im Messskript Fassung 3
NAME = r"[A-ZÄÖÜ][A-Za-zÄÖÜäöüßéèáíóúñ'’\-]+"
AUTOR = r"(?:R Core Team|" + NAME + r"(?:,\s" + NAME + r",\set al\.|\set al\.|\s(?:&|und)\s" + NAME + r")?)"


def schluessel(autor, jahr):
    a = re.sub(r'\s+und\s+', ' & ', autor.strip().rstrip(','))
    a = a.replace('’', "'")
    return a + ' ' + jahr


def belege(text):
    """Verschiedene Autor-Jahr-Belege eines Textes. Sekundärquellen vor „zitiert nach“ werden getrennt geführt."""
    gefunden, sekundaer = set(), set()
    for m in re.finditer('(' + AUTOR + r')\s\((\d{4}[a-z]?)((?:,\s\d{4}[a-z]?)*)(?:,\s[^)]*)?\)', text):
        for j in [m.group(2)] + re.findall(r'\d{4}[a-z]?', m.group(3) or ''):
            gefunden.add(schluessel(m.group(1), j))
    for g in re.finditer(r'\(([^()]*?\d{4}[^()]*?)\)', text):
        for teil in g.group(1).split(SEMI):
            stuecke = re.split(r',\szitiert nach\s', teil.strip())
            for i, st in enumerate(stuecke):
                st = re.sub(r'^(?:vgl\.|nach)\s', '', st.strip())
                m = re.match('(' + AUTOR + r'),\s(\d{4}[a-z]?)((?:,\s\d{4}[a-z]?)*)', st)
                if not m:
                    continue
                for j in [m.group(2)] + re.findall(r'\d{4}[a-z]?', m.group(3) or ''):
                    k = schluessel(m.group(1), j)
                    if len(stuecke) > 1 and i == 0:
                        sekundaer.add(k)
                    else:
                        gefunden.add(k)
    return gefunden - sekundaer, sekundaer


# ---------------------------------------------------------------- Rendern
def rendern(docx_pfad, arbeit):
    ziel = os.path.join(arbeit, os.path.splitext(os.path.basename(docx_pfad))[0] + '.pdf')
    subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', docx_pfad, '--outdir', arbeit],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300, check=False)
    if not os.path.exists(ziel):
        raise SystemExit('Abbruch: LibreOffice hat kein PDF erzeugt (' + docx_pfad + ')')
    return ziel


def zeilen_je_seite(pdf_pfad, oben, unten):
    """Textzeilen je Seite innerhalb des Satzspiegels (Kopfzeile mit Seitenzahl ausgeschlossen)."""
    seiten = []
    with pdfplumber.open(pdf_pfad) as pdf:
        for page in pdf.pages:
            ws = page.extract_words(keep_blank_chars=False, use_text_flow=False)
            zeilen = {}
            for w in ws:
                if w['top'] < oben - 2 or w['bottom'] > unten + 2:
                    continue
                key = round(w['top'] / 2.0)
                zeilen.setdefault(key, []).append(w)
            zl = []
            for key in sorted(zeilen):
                ww = sorted(zeilen[key], key=lambda x: x['x0'])
                zl.append({'text': ' '.join(x['text'] for x in ww), 'top': min(x['top'] for x in ww),
                           'bottom': max(x['bottom'] for x in ww)})
            seiten.append(zl)
    return seiten


def finde_ueberschrift(seiten, titel, ab_seite):
    """Erste Zeile ab Seite ab_seite, die mit dem Überschriftentext beginnt und keine Punktleiste trägt."""
    kopf = re.sub(r'\s+', ' ', titel)[:28]
    for si in range(ab_seite, len(seiten)):
        for zi, z in enumerate(seiten[si]):
            t = re.sub(r'\s+', ' ', z['text'])
            if t.startswith(kopf) and '....' not in t and '…' not in t:
                return si, zi
    return None


def main():
    arbeit = tempfile.mkdtemp(prefix='seitenmodell_')
    log('Seitenmodell_2026-09-28.py, Lauf', datetime.datetime.now().strftime('%d.%m.%Y %H:%M'), '(Uhr des Rechners)')
    log('Master:', os.path.basename(MASTER), os.path.getsize(MASTER), 'Byte, MD5', md5(MASTER))
    log('')

    doc = Document(MASTER)
    sec = doc.sections[-1]
    oben = sec.top_margin.pt
    unten = sec.page_height.pt - sec.bottom_margin.pt
    hoehe = unten - oben
    log('Satzspiegel aus dem Master: Seitenhöhe %s pt, Rand oben %s pt, unten %s pt, Textbreite %s cm'
        % (de(sec.page_height.pt), de(oben), de(sec.bottom_margin.pt),
           de((sec.page_width - sec.left_margin - sec.right_margin) / 360000.0)))

    # ---------------------------------------------------------- (a) Dichte am Master
    kopie = os.path.join(arbeit, 'master.docx')
    shutil.copy(MASTER, kopie)
    pdf = rendern(kopie, arbeit)
    seiten = zeilen_je_seite(pdf, oben, unten)
    log('Gerendert mit LibreOffice:', len(seiten), 'Seiten')
    abschnitte = absatztext(doc)

    pos = []
    ab = 0
    start_gefunden = False
    for a in abschnitte:
        if not start_gefunden and not a['titel'].startswith('1 Einleitung'):
            continue
        start_gefunden = True
        f = finde_ueberschrift(seiten, a['titel'], ab)
        if f is None:
            raise SystemExit('Abbruch: Überschrift nicht im PDF gefunden: ' + a['titel'])
        pos.append((a, f[0], f[1]))
        ab = f[0]
    kap1 = [p for p in pos if p[0]['ebene'] == 1]
    einl_seite = kap1[0][1] + 1
    lit = [p for p in kap1 if p[0]['titel'].startswith('Literaturverzeichnis')][0]
    log('Einleitung beginnt auf Seite %d, Literaturverzeichnis auf Seite %d (LibreOffice-Zählung ab Titelblatt)'
        % (einl_seite, lit[1] + 1))
    vorspann = kap1[0][1]
    log('Vorspann (Titelblatt bis Verzeichnisse): %d Seiten, zählt nach K1 nicht zu den 33 Seiten' % vorspann)
    log('')

    def letzte_textzeile(i_von, i_bis):
        """Letzte Zeile mit Absatztext zwischen den Überschriften pos[i_von] und pos[i_bis] (ausschließlich)."""
        s_ende, z_ende = pos[i_bis][1], pos[i_bis][2]
        for i in range(i_bis - 1, i_von - 1, -1):
            if woerter(pos[i][0]) > 0:
                j = i + 1
                s2, z2 = pos[j][1], pos[j][2]
                # letzte Zeile vor der nächsten Überschrift
                if z2 > 0:
                    return s2, seiten[s2][z2 - 1]
                s = s2 - 1
                while s >= 0 and not seiten[s]:
                    s -= 1
                return s, seiten[s][-1]
        return None

    bloecke = []
    for k, (a, s, z) in enumerate(pos):
        if a['ebene'] != 1 or not re.match(r'^\d', a['titel']):
            continue
        nxt = [i for i in range(k + 1, len(pos)) if pos[i][0]['ebene'] == 1]
        if not nxt:
            continue
        i_bis = nxt[0]
        w = sum(woerter(pos[i][0]) for i in range(k, i_bis))
        if w == 0:
            continue
        ende = letzte_textzeile(k, i_bis)
        s_e, zeile = ende
        anteil = (zeile['bottom'] - oben) / hoehe
        seiten_block = (s_e - s) + anteil
        rest = 1.0 - anteil
        bloecke.append({'titel': a['titel'], 'woerter': w, 'von': s + 1, 'bis': s_e + 1, 'seiten': seiten_block,
                        'dichte': w / seiten_block, 'rest': rest})
    log('(a) Dichte am Master (Messung am LibreOffice-Render, Absatztext nach der Zählregel des Messskripts):')
    for b in bloecke:
        log('    %-52s %5d Wörter auf Seite %d bis %d = %s Seiten bis zur letzten Textzeile -> %s Wörter je Seite, '
            'Rest der letzten Seite %s' % (b['titel'][:52], b['woerter'], b['von'], b['bis'], de(b['seiten'], 2),
                                            de(b['dichte'], 0), de(b['rest'], 2)))
    w_ges = sum(b['woerter'] for b in bloecke)
    s_ges = sum(b['seiten'] for b in bloecke)
    dichte = w_ges / s_ges
    b_einl = [b for b in bloecke if b['titel'].startswith('2 ')]
    b_rest = [b for b in bloecke if b['titel'].startswith('4 ')]
    if len(b_einl) != 1 or len(b_rest) != 1:
        raise SystemExit('Abbruch: Kapitel 2 (Altbestand) oder Kapitel 4 im Render nicht eindeutig gefunden')
    d_einl = b_einl[0]['dichte']
    d_rest = b_rest[0]['dichte']
    log('    zusammen %d Wörter auf %s Seiten -> %s Wörter je Seite (nur zum Vergleich)' % (w_ges, de(s_ges, 2), de(dichte, 0)))
    log('    Parameter des Modells nach Kapiteltyp: Einleitung wie der Altbestand Kapitel 2 (Fließtext mit Belegen,'
        ' wenige Überschriften) %s Wörter je Seite · Methodik bis Fazit wie Kapitel 4 (mehr Unterüberschriften) %s Wörter'
        ' je Seite' % (de(d_einl, 0), de(d_rest, 0)))
    log('    Vergleich: F16 § 1.1 führte 327 Wörter je Seite einschließlich Umbrüchen und drei Objekten (Messung 13.09.).'
        ' Die neue Dichte ist ohne Objekte und ohne Kapitelende gemessen, beides steht unten als eigene Zeile.')
    log('    Rest der letzten Seite je Kapitelende (gemessen): ' + ', '.join(de(b['rest'], 2) for b in bloecke)
        + '. Im Modell gilt der Erwartungswert 0,5 Seiten je Kapitelende (Annahme, gleichverteilter Rest).')
    log('')

    # Plausibilität der Objektgröße: Abbildungen aus der Objektvorlage auf Textbreite
    textbreite_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000.0
    satz_cm = hoehe / 72.0 * 2.54
    log('    Plausibilität 0,4 Seiten je Objekt (F17 § 5.3): Abb. 1 (Seitenverhältnis 0,75) auf %s cm Textbreite '
        '%s cm hoch = %s des Satzspiegels, Abb. 2 (0,406) %s cm = %s, jeweils ohne Beschriftung. Tabellen: keine Messung '
        '(Objektvorlage im Querformat), Annahme bleibt.' % (de(textbreite_cm), de(textbreite_cm * 0.75),
                                                            de(textbreite_cm * 0.75 / satz_cm, 2), de(textbreite_cm * 0.406),
                                                            de(textbreite_cm * 0.406 / satz_cm, 2)))
    log('')

    # ---------------------------------------------------------- (c) Zahl der Einträge
    k4 = set()
    sek = set()
    for a in abschnitte:
        if re.match(r'^4(\.|$)', a['nr']):
            g, s2 = belege(' '.join(a['absaetze']))
            k4 |= g
            sek |= s2
    text_einl = open(UEB_EINL, encoding='utf-8').read()
    if '## 1 Wortlaut' in text_einl:
        # Fassung 2: Textvorschlag Einleitung, Belege aus dem Wortlaut (§ 1, neun Absätze)
        wortlaut = text_einl.split('## 1 Wortlaut')[1].split('\n## 2 ')[0]
        einl, _ = belege(wortlaut)
        QUELLE_EINL = 'Belege im Wortlaut des Textvorschlags Einleitung (§ 1)'
    else:
        m = re.search(r'## 3 Bauplan.*?\n(\|.*?)\n\n', text_einl, re.S)
        if not m:
            raise SystemExit('Abbruch: Tabelle „3 Bauplan“ in der Übergabe Einleitung nicht gefunden')
        einl = set()
        for zeile in m.group(1).split('\n'):
            zellen = [c.strip() for c in zeile.strip('|').split('|')]
            if len(zellen) < 4 or not re.match(r'^\d', zellen[0]):
                continue
            g, _ = belege(zellen[3])
            einl |= g
        QUELLE_EINL = 'Quellen der Übergabe Einleitung § 3'
    vorab = {schluessel('Lloyd et al.', '2016'), schluessel('Ramirez-Campillo et al.', '2020'),
             schluessel('Moran et al.', '2017')}                       # F17 § 6.5, Stärkste Vorab-Erwartungen
    obere = {schluessel('Hilska et al.', '2021'),                       # G34 (c): Adhärenzreferenz 6.3
             schluessel('Morgan et al.', '2022'), schluessel('Havanecz et al.', '2026'),
             schluessel('Marín-Jiménez et al.', '2024'),                # G33 (f): Task 12
             schluessel('Taylor et al.', '2018'),                       # G33 (d): 6.2
             schluessel('Ferguson et al.', '2024'), schluessel('Sáez-Sáez de Villarreal et al.', '2009'),
             schluessel('Behm et al.', '2017'), schluessel('Moran et al.', '2024'),   # F17 § 12 G4, G8
             schluessel('Hedges', '1981'), schluessel('Al Haddad et al.', '2015')}    # H9, K18 nach Beschaffung
    basis = k4 | einl | vorab
    obere_menge = basis | obere
    log('(c) Zahl der Einträge des Literaturverzeichnisses (Schätzung, keine Messung):')
    log('    Belege in Kapitel 4 des Masters: %d verschiedene Autor-Jahr-Belege, Sekundärquellen ohne Eintrag: %s'
        % (len(k4), ', '.join(sorted(sek)) or 'keine'))
    log('      ' + ' · '.join(sorted(k4)))
    log('    %s: %d' % (QUELLE_EINL, len(einl)))
    log('      ' + ' · '.join(sorted(einl)))
    log('    Vorab-Erwartungen für Kapitel 6 (F17 § 6.5): %d, davon neu gegenüber den beiden Mengen: %d'
        % (len(vorab), len(vorab - k4 - einl)))
    log('    Basis ohne Doppelzählung: %d Einträge' % len(basis))
    log('    Obere Variante mit den für Kapitel 6 vorgemerkten weiteren Quellen (G33 d, f, G34 c, F17 § 12, H9, K18): '
        '%d Einträge (+%d)' % (len(obere_menge), len(obere_menge) - len(basis)))
    log('')

    # ---------------------------------------------------------- (b) Probe Literaturverzeichnis
    t1 = list(csv.DictReader(open(T1, encoding='utf-8-sig'), delimiter=SEMI))
    kandidaten = []
    for r in t1:
        vz = r['vollzitat'].strip()
        if '⟨' in vz or 'https://doi.org/' not in vz:
            continue
        kandidaten.append(vz)
    kandidaten.sort(key=lambda x: x.lower())
    probe = kandidaten[:20]
    if len(probe) < 20:
        raise SystemExit('Abbruch: weniger als 20 vollständige Einträge in T1')
    pdoc = Document(MASTER)
    body = pdoc.element.body
    for ch in list(body):
        if not ch.tag.endswith('}sectPr'):
            body.remove(ch)
    pdoc.add_paragraph('Literaturverzeichnis', style='Heading 1')
    for vz in probe:
        p = pdoc.add_paragraph(vz, style='Normal')
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.first_line_indent = Cm(-1.25)
    probe_docx = os.path.join(arbeit, 'probe_literatur.docx')
    pdoc.save(probe_docx)
    probe_pdf = rendern(probe_docx, arbeit)
    ps = zeilen_je_seite(probe_pdf, oben, unten)
    ps = [s for s in ps if s]
    anteil = (ps[-1][-1]['bottom'] - oben) / hoehe
    probe_seiten = (len(ps) - 1) + anteil
    je_seite = 20 / probe_seiten
    zeichen = sum(len(x) for x in probe) / 20.0
    log('(b) Probe Literaturverzeichnis (Messung am LibreOffice-Render):')
    log('    20 vollständige Einträge aus T1 (APA 7 mit DOI, alphabetisch die ersten 20), mittlere Länge %s Zeichen,'
        ' Formatvorlage Standard des Masters (Arial 11, Zeilenabstand 1,5, Blocksatz), hängender Einzug 1,25 cm,'
        ' Überschrift „Literaturverzeichnis“ in Überschrift 1' % de(zeichen, 0))
    log('    belegt %s Seiten -> %s Einträge je Seite einschließlich Überschrift' % (de(probe_seiten, 2), de(je_seite, 1)))
    log('    Hinweis: Der SMK-Leitfaden regelt für das Literaturverzeichnis nur den hängenden Absatz (1 bis 1,5 cm, Abschn.'
        ' 4.6). Der Zeilenabstand 1,5 gilt für den Fließtext (Abschn. 4.4.2). Ein engerer Satz des Verzeichnisses ist dort'
        ' nicht vorgesehen und wird hier nicht angenommen.')
    log('')

    # ---------------------------------------------------------- (d) Seitenrechnung
    WORT = 6350
    W_EINL = 1500
    OBJ, JE_OBJ = 5, 0.4
    KAP_ENDEN, JE_ENDE = 5, 0.5
    text = W_EINL / d_einl + (WORT - W_EINL) / d_rest
    objekte = OBJ * JE_OBJ
    enden = KAP_ENDEN * JE_ENDE
    textteil = text + objekte + enden
    lit_b = len(basis) / je_seite + JE_ENDE
    lit_o = len(obere_menge) / je_seite + JE_ENDE
    ges_b = textteil + lit_b
    ges_o = textteil + lit_o
    zeilen = [
        ('6.350 Wörter Absatztext: Einleitung 1.500 bei %s, Methodik bis Fazit 4.850 bei %s Wörtern je Seite'
         % (de(d_einl, 0), de(d_rest, 0)), text, 'Modellrechnung, Dichten gemessen (a)'),
        ('fünf Objekte im Textteil zu je 0,4 Seiten', objekte, 'Modellrechnung, Annahme F17 § 5.3'),
        ('fünf Kapitelenden zu je 0,5 Seiten (Umbruch vor jeder Hauptüberschrift)', enden,
         'Modellrechnung, Erwartungswert, gemessen ' + ', '.join(de(b['rest'], 2) for b in bloecke)),
        ('Textteil (Einleitung bis Fazit)', textteil, 'Modellrechnung'),
        ('Literaturverzeichnis, %d Einträge bei %s je Seite, dazu 0,5 für die letzte Seite' % (len(basis), de(je_seite, 1)),
         lit_b, 'Modellrechnung, Einträge geschätzt (c), je Seite gemessen (b)'),
        ('Einleitung bis Ende des Literaturverzeichnisses (Zählweise K1)', ges_b, 'Modellrechnung'),
        ('obere Variante mit %d Einträgen' % len(obere_menge), ges_o, 'Modellrechnung'),
    ]
    log('(d) Seitenrechnung — Modellrechnung, keine Messung (Zählweise nach K1: ohne Vorspann und Anhang):')
    for z in zeilen:
        log('    %-86s %6s  [%s]' % (z[0], de(z[1], 1), z[2]))
    log('    %-86s %6s' % ('Schwelle der Tauschregel (K3, Verfasser 28.09.)', '32'))
    log('    %-86s %6s' % ('Grenze des Verfassers (28.09., 15:14 Sitzungsuhr), mit Literaturverzeichnis', '33'))
    log('    %-86s %6s' % ('Grenze des Betreuers (mündlich), Textseiten', '37'))
    log('    Vorspann nach K1 nicht gezählt, gerendert heute %d Seiten (nur zur Information).' % vorspann)
    log('    Puffer bis 33: %s Seiten (Basis), %s Seiten (obere Variante).' % (de(33 - ges_b, 1), de(33 - ges_o, 1)))
    log('    Rahmen der Prüfungsordnung § 15 (1): „30 bis 50 Textseiten nicht überschreiten“, Obergrenze. Der Textteil liegt'
        ' mit %s Seiten darunter, bewusst (K2: R6 entfällt).' % de(textteil, 1))
    log('')
    faktor_word = 1.21
    log('Weitere Rechnungen: Word zählt nach F16 § 1.1 rund 21 %% mehr, bei 6.350 also etwa %s Wörter (Faktor aus F16,'
        ' nicht neu gemessen). Faktor zum Korpusmittel 6.350 / 4.127 = %s.' % (de(WORT * faktor_word, 0), de(WORT / 4127.0, 2)))
    log('Kontrolle: Word-Messung am fertigen Dokument nach Kapitel 5 und 6 und nach dem Literaturverzeichnis (Task 15).')

    par = [
        ('dichte_einleitung', '%.4f' % d_einl, 'Wörter je Seite', 'gemessen am Altbestand Kapitel 2 (LibreOffice-Render)', 'Seitenmodell (a)'),
        ('dichte_kapitel_4_bis_7', '%.4f' % d_rest, 'Wörter je Seite', 'gemessen an Kapitel 4 (LibreOffice-Render)', 'Seitenmodell (a)'),
        ('dichte_gesamt_vergleich', '%.4f' % dichte, 'Wörter je Seite', 'gemessen, nur zum Vergleich', 'Seitenmodell (a)'),
        ('seiten_je_objekt', '0.4', 'Seiten', 'Annahme', 'F17 § 5.3'),
        ('objekte_textteil', '5', 'Objekte', 'Festlegung', 'F17 § 5.3'),
        ('seiten_je_kapitelende', '0.5', 'Seiten', 'Annahme (Erwartungswert)', 'Seitenmodell (a)'),
        ('kapitelenden_textteil', '5', 'Kapitel', 'Festlegung', 'Gliederung v5'),
        ('eintraege_je_seite', '%.4f' % je_seite, 'Einträge je Seite', 'gemessen (Probe, LibreOffice-Render)', 'Seitenmodell (b)'),
        ('eintraege_basis', str(len(basis)), 'Einträge', 'geschätzt', 'Seitenmodell (c)'),
        ('eintraege_obere_variante', str(len(obere_menge)), 'Einträge', 'geschätzt', 'Seitenmodell (c)'),
        ('vorspann_seiten', str(vorspann), 'Seiten', 'gemessen, zählt nach K1 nicht', 'Seitenmodell (d)'),
        ('schwelle_tauschregel', '32', 'Seiten', 'Festlegung', 'Klick K3, 28.09.'),
        ('grenze_verfasser', '33', 'Seiten', 'Festlegung', 'Verfasser 28.09., 15:14'),
        ('grenze_betreuer', '37', 'Seiten', 'Festlegung', 'Betreuer, mündlich'),
        ('woerter_budget', '6350', 'Wörter', 'Festlegung', 'F17 § 5.2'),
        ('word_faktor', '1.21', 'Faktor', 'aus F16 § 1.1, nicht neu gemessen', 'F16 § 1.1'),
        ('prognose_textteil', '%.4f' % textteil, 'Seiten', 'Modellrechnung', 'Seitenmodell (d)'),
        ('prognose_gesamt_basis', '%.4f' % ges_b, 'Seiten', 'Modellrechnung', 'Seitenmodell (d)'),
        ('prognose_gesamt_obere', '%.4f' % ges_o, 'Seiten', 'Modellrechnung', 'Seitenmodell (d)'),
    ]
    with open(OUT_CSV, 'w', encoding='utf-8', newline='') as f:
        wr = csv.writer(f, delimiter=SEMI)
        wr.writerow(['parameter', 'wert', 'einheit', 'art', 'quelle'])
        wr.writerows(par)
    with open(OUT_TXT, 'w', encoding='utf-8', newline='\n') as f:
        f.write('\n'.join(LOG) + '\n')
    shutil.rmtree(arbeit, ignore_errors=True)


if __name__ == '__main__':
    main()
