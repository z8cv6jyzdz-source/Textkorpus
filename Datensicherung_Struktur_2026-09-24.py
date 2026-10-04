# -*- coding: utf-8 -*-
"""
Datensicherung_Struktur_2026-09-24.py

Zweck:    Strukturprüfung von Workbook und Fragebogen-Rohexport vor dem Einfrieren
          des Datenstands. Prüft die Regeln aus Spezifikation_2026-09-24 S01 (Regeln 2
          bis 4), S03 mit K5 (Bemerkung bei jeder ungültigen Zeile, außer bei Spielern
          mit Status "nicht angetreten"), S06 Regel 3 (kein Listenlabel mit Meldung ohne
          Analysecode) und S06 Regel 7 (Meldungen nur in W1 bis W6), dazu Pflichtfelder,
          Text in Zahlenspalten, gleiche Codemenge in beiden Blättern, Vokabular der
          Bemerkungen und Fremdzeilen unter den Spielerzeilen. Berichtet die Kontroll-
          fälle a bis g aus der Übergabe Datensicherung (Rev. 87).
Schritt:  Auswertungsverfahren_2026-09-24, Phase 1, Schritt 1.3
Eingang:  Statistik/Studiendaten_U15_gesamt.xlsx (Blätter 01_Personen, 02_Rohdaten)
              SHA-256 beim ersten Lauf: a3f665f83a97fef99e83b1e75c09cef2887f4de279c93e5467ebb536d80da6f7
          SoSci_Fragebogen/data_test546007_2026-08-31_11-37.json
              SHA-256: 7faf00a66b1bfcdfe78e100c8259b81b426bd797c006d10d756f243c1b5b40ab
          Claude/02_Befunde/Spezifikation_2026-09-24_Vokabular.csv
              SHA-256: 7afaa01d82dc9f5ca10534a09f916dca5475842caa32e71aad866ba43612860f
          Claude/02_Befunde/Spezifikation_2026-09-24_Fragebogen.csv
              SHA-256: 04318c066939248c5feca5077c65a703508a1137738310fdfa68b3a329003f04
          Die Prüfsummen werden bei jedem Lauf neu berechnet und ausgegeben. Nach einer
          Korrektur des Workbooks ändert sich dessen Prüfsumme, das ist erwartet.
Fassung:  2026-09-24, Fassung 2 (Claude, Task Datensicherung). Fassung 2: Körperhöhe, Körpermasse und
          Elterngrößen sind keine Pflichtfelder mehr (leer zulässig nach S05 Regel 9, K2), nur Hinweis.
Aufruf:   python Datensicherung_Struktur_2026-09-24.py [Basisordner]
          Basisordner = Ordner "Bachelorarbeit" (Voreinstellung: übergeordneter Ordner
          des Skripts, zwei Ebenen über Claude/03_Skripte). Ausgabe in die gleichnamige
          .txt neben dem Skript und auf die Konsole.
Regeln:   Keine Namen, kein Geburtsdatum, kein Testdatum je Spieler in der Ausgabe.
          Datumsspalten werden nur auf Vorhandensein und Typ geprüft.
"""

import sys, os, io, re, csv, json, hashlib, datetime, collections

FASSUNG = "2026-09-24, Fassung 2"

# ----------------------------------------------------------------------------
# Pfade
# ----------------------------------------------------------------------------
if len(sys.argv) > 1:
    BASIS = os.path.abspath(sys.argv[1])
else:
    BASIS = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))

P_WB   = os.path.join(BASIS, "Statistik", "Studiendaten_U15_gesamt.xlsx")
P_JSON = os.path.join(BASIS, "SoSci_Fragebogen", "data_test546007_2026-08-31_11-37.json")
P_VOK  = os.path.join(BASIS, "Claude", "02_Befunde", "Spezifikation_2026-09-24_Vokabular.csv")
P_FB   = os.path.join(BASIS, "Claude", "02_Befunde", "Spezifikation_2026-09-24_Fragebogen.csv")
P_OUT  = os.path.splitext(os.path.abspath(__file__))[0] + ".txt"

out = io.StringIO()
def w(s=""):
    out.write(s + "\n")

def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()

# ----------------------------------------------------------------------------
# Feste Umsetzungen (Spezifikation 0.4 und Übergabe 1.5)
# ----------------------------------------------------------------------------
VEREIN_MAP = {"Hohenlind": "A", "Blau-Weiß Köln": "B", "Vorwärts Spoho": "C"}
GRUPPE_MAP = {"Intervention": "IG", "Kontrolle": "KG"}
GRUPPE_SOLL = {"A": "IG", "B": "IG", "C": "KG"}
TESTS = ["Sprint_5m", "Sprint_10m", "Sprint_30m", "COD_505", "Standweitsprung"]
EINHEIT_SOLL = {"Sprint_5m": "s", "Sprint_10m": "s", "Sprint_30m": "s", "COD_505": "s", "Standweitsprung": "cm"}
ZEITPUNKTE = ["prä", "post"]
CODE_RE = re.compile(r"^(HL|BW|VS)-\d{2}$")
PFLICHT_PERSONEN = ["Code", "Verein", "Gruppe", "Geburtsdatum", "Testdatum_prä", "Status"]
# Anthropometrie und Elterngrößen dürfen fehlen (Spezifikation S05 Regel 9, K2: %PAH dann fehlend), werden aber gemeldet
OPTIONAL_PERSONEN = ["Größe_prä_cm", "Gewicht_prä_kg", "Größe_Mutter_cm", "Größe_Vater_cm"]
ZAHLENSPALTEN_PERSONEN = ["Größe_prä_cm", "Gewicht_prä_kg", "Größe_Mutter_cm", "Größe_Vater_cm", "Familiarisierung"]
DATUMSSPALTEN_PERSONEN = ["Geburtsdatum", "Testdatum_prä", "Testdatum_post"]
STATUS_NANG_PRAEFIX = "Post-Testung nicht angetreten"

# Zuordnungstabelle Fragebogen (Vorschlag Schritt 1.2 b, Übergabe: Listenlabel = Workbook-Code
# altes Schema, Ausnahme HL-04 -> HL-05, ohne realen Spieler: HL-05, BW-03, BW-12 bis BW-14)
ZUORDNUNG = {
    "HL-01": "HL-01", "HL-02": "HL-02", "HL-03": "HL-03", "HL-04": "HL-05", "HL-05": "",
    "HL-06": "HL-06", "HL-07": "HL-07", "HL-08": "HL-08",
    "BW-01": "BW-01", "BW-02": "BW-02", "BW-03": "", "BW-04": "BW-04", "BW-05": "BW-05",
    "BW-06": "BW-06", "BW-07": "BW-07", "BW-08": "BW-08", "BW-09": "BW-09", "BW-10": "BW-10",
    "BW-11": "BW-11", "BW-12": "", "BW-13": "", "BW-14": "",
}
# Programmwochen nach Kalenderregel (Spezifikation S06 Regel 7)
W_START = datetime.datetime(2026, 7, 20, 0, 0, 0)
W_ENDE  = datetime.datetime(2026, 8, 30, 23, 59, 59)

verstoesse = []   # (Kürzel, Text) -> führen in der Spezifikation zum Abbruch
hinweise = []     # (Kürzel, Text) -> zu klären oder zu korrigieren, kein Abbruch der Spezifikation
def V(k, t): verstoesse.append((k, t))
def H(k, t): hinweise.append((k, t))

# ----------------------------------------------------------------------------
# Kopf
# ----------------------------------------------------------------------------
w("Datensicherung, Schritt 1.3 Strukturprüfung")
w("Skriptfassung " + FASSUNG + " · Lauf " + datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
w("Basisordner: " + BASIS)
w()
w("Eingangsdateien mit SHA-256 (zum Laufzeitpunkt):")
for p in (P_WB, P_JSON, P_VOK, P_FB):
    if not os.path.exists(p):
        w("  FEHLT: " + p)
        print(out.getvalue()); sys.exit(2)
    w("  " + os.path.relpath(p, BASIS) + "  " + sha256(p) + "  " + str(os.path.getsize(p)) + " B  geändert " +
      datetime.datetime.fromtimestamp(os.path.getmtime(p)).strftime("%Y-%m-%d %H:%M:%S"))
w()

import openpyxl
wb = openpyxl.load_workbook(P_WB, data_only=False)
w("Blätter im Workbook: " + ", ".join(wb.sheetnames))
for b in ("01_Personen", "02_Rohdaten"):
    if b not in wb.sheetnames:
        V("S01", "Blatt fehlt: " + b)
if verstoesse:
    for k, t in verstoesse: w("VERSTOSS " + k + ": " + t)
    print(out.getvalue()); sys.exit(2)

# ----------------------------------------------------------------------------
# Vokabular und Fragebogen-Codes
# ----------------------------------------------------------------------------
with open(P_VOK, encoding="utf-8", newline="") as f:
    VOK = {r["bemerkung_exakt"]: r["kategorie"] for r in csv.DictReader(f)}
with open(P_FB, encoding="utf-8", newline="") as f:
    FBROWS = list(csv.DictReader(f))
FB_CODES = collections.defaultdict(set)
CASE_KORR = {}
for r in FBROWS:
    if r["variable"] == "CASE_KORREKTUR":
        CASE_KORR[int(r["code"])] = r["bedeutung"].replace("Listenlabel ", "")
    else:
        FB_CODES[r["variable"]].add(int(r["code"]))
H010_LABEL = {int(r["code"]): r["bedeutung"] for r in FBROWS if r["variable"] == "H010"}
w("Vokabular der Bemerkungen (Anlage): " + str(len(VOK)) + " Einträge")
w("Fallkorrekturen (Anlage): " + ", ".join("CASE %d -> %s" % (c, l) for c, l in sorted(CASE_KORR.items())))
w()

# ----------------------------------------------------------------------------
# 01_Personen
# ----------------------------------------------------------------------------
w("=" * 78)
w("1  Blatt 01_Personen")
w("=" * 78)
ws = wb["01_Personen"]
hdr = [c.value for c in ws[1]]
col = {h: i + 1 for i, h in enumerate(hdr) if h is not None}
w("Kopfzeile: " + ", ".join(str(h) for h in hdr if h is not None))
fehlende_spalten = [h for h in PFLICHT_PERSONEN + ["Familiarisierung", "Bemerkung", "Testdatum_post"] if h not in col]
if fehlende_spalten:
    V("S01", "Spalten fehlen in 01_Personen: " + ", ".join(fehlende_spalten))
    for k, t in verstoesse: w("VERSTOSS " + k + ": " + t)
    print(out.getvalue()); sys.exit(2)

personen = []      # dict je Spielerzeile
fremdzeilen = []   # (Zeile, Inhalt Spalte A/B) unterhalb oder zwischen den Spielerzeilen
for r in range(2, ws.max_row + 1):
    code = ws.cell(r, col["Code"]).value
    zeile_leer = all(ws.cell(r, c).value in (None, "") for c in range(1, ws.max_column + 1))
    if zeile_leer:
        continue
    if isinstance(code, str) and CODE_RE.match(code.strip()):
        d = {"_zeile": r}
        for h, c in col.items():
            d[h] = ws.cell(r, c).value
        personen.append(d)
    else:
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        fremdzeilen.append((r, str(a)[:40] if a is not None else "", str(b)[:60] if b is not None else ""))

w("Spielerzeilen (Code nach Muster XX-nn): %d" % len(personen))
w("Fremdzeilen (Hinweis oder Legende, kein Spieler): %d" % len(fremdzeilen))
if fremdzeilen:
    H("g", "01_Personen enthält %d Hinweis- und Legendenzeilen unterhalb der Spieler (Zeilen %s). "
           "Sie werden beim Export nicht als Spieler gelesen, im Workbook können sie bleiben."
      % (len(fremdzeilen), ", ".join(str(z) for z, _, _ in fremdzeilen)))
    for z, a, b in fremdzeilen:
        w("   Zeile %d: %s | %s" % (z, a, b))

codes_p = [d["Code"].strip() for d in personen]
dupl = [c for c, n in collections.Counter(codes_p).items() if n > 1]
if dupl:
    V("S01.3", "Code nicht eindeutig in 01_Personen: " + ", ".join(dupl))
else:
    w("Code eindeutig: ja")

def ist_zahl(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)

for d in personen:
    c = d["Code"].strip()
    z = d["_zeile"]
    # Pflichtfelder
    for h in PFLICHT_PERSONEN:
        if d.get(h) in (None, ""):
            V("Pflicht", "%s (Zeile %d): Pflichtfeld %s leer" % (c, z, h))
    for h in OPTIONAL_PERSONEN:
        if d.get(h) in (None, ""):
            H("leer", "%s (Zeile %d): %s leer, %%PAH für diesen Spieler nicht berechenbar (zulässig, S05 Regel 9)" % (c, z, h))
    # Verein und Gruppe
    ver = d.get("Verein")
    if ver not in VEREIN_MAP:
        V("S01.3", "%s: Verein unbekannt: %r" % (c, ver))
    gr = d.get("Gruppe")
    if gr not in GRUPPE_MAP:
        V("S01.3", "%s: Gruppe unbekannt: %r" % (c, gr))
    if ver in VEREIN_MAP and gr in GRUPPE_MAP:
        if GRUPPE_MAP[gr] != GRUPPE_SOLL[VEREIN_MAP[ver]]:
            V("S01.3", "%s: Gruppe %s passt nicht zum Verein %s (Soll %s)" % (c, GRUPPE_MAP[gr], VEREIN_MAP[ver], GRUPPE_SOLL[VEREIN_MAP[ver]]))
    # Code-Präfix gegen Verein
    praefix = {"HL": "Hohenlind", "BW": "Blau-Weiß Köln", "VS": "Vorwärts Spoho"}[c[:2]]
    if ver != praefix:
        V("S01.3", "%s: Code-Präfix passt nicht zum Verein %r" % (c, ver))
    # Text in Zahlenspalten
    for h in ZAHLENSPALTEN_PERSONEN:
        v = d.get(h)
        if v is None or v == "":
            continue
        if isinstance(v, str) and v.startswith("="):
            V("Text", "%s: Formel in Zahlenspalte %s" % (c, h))
        elif not ist_zahl(v):
            V("Text", "%s: Text in Zahlenspalte %s: %r" % (c, h, v))
    # Datumsspalten: nur Typ
    for h in DATUMSSPALTEN_PERSONEN:
        v = d.get(h)
        if v is None or v == "":
            continue
        if not isinstance(v, (datetime.date, datetime.datetime)):
            V("Text", "%s: %s ist kein Datumswert (Typ %s)" % (c, h, type(v).__name__))
    # Familiarisierung
    fam = d.get("Familiarisierung")
    if fam not in (None, "", 1, 2):
        V("S01.3", "%s: Familiarisierung %r nicht in {1, 2, leer}" % (c, fam))
    if fam in (None, ""):
        H("f", "%s: Familiarisierung leer (zulässig nach S01, Spieler fällt aus dem Familiarisierungsset FAMS). "
               "Laut Anwesenheitsliste nachtragen oder leer lassen (Belegprüfung)." % c)
    # Status
    st = d.get("Status")
    if st is None or st == "":
        pass  # schon als Pflichtfeld gemeldet
    elif st == "ausgewertet":
        pass
    elif isinstance(st, str) and st.startswith(STATUS_NANG_PRAEFIX):
        pass
    else:
        V("S01.3", "%s: Status %r ist weder \"ausgewertet\" noch beginnt er mit \"%s\"" % (c, st, STATUS_NANG_PRAEFIX))
    # Bemerkung (Freitext, nur Vorhandensein melden)
    if d.get("Bemerkung") not in (None, ""):
        H("Bem", "%s: Bemerkung in 01_Personen vorhanden (Freitext, geht nicht in den Datenstand). Prüfen, ob überholt." % c)

# Status-Zusammenfassung
st_cnt = collections.Counter()
for d in personen:
    st = d.get("Status")
    if st == "ausgewertet": k = "ausgewertet"
    elif isinstance(st, str) and st.startswith(STATUS_NANG_PRAEFIX): k = "nicht angetreten"
    else: k = "ANDERER TEXT"
    st_cnt[(VEREIN_MAP.get(d.get("Verein"), "?"), k)] += 1
w("Status je Verein: " + ", ".join("%s %s %d" % (v, k, n) for (v, k), n in sorted(st_cnt.items())))
andere = [d["Code"] for d in personen if not (d.get("Status") == "ausgewertet" or (isinstance(d.get("Status"), str) and d.get("Status").startswith(STATUS_NANG_PRAEFIX)))]
if andere:
    status_texte = sorted({d.get("Status") for d in personen if d["Code"] in andere})
    H("a", "Status außerhalb der festen Umsetzung bei %d Spielern (%s). Text: %r. Der Export bricht bei diesem Text ab. "
           "Nach Beleg auf \"ausgewertet\" oder \"Post-Testung nicht angetreten – Grund\" setzen."
      % (len(andere), ", ".join(andere), status_texte))
nang_codes = {d["Code"].strip() for d in personen if isinstance(d.get("Status"), str) and d.get("Status").startswith(STATUS_NANG_PRAEFIX)}
w("Status \"nicht angetreten\": " + (", ".join(sorted(nang_codes)) if nang_codes else "keiner"))
w()

# ----------------------------------------------------------------------------
# 02_Rohdaten
# ----------------------------------------------------------------------------
w("=" * 78)
w("2  Blatt 02_Rohdaten")
w("=" * 78)
ws2 = wb["02_Rohdaten"]
hdr2 = [c.value for c in ws2[1]]
w("Kopfzeile: " + ", ".join(str(h) for h in hdr2 if h is not None))
SOLL2 = ["Code", "Verein", "Zeitpunkt", "Test", "Seite", "Versuch", "Wert", "Einheit", "ungültig", "Bemerkung", "gilt"]
if hdr2[:len(SOLL2)] != SOLL2:
    V("S01", "Kopfzeile 02_Rohdaten weicht ab. Soll: " + ", ".join(SOLL2))
    for k, t in verstoesse: w("VERSTOSS " + k + ": " + t)
    print(out.getvalue()); sys.exit(2)

roh = []
for r in range(2, ws2.max_row + 1):
    vals = [ws2.cell(r, c).value for c in range(1, 12)]
    if all(v in (None, "") for v in vals[:10]):
        continue
    d = dict(zip(SOLL2, vals)); d["_zeile"] = r
    roh.append(d)
w("Datenzeilen: %d" % len(roh))
verein_von_code = {d["Code"].strip(): d.get("Verein") for d in personen}

schluessel = collections.Counter()
bem_cnt = collections.Counter()
ung_cnt = collections.Counter()
k5_faelle = []
nang_mit_wert = []
vok_fremd = collections.Counter()
for d in roh:
    z = d["_zeile"]
    c = d["Code"]
    if not (isinstance(c, str) and CODE_RE.match(c.strip())):
        V("S01.2", "Zeile %d: Code %r ohne gültiges Muster" % (z, c)); continue
    c = c.strip()
    if c not in verein_von_code:
        V("S01.2", "Zeile %d: Code %s steht nicht in 01_Personen" % (z, c))
    elif d.get("Verein") != verein_von_code[c]:
        V("S01.2", "Zeile %d: Verein %r weicht von 01_Personen ab (%r)" % (z, d.get("Verein"), verein_von_code[c]))
    if d.get("Zeitpunkt") not in ZEITPUNKTE:
        V("S01.2", "Zeile %d: Zeitpunkt %r" % (z, d.get("Zeitpunkt")))
    if d.get("Test") not in TESTS:
        V("S01.2", "Zeile %d: Test %r" % (z, d.get("Test")))
    seite = d.get("Seite")
    if d.get("Test") == "COD_505":
        if seite not in ("L", "R"):
            V("S01.2", "Zeile %d: Seite %r beim COD_505 (Soll L oder R)" % (z, seite))
    else:
        if seite not in (None, "", "–", "-"):
            V("S01.2", "Zeile %d: Seite %r bei %s (Soll leer oder –)" % (z, seite, d.get("Test")))
    if d.get("Versuch") not in (1, 2, 3):
        V("S01.2", "Zeile %d: Versuch %r" % (z, d.get("Versuch")))
    wert = d.get("Wert")
    if wert is not None and wert != "":
        if isinstance(wert, str):
            V("Text", "Zeile %d: %s %s %s Versuch %s: Wert ist Text: %r" % (z, c, d.get("Zeitpunkt"), d.get("Test"), d.get("Versuch"), wert))
        elif not ist_zahl(wert) or wert <= 0:
            V("S01.2", "Zeile %d: Wert %r nicht größer 0" % (z, wert))
    if d.get("Einheit") != EINHEIT_SOLL.get(d.get("Test")):
        V("Einheit", "Zeile %d: Einheit %r bei %s (Soll %s)" % (z, d.get("Einheit"), d.get("Test"), EINHEIT_SOLL.get(d.get("Test"))))
    ung = d.get("ungültig")
    ung_cnt[repr(ung)] += 1
    if ung not in (None, "", "x"):
        H("ung", "Zeile %d: Spalte ungültig trägt %r (erwartet leer oder x). Nach S02 zählt jeder Eintrag als ungültig." % (z, ung))
    bem = d.get("Bemerkung")
    if bem not in (None, ""):
        bem_cnt[bem] += 1
        if bem not in VOK:
            vok_fremd[bem] += 1
    key = (c, d.get("Zeitpunkt"), d.get("Test"), seite if d.get("Test") == "COD_505" else "–", d.get("Versuch"))
    schluessel[key] += 1
    # gültig_roh und K5
    gueltig_roh = (wert not in (None, "")) and (ung in (None, ""))
    d["_gueltig_roh"] = gueltig_roh
    if not gueltig_roh:
        ist_nang = (d.get("Zeitpunkt") == "post" and c in nang_codes)
        if ist_nang:
            if wert not in (None, ""):
                nang_mit_wert.append((z, c))
        elif bem in (None, ""):
            k5_faelle.append((z, c, d.get("Zeitpunkt"), d.get("Test"), seite, d.get("Versuch"), wert, ung))

dupl_keys = [k for k, n in schluessel.items() if n > 1]
if dupl_keys:
    for k in dupl_keys: V("S01.2", "Schlüssel doppelt: %s" % (k,))
else:
    w("Schlüssel (Code, Zeitpunkt, Test, Seite, Versuch) eindeutig: ja")

# Vollständigkeit des Rasters
soll_keys = set()
for c in codes_p:
    for zp in ZEITPUNKTE:
        for t in TESTS:
            for s in (("L", "R") if t == "COD_505" else ("–",)):
                for v in (1, 2, 3):
                    soll_keys.add((c, zp, t, s, v))
fehlt = soll_keys - set(schluessel)
zuviel = set(schluessel) - soll_keys
w("Raster Spieler x Zeitpunkt x Test x Seite x Versuch: Soll %d, Ist %d, fehlend %d, überzählig %d" % (len(soll_keys), len(schluessel), len(fehlt), len(zuviel)))
for k in sorted(fehlt)[:30]: H("Raster", "Zeile fehlt im Raster: %s" % (k,))
for k in sorted(zuviel)[:30]: H("Raster", "Zeile außerhalb des Rasters: %s" % (k,))

codes_r = {d["Code"].strip() for d in roh if isinstance(d.get("Code"), str)}
nur_p = set(codes_p) - codes_r
nur_r = codes_r - set(codes_p)
if nur_p: V("Codes", "Codes nur in 01_Personen: " + ", ".join(sorted(nur_p)))
if nur_r: V("Codes", "Codes nur in 02_Rohdaten: " + ", ".join(sorted(nur_r)))
if not nur_p and not nur_r:
    w("Codemenge in 01_Personen und 02_Rohdaten identisch: ja (%d Codes)" % len(codes_r))

w("Spalte ungültig, Verteilung: " + ", ".join("%s %d" % (k, n) for k, n in ung_cnt.most_common()))
w("Bemerkungen, Verteilung:")
for b, n in bem_cnt.most_common():
    w("   %-32s %4d   %s" % (b, n, ("Vokabular " + VOK[b]) if b in VOK else "NICHT IM VOKABULAR"))
for b, n in vok_fremd.items():
    V("K5", "Bemerkungstext nicht im Vokabular der Anlage: %r (%d Zeilen)" % (b, n))
    H("d", "Bemerkung %r (%d Zeilen) steht nicht im Vokabular. Vereinheitlichen auf den nächstliegenden Vokabulareintrag." % (b, n))
if "Fehversuch" in bem_cnt:
    H("d", "Bemerkung 'Fehversuch' (%d Zeile) ist nur als Schreibvariante gelistet. Vereinheitlichen auf 'Fehlversuch'." % bem_cnt["Fehversuch"])

w()
w("Ungültige Zeilen ohne Bemerkung außerhalb der Menge \"nicht angetreten\" (S03 mit K5, Abbruch der Spezifikation): %d" % len(k5_faelle))
k5_nach_code = collections.Counter((c, zp) for _, c, zp, *_ in k5_faelle)
for (c, zp), n in sorted(k5_nach_code.items()):
    w("   %s %s: %d Zeilen" % (c, zp, n))
einzeln =[f for f in k5_faelle if k5_nach_code[(f[1], f[2])] < 18]
for z, c, zp, t, s, v, wert, ung in einzeln:
    w("   Zeile %d: %s %s %s %s Versuch %s, Wert %r, ungültig %r, Bemerkung leer" % (z, c, zp, t, s, v, wert, ung))
komplett = sorted({(c, zp) for (c, zp), n in k5_nach_code.items() if n >= 18})
if komplett:
    H("b", "Spieler ohne jeden Post-Wert und ohne Bemerkung, Status nicht \"nicht angetreten\": " +
           ", ".join("%s (%s, 18 Zeilen)" % (c, zp) for c, zp in komplett) +
           ". Nach Beleg Status \"Post-Testung nicht angetreten – Grund\" setzen, dann fallen die Zeilen unter NANG.")
if einzeln:
    H("c", "Ungültige Einzelzeilen ohne Bemerkung: " +
           "; ".join("%s %s %s %s Versuch %s (Zeile %d)" % (c, zp, t, s, v, z) for z, c, zp, t, s, v, wert, ung in einzeln).replace(";", " ·") +
           ". Grund am Papierbogen prüfen und Bemerkung aus dem Vokabular eintragen.")
for z, c in nang_mit_wert:
    V("S03.1", "Zeile %d: %s post trägt einen Wert, obwohl der Status \"nicht angetreten\" lautet (Widerspruch)" % (z, c))
for k in k5_faelle:
    V("K5", "Zeile %d: ungültige Zeile ohne Bemerkung (%s %s %s %s Versuch %s)" % (k[0], k[1], k[2], k[3], k[4], k[5]))

# Sprint-Auslöseprüfung (S02 Regel 2), nur zur Information
laeufe = collections.defaultdict(dict)
for d in roh:
    if d.get("Test") in ("Sprint_5m", "Sprint_10m", "Sprint_30m") and d.get("_gueltig_roh"):
        dist = {"Sprint_5m": 5, "Sprint_10m": 10, "Sprint_30m": 30}[d["Test"]]
        laeufe[(d["Code"].strip(), d["Zeitpunkt"], d["Versuch"])][dist] = d["Wert"]
ausl = []
for k, tz in laeufe.items():
    pts = sorted(tz.items())
    d0, t0 = 0, 0.0
    for dist, t in pts:
        dt = t - t0
        if dt <= 0 or (dist - d0) / dt > 10.0:
            ausl.append((k, dist, dt))
            break
        d0, t0 = dist, t
w()
w("Sprintläufe mit mindestens einer gültigen Teilzeit: %d, davon auslösegestört nach S02 Regel 2: %d" % (len(laeufe), len(ausl)))
for k, dist, dt in ausl:
    w("   %s bis %d m, Abschnittszeit %.3f s" % (k, dist, dt))
w()

# ----------------------------------------------------------------------------
# Fragebogen-Rohexport
# ----------------------------------------------------------------------------
w("=" * 78)
w("3  Fragebogen-Rohexport (JSON)")
w("=" * 78)
with open(P_JSON, encoding="utf-8") as f:
    js = json.load(f)
meta = js.get("metadata", {})
w("Metadaten: Projekt %s, Export %s, Filter %s" % (meta.get("project"), meta.get("datetime"), meta.get("filter")))
data = js["data"]
w("Meldungen: %d" % len(data))
COLS = ["CASE", "STARTED", "H010", "H002", "H003", "H004", "H005", "H006", "H007", "H008", "H009"]
cases = collections.Counter()
labels_mit_meldung = collections.Counter()
labels_korr = collections.Counter()
for key, rec in data.items():
    case = rec.get("CASE")
    cases[case] += 1
    for cname in COLS:
        if cname not in rec:
            V("S01.4", "CASE %s: Spalte %s fehlt" % (case, cname))
    for cname in ("H010", "H002", "H003", "H004", "H005", "H006", "H007", "H008", "H009"):
        v = rec.get(cname)
        if not ist_zahl(v) or int(v) != v or int(v) not in FB_CODES[cname]:
            V("S01.4", "CASE %s: %s = %r außerhalb des Codebuchs" % (case, cname, v))
    st = rec.get("STARTED")
    try:
        ts = datetime.datetime.strptime(st, "%Y-%m-%d %H:%M:%S")
        if not (W_START <= ts <= W_ENDE):
            V("S06.7", "CASE %s: STARTED %s außerhalb W1 bis W6" % (case, st))
    except Exception:
        V("S01.4", "CASE %s: STARTED %r nicht im Format JJJJ-MM-TT hh:mm:ss" % (case, st))
    if ist_zahl(rec.get("H010")):
        lab = H010_LABEL.get(int(rec["H010"]))
        labels_mit_meldung[lab] += 1
        lab_k = CASE_KORR.get(int(case), lab)
        labels_korr[lab_k] += 1
        if int(case) in CASE_KORR:
            w("   Fallkorrektur CASE %s: Listenlabel roh %s -> %s" % (case, lab, lab_k))
dupl_cases = [c for c, n in cases.items() if n > 1]
if dupl_cases: V("S01.4", "CASE nicht eindeutig: %s" % dupl_cases)
else: w("CASE eindeutig: ja")
w("Meldungen je Listenlabel (roh -> nach Fallkorrektur):")
for lab in sorted(ZUORDNUNG):
    w("   %s  %3d -> %3d   Analysecode: %s" % (lab, labels_mit_meldung.get(lab, 0), labels_korr.get(lab, 0), ZUORDNUNG[lab] or "(kein realer Spieler)"))
for lab, n in labels_korr.items():
    if n > 0 and not ZUORDNUNG.get(lab):
        V("S06.3", "Listenlabel %s trägt %d Meldungen, hat aber keinen Analysecode" % (lab, n))
    ac = ZUORDNUNG.get(lab)
    if ac and ac not in codes_p:
        V("S06.3", "Analysecode %s (Listenlabel %s) steht nicht in 01_Personen" % (ac, lab))
    if ac and verein_von_code.get(ac) == "Vorwärts Spoho":
        V("S06.3", "Analysecode %s gehört zur KG" % ac)
ig_ohne_meldung = [c for c in codes_p if verein_von_code[c] != "Vorwärts Spoho" and labels_korr.get({v: k for k, v in ZUORDNUNG.items()}.get(c, ""), 0) == 0]
w("IG-Spieler ohne zugeordnete Meldung: " + (", ".join(ig_ohne_meldung) if ig_ohne_meldung else "keiner"))
w()

# ----------------------------------------------------------------------------
# Ergebnis
# ----------------------------------------------------------------------------
w("=" * 78)
w("4  Ergebnis")
w("=" * 78)
w("Verstöße, an denen die Spezifikation abbricht: %d" % len(verstoesse))
for k, t in verstoesse:
    w("   VERSTOSS [%s] %s" % (k, t))
w()
w("Hinweise und Kontrollfälle (a bis g nach der Übergabe): %d" % len(hinweise))
for k, t in hinweise:
    w("   [%s] %s" % (k, t))
w()
w("Kontrollfälle laut Übergabe, gefunden ja/nein:")
kf = {k for k, _ in hinweise}
for k, txt in (("a", "13 Spieler von Verein C mit Status ausstehend trotz Post-Werten"),
               ("b", "VS-07 und VS-16 ohne Post-Werte und ohne Grund"),
               ("c", "zwei ungültige Zeilen ohne Bemerkung (VS-02 post 505 R V2, VS-06 post 505 R V1)"),
               ("d", "Sytem nicht aufgenommen (2x) nicht im Vokabular, Fehversuch (1x) nur Schreibvariante"),
               ("e", "VS-18 Körperhöhe und Körpermasse als Text"),
               ("f", "VS-11 Familiarisierung leer"),
               ("g", "Hinweis- und Legendenzeilen unter den Spielerzeilen")):
    gef = (k in kf) or (k == "e" and any(t.startswith("VS-18") for kk, t in verstoesse if kk == "Text"))
    w("   %s  %-75s %s" % (k, txt, "gefunden" if gef else "NICHT gefunden"))
w()
w("Abbruchprobe bestanden: " + ("ja" if not verstoesse else "nein"))

text = out.getvalue()
with open(P_OUT, "w", encoding="utf-8", newline="\n") as f:
    f.write(text)
print(text)
sys.exit(0 if not verstoesse else 1)
