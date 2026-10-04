# -*- coding: utf-8 -*-
"""
Handprobe_Pruefung_2026-09-25.py — Rechenprobe des Handprobenblatts (Formelprobe) mit LibreOffice
Bachelorarbeit U15-Plyometrie · DSHS Köln · Auswertungsverfahren 2026-09-24 (Rev. 87), Schritt 6.3

Zweck: Das von Handprobe_2026-09-25.py (Fassung 2) erzeugte Arbeitsblatt trägt Formeln ohne gespeicherte
Werte. Dieses Skript lässt LibreOffice Calc (headless) alle Formeln berechnen, liest die Urteile der
34 Kennwerte und der Zusatzprüfungen Z1 bis Z4 zurück und schreibt das Ergebnis in eine Textdatei.
Die Prüfung ersetzt nicht das Öffnen in Excel durch den Verfasser, sie belegt nur, dass die Formeln
in einer zweiten Tabellenkalkulation dieselben Urteile liefern.

Eingang: Handprobe_2026-09-25.xlsx (Fassung 2)
Aufruf: python Handprobe_Pruefung_2026-09-25.py <Handprobe.xlsx> <Ausgabe.txt>
Fassung: 2026-09-25, erste Fassung.
"""
import sys
import os
import subprocess
import tempfile
import hashlib
from openpyxl import load_workbook

QUELLE, AUSGABE = sys.argv[1], sys.argv[2]
sha = hashlib.sha256(open(QUELLE, 'rb').read()).hexdigest()
tmp = tempfile.mkdtemp(prefix='handprobe_')
lauf = subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', tmp, QUELLE],
                      capture_output=True, text=True, timeout=600)
version = subprocess.run(['soffice', '--version'], capture_output=True, text=True).stdout.strip().split('\n')[0]
neu = os.path.join(tmp, os.path.basename(QUELLE))
if not os.path.exists(neu):
    raise SystemExit('LibreOffice hat keine Datei erzeugt: ' + lauf.stderr)
wb = load_workbook(neu, data_only=True)
sh = wb['Handprobe']

zeilen = []
zeilen.append('Rechenprobe des Handprobenblatts (Excel-Formelprobe, Phase 6.3) mit LibreOffice Calc')
zeilen.append('Datei: %s (SHA-256 %s)' % (os.path.basename(QUELLE), sha))
zeilen.append('Rechner der Probe: %s (headless, Konvertierung xlsx nach xlsx mit Neuberechnung aller Formeln)' % version)
zeilen.append('')
zeilen.append('Kennwert | Kennung | Toleranz | Urteil | relative Abweichung Formel gegen R')
maxrel = 0.0
n_stimmt = 0
n_gesamt = 0
r = 7
while sh.cell(r, 1).value is not None and isinstance(sh.cell(r, 1).value, (int, float)):
    nr, kenn, wr, wh, tol, urteil = (sh.cell(r, c).value for c in (1, 2, 6, 8, 10, 11))
    n_gesamt += 1
    if urteil == 'stimmt':
        n_stimmt += 1
    rel = ''
    if isinstance(wr, (int, float)) and isinstance(wh, (int, float)):
        d = abs(wh - wr) / max(1.0, abs(wr))
        maxrel = max(maxrel, d)
        rel = '%.3g' % d
    zeilen.append('%s | %s | %s | %s | %s' % (nr, kenn, tol, urteil, rel))
    r += 1
zeilen.append('')
zeilen.append('Kennwerte mit Urteil „stimmt“: %d von %d' % (n_stimmt, n_gesamt))
zeilen.append('Größte relative Abweichung einer Formel gegen die berichtete Rechnung: %.3g (Toleranz 1e-6, Anzahlen exakt)' % maxrel)
zeilen.append('')
zeilen.append('Zusatzprüfungen:')
n_z_ok = 0
n_z = 0
for rr in range(r, r + 8):
    a = sh.cell(rr, 1).value
    if isinstance(a, str) and a.startswith('Z') and len(a) == 2:
        n_z += 1
        u = sh.cell(rr, 11).value
        if u == 'stimmt':
            n_z_ok += 1
        zeilen.append('%s | Formelwert %s | Sollwert %s | %s' % (a, sh.cell(rr, 8).value, sh.cell(rr, 6).value, u))
zeilen.append('Zusatzprüfungen bestanden: %d von %d' % (n_z_ok, n_z))
zeilen.append('')
fw = wb['Formelwerte']
zeilen.append('Blatt Formelwerte: Summe der Abweichungen je Spieler gegen die berichtete Rechnung = %s' % fw.cell(fw.max_row, 45).value)
zeilen.append('')
zeilen.append('Gesamturteil der Rechenprobe: ' + ('bestanden' if (n_stimmt == n_gesamt and n_z_ok == n_z and n_gesamt == 34) else 'NICHT bestanden'))
with open(AUSGABE, 'w', encoding='utf-8') as f:
    f.write('\n'.join(zeilen) + '\n')
print('geschrieben:', AUSGABE, '· Urteile stimmt:', n_stimmt, 'von', n_gesamt, '· Zusatzprüfungen:', n_z_ok, 'von', n_z)
