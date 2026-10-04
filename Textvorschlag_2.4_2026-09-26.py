# -*- coding: utf-8 -*-
"""
Textvorschlag_2.4_2026-09-26.py — Task 7 des Plans (2.4 mit 2.4.1 bis 2.4.3 auf höchstens 700 Wörter)
Bachelorarbeit U15-Plyometrie · DSHS Köln

Zweck: (1) Messung des Masterbestands und des Textvorschlags je Abschnitt (Wörter als Leerraum-Token,
Sätze nach der Satzteilung von Manuskriptstand_2026-09-25.py, Median, längster Satz, Sätze über 32 und 40 Wörter,
Semikola außerhalb der Zitierklammern, nummerierte Abschnittsverweise, Belegklammern, Objekt- und Anhangsverweise),
(2) Kürzungsleiter: prüft, dass die Satz- und Teilsatzzuordnung den Masterbestand Wort für Wort wiedergibt,
summiert je Kategorie und rechnet die Stufen, dazu die optionalen Stufen unterhalb der Empfehlung,
(3) Probeexemplar NUR IM ARBEITSORDNER: Kopie des Masters mit 2.4 neu, darauf Manuskriptstand_2026-09-25.py
und Endabgleich_Manuskript_2026-09-25.py.
Der Master wird nicht verändert (Prozessregel F16 § 1.2, Verfasser 26.09.: keine Änderung am Master durch Claude).

Aufruf:
python Textvorschlag_2.4_2026-09-26.py <Master.docx> <Textvorschlag_2.4_2026-09-26.txt> <Kuerzungsleiter_2.4_2026-09-26.tsv>
       <Optionale_Stufen_2.4_2026-09-26.tsv> <Ordner 03_Skripte> <Ordner 02_Befunde> <Kennzahlen_2026-09-22.md aus dem Archiv>
       <Arbeitsordner> <Laufprotokoll.txt>
Übertragungsvorlage: Zeilen mit # sind Kommentare, eine Zeile „@@ 2.4.1“ eröffnet den Abschnitt, Absätze durch Leerzeilen getrennt.
Ohne Semikolon im Skript (chr(59)).
Fassung: 2026-09-26, erste Fassung.
"""
import sys
import os
import re
import copy
import difflib
import hashlib
import statistics
import subprocess
from collections import OrderedDict
from docx import Document
from docx.text.paragraph import Paragraph

MASTER, TXT, LEITER, OPT, SKR, BEF, KALT, ARB, LOG = sys.argv[1:10]
SEMI = chr(59)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
ABSCHNITTE = ['2.4', '2.4.1', '2.4.2', '2.4.3']
BUDGET = 700
os.makedirs(ARB, exist_ok=True)
AUS = []


def log(*t):
    AUS.append(' '.join(str(x) for x in t))


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()


def saetze(text):
    """Satzteilung wörtlich wie Manuskriptstand_2026-09-25.py (Fassung 2)"""
    t = re.sub(r'(\d)\.(\d)', r'\1<P>\2', text)
    t = re.sub(r'\b(et al|Abschn|Tab|Abb|vgl|bzw|ca|Nr|Aufl|Hrsg|Jg)\.', r'\1<P>', t)
    t = re.sub(r'\b([A-Z])\.\s', r'\1<P> ', t)
    t = re.sub(r'\bS\.\s', 'S<P> ', t)
    t = re.sub(r'\b(u|z|d)\.\s?(a|B|h)\.', r'\1<P>\2<P>', t)
    t = re.sub(r'(\d{2})\.(\d{2})\.(\d{4})', r'\1<P>\2<P>\3', t)
    t = re.sub(r'(\d{2})\.(\d{2})\.', r'\1<P>\2<P>', t)
    t = re.sub(r'(\d)\.\s', r'\1<P> ', t)
    teile = [s.strip() for s in re.split(r'(?<=[.!?])\s+(?=[A-ZÄÖÜ„(⟨])', t) if s.strip()]
    return [s.replace('<P>', '.') for s in teile]


def ohne_zitierklammern(text):
    return re.sub(r'\([^)]*\d{4}[^)]*\)', '', text)


RE_REF = re.compile(r'Abschn(?:itt)?\.?\s*\d|Kapitel\s*\d|siehe (?:oben|unten)')


def kennwerte(absaetze):
    text = ' '.join(absaetze)
    ss = [s for a in absaetze for s in saetze(a)]
    lang = [len(s.split()) for s in ss]
    return OrderedDict([
        ('Wörter', sum(len(a.split()) for a in absaetze)),
        ('Absätze', len(absaetze)),
        ('Sätze', len(ss)),
        ('Median', statistics.median(lang) if lang else 0),
        ('längster', max(lang) if lang else 0),
        ('über 32', sum(1 for x in lang if x > 32)),
        ('über 40', sum(1 for x in lang if x > 40)),
        ('Semikola', ohne_zitierklammern(text).count(SEMI)),
        ('Verweise', len(RE_REF.findall(text))),
        ('Belegklammern', len(re.findall(r'\((?:[A-ZÄÖÜ][^()]*?\d{4}[^()]*)\)', text))),
        ('Objektverweise', len(re.findall(r'(?:Tab|Abb)\.\s*H?\d+', text))),
        ('Anhangsverweise', len(re.findall(r'Anhang [A-H]', text))),
    ])


def abschnitt(doc, nr):
    ps = doc.paragraphs
    kopf = [k for k, p in enumerate(ps) if p.style.name.startswith('Heading') and p.text.strip().startswith(nr + ' ')]
    if len(kopf) != 1:
        raise SystemExit('Überschrift ' + nr + ' nicht genau einmal gefunden')
    j = kopf[0] + 1
    body = []
    while j < len(ps) and not ps[j].style.name.startswith('Heading'):
        if ps[j].text.strip():
            body.append(ps[j])
        j += 1
    return body


def vorlage_lesen(pfad):
    sek = OrderedDict()
    akt = None
    for block in re.split(r'\n\s*\n', open(pfad, encoding='utf-8').read()):
        zeilen = []
        for z in block.split('\n'):
            if z.startswith('@@'):
                akt = z[2:].strip()
                sek.setdefault(akt, [])
                continue
            if z.startswith('#') or not z.strip():
                continue
            zeilen.append(z.strip())
        if zeilen:
            sek[akt].append(' '.join(zeilen))
    return sek


# ------------------------------------------------------------------ Eingänge
log('Textvorschlag_2.4_2026-09-26.py, Laufprotokoll')
for p in (MASTER, TXT, LEITER, OPT, KALT):
    log('Eingang', os.path.basename(p), os.path.getsize(p), 'Byte, SHA-256', sha(p))
if 'word/comments.xml' in __import__('zipfile').ZipFile(MASTER).namelist():
    log('WARNUNG: Master trägt Kommentare')
else:
    log('Master ohne comments.xml')
d = Document(MASTER)
alt = OrderedDict((nr, [p.text.strip() for p in abschnitt(d, nr)]) for nr in ABSCHNITTE)
neu = vorlage_lesen(TXT)
if list(neu.keys()) != ABSCHNITTE:
    raise SystemExit('Vorlage trägt nicht genau die Abschnitte ' + ', '.join(ABSCHNITTE))

# ------------------------------------------------------------------ Messung
log('')
log('Messung je Abschnitt (Master gegen Textvorschlag)')
for nr in ABSCHNITTE:
    ka, kn = kennwerte(alt[nr]), kennwerte(neu[nr])
    log('  %-6s Master:     %s' % (nr, ' · '.join('%s %s' % (k, v) for k, v in ka.items())))
    log('  %-6s Vorschlag:  %s' % (nr, ' · '.join('%s %s' % (k, v) for k, v in kn.items())))
ka = kennwerte([a for nr in ABSCHNITTE for a in alt[nr]])
kn = kennwerte([a for nr in ABSCHNITTE for a in neu[nr]])
log('  gesamt Master:     %s' % ' · '.join('%s %s' % (k, v) for k, v in ka.items()))
log('  gesamt Vorschlag:  %s' % ' · '.join('%s %s' % (k, v) for k, v in kn.items()))
log('  Budget %d, Vorschlag %d, Abstand %+d, Kürzung gegenüber dem Master %+d' % (BUDGET, kn['Wörter'], kn['Wörter'] - BUDGET, kn['Wörter'] - ka['Wörter']))
log('')
log('Sätze des Vorschlags mit Wortzahl')
for nr in ABSCHNITTE:
    for i, a in enumerate(neu[nr], 1):
        for j, s in enumerate(saetze(a), 1):
            log('  %s Abs. %d S%d (%d): %s' % (nr, i, j, len(s.split()), s))

# ------------------------------------------------------------------ Kürzungsleiter
log('')
log('Kürzungsleiter')
zuordnung = []
for z in open(LEITER, encoding='utf-8').read().split('\n'):
    if z.strip() and not z.startswith('#'):
        a, b, s, k, g, t = z.split('\t')
        zuordnung.append((a, int(b), int(s), k, g, t))
bestand = ' '.join(a for nr in ABSCHNITTE for a in alt[nr]).split()
stuecke = ' '.join(z[5] for z in zuordnung).split()
if bestand != stuecke:
    raise SystemExit('Die Zuordnung gibt den Masterbestand nicht Wort für Wort wieder')
log('  Zuordnung gibt den Masterbestand Wort für Wort wieder: %d Wörter in %d Teilstücken' % (len(bestand), len(zuordnung)))
kat = OrderedDict([('K1', 'Selbstkommentar, Überleitung und Abschnittsverweise streichen'),
                   ('K2', 'Fehler und Zitierfallen streichen (W14, W16, T4)'),
                   ('K3', 'Methodendiskussion und Bezüge auf die eigene Studie nach 6.1 und 6.2 verlagern'),
                   ('K4', 'Rollentrennung: Mechanismus nach 2.1, Testaufbau steht in 4.4'),
                   ('K5', 'Doppelungen und entbehrliche Einzelbefunde streichen')])
summe = OrderedDict((k, sum(len(z[5].split()) for z in zuordnung if z[3] == k)) for k in list(kat) + ['U'])
stand = len(bestand)
log('  Stufe 0: Master %d' % stand)
for i, (k, text) in enumerate(kat.items(), 1):
    stand -= summe[k]
    log('  Stufe %d: %s, −%d → %d' % (i, text, summe[k], stand))
log('  Stufe 6: Übernommenes verdichten, korrigieren und ergänzen (Empfehlung), %+d → %d' % (kn['Wörter'] - summe['U'], kn['Wörter']))
if stand != summe['U']:
    raise SystemExit('Leiterrechnung stimmt nicht')
log('  Teilstücke je Kategorie:')
for k in list(kat) + ['U']:
    for z in zuordnung:
        if z[3] == k:
            log('    %s %s Abs. %d S%d (%d): %s … [%s]' % (k, z[0], z[1], z[2], len(z[5].split()), ' '.join(z[5].split()[:9]), z[4]))

# ------------------------------------------------------------------ optionale Stufen
log('')
log('Optionale Stufen unterhalb der Empfehlung (kumuliert)')
opt_text = OrderedDict((nr, list(neu[nr])) for nr in ABSCHNITTE)
stand = kn['Wörter']
for z in open(OPT, encoding='utf-8').read().split('\n'):
    if not z.strip() or z.startswith('#'):
        continue
    stufe, nr, a_alt, a_neu, preis = z.split('\t')
    treffer = [i for i, a in enumerate(opt_text[nr]) if a_alt in a]
    if len(treffer) != 1 or opt_text[nr][treffer[0]].count(a_alt) != 1:
        raise SystemExit('Optionale Stufe ' + stufe + ': Wortlaut nicht genau einmal gefunden')
    i = treffer[0]
    opt_text[nr][i] = opt_text[nr][i].replace(a_alt, a_neu, 1).strip()
    neu_stand = kennwerte([a for n in ABSCHNITTE for a in opt_text[n]])['Wörter']
    log('  %s (%s): %+d → %d · Preis: %s' % (stufe, nr, neu_stand - stand, neu_stand, preis))
    stand = neu_stand

# ------------------------------------------------------------------ Satzvergleich
log('')
log('Satzvergleich je Abschnitt (Master gegen Vorschlag)')
for nr in ABSCHNITTE:
    sa = [s for a in alt[nr] for s in saetze(a)]
    sn = [s for a in neu[nr] for s in saetze(a)]
    gleich = sum(1 for x in sn if x in sa)
    log('  %s: %d Sätze im Master, %d im Vorschlag, davon wortgleich übernommen %d' % (nr, len(sa), len(sn), gleich))

# ------------------------------------------------------------------ Probeexemplar (nur Arbeitsordner)
probe = os.path.join(ARB, 'Probeexemplar_2.4_nicht_zurueckschreiben.docx')
d2 = Document(MASTER)
for nr in ABSCHNITTE:
    body = abschnitt(d2, nr)
    if not all(p.style.name == 'Normal' for p in body):
        raise SystemExit('Abschnitt ' + nr + ' enthält Absätze außerhalb der Formatvorlage Standard')
    vorlage = body[0]
    for p in body[1:]:
        p._p.getparent().remove(p._p)
    letzte = vorlage._p
    for k, txt in enumerate(neu[nr]):
        el = vorlage._p if k == 0 else copy.deepcopy(vorlage._p)
        if k > 0:
            letzte.addnext(el)
        for r in el.findall(W + 'r'):
            el.remove(r)
        Paragraph(el, vorlage._parent).add_run(txt)
        letzte = el
d2.save(probe)
kontrolle = Document(probe)
for nr in ABSCHNITTE:
    if [p.text.strip() for p in abschnitt(kontrolle, nr)] != neu[nr]:
        raise SystemExit('Probeexemplar gibt ' + nr + ' nicht wortgleich wieder')
log('')
log('Probeexemplar', probe, os.path.getsize(probe), 'Byte (nur Arbeitsordner, wird nicht zurückgeschrieben), Abschnitte wortgleich mit der Vorlage')

ms_txt = os.path.join(ARB, 'Manuskriptstand_Probe_2.4.txt')
subprocess.run([sys.executable, os.path.join(SKR, 'Manuskriptstand_2026-09-25.py'), probe, ms_txt], check=True, capture_output=True)
ms = open(ms_txt, encoding='utf-8').read()
log('Manuskriptstand am Probeexemplar:')
for z in ms.split('\n'):
    zs = z.strip()
    if re.match(r'^2\.4(\.\d)?\s', zs) or zs.startswith('Kapitel 2') or zs.startswith('Absatztext') or zs.startswith('Semikola') or zs.startswith('Nummerierte'):
        log('  ' + zs)

ea = os.path.join(ARB, 'Endabgleich_Probe_2.4')
subprocess.run([sys.executable, os.path.join(SKR, 'Endabgleich_Manuskript_2026-09-25.py'), probe,
                os.path.join(BEF, 'Kennzahlen_2026-09-25_Werte.csv'), os.path.join(BEF, 'Kennzahlen_2026-09-25.md'),
                KALT, os.path.join(SKR, 'Programmkennzahlen_2026-09-23.txt'), os.path.join(SKR, 'Objekte_2026-09-25'), ea],
               check=True, capture_output=True)
lauf = open(os.path.join(ea, 'Endabgleich_Manuskript_2026-09-25.txt'), encoding='utf-8').read()
log('Endabgleich am Probeexemplar:')
for z in lauf.split('\n'):
    if z.startswith('Zahlen im Master') or z.startswith('Satzprüfungen:') or 'ABWEICHUNG' in z or z.strip().startswith('Wortprüfung') or z.startswith('Vorschläge'):
        log('  ' + z.strip()[:300])

with open(LOG, 'w', encoding='utf-8', newline='\n') as f:
    f.write('\n'.join(AUS) + '\n')
print('\n'.join(AUS))
