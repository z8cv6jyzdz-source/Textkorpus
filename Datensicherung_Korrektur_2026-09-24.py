# -*- coding: utf-8 -*-
"""
Datensicherung_Korrektur_2026-09-24.py

Zweck:    Schreibt die vom Verfasser freigegebenen Korrekturen (Belegprüfung) in das Workbook
          und legt das Blatt 00_Aenderungen an. Jede Änderung ist eine Zeile der Korrekturliste
          mit Blatt, Zelle, alt, neu, Grund, Beleg, Datum. Geschrieben wird auf Paketebene
          (XML der betroffenen Zellen), damit Formeln, Formate, zwischengespeicherte Werte und
          alle übrigen Zellen unverändert bleiben. Vor dem Schreiben wird der alte Inhalt jeder
          Zelle gegen die Liste geprüft, nach dem Schreiben Zelle für Zelle gegen die Sicherung
          verglichen: Geändert sind genau die protokollierten Zellen, Formeln unverändert.
Schritt:  Auswertungsverfahren_2026-09-24, Phase 1, Schritt 1.4
Eingang:  Statistik/Studiendaten_U15_gesamt.xlsx (Stand vor Phase 1,
              SHA-256 a3f665f83a97fef99e83b1e75c09cef2887f4de279c93e5467ebb536d80da6f7)
          Claude/03_Skripte/Datensicherung_Korrektur_2026-09-24_Liste.csv (Korrekturliste,
              freigegeben vom Verfasser, Spalten: Nr, Blatt, Zelle, Code, Feld, alt, neu, Typ,
              Grund, Beleg, Datum, Ausfuehrung). Typ: text, zahl oder leer. Ausfuehrung: "Skript"
              (das Skript schreibt, "alt" muss dem Zellinhalt entsprechen) oder "Verfasser" (vom
              Verfasser schon von Hand in Excel geschrieben, "neu" muss dem Zellinhalt entsprechen,
              die Zeile wird nur im Blatt 00_Aenderungen protokolliert). Sonst Abbruch. Beleg: Papierbogen, Anwesenheitsliste,
              Einwilligung oder Verfasser. Werte ohne Papierbeleg tragen "Verfasser".
Fassung:  2026-09-24, Fassung 2 (Claude, Task Datensicherung). Fassung 2: Spalte Ausfuehrung
          (Skript oder Verfasser), weil der Verfasser drei Zellen selbst in Excel geändert hat.
Aufruf:   python Datensicherung_Korrektur_2026-09-24.py [Basisordner] --pruefen
              prüft nur, ob "alt" zu jeder Zelle passt, schreibt nichts
          python Datensicherung_Korrektur_2026-09-24.py [Basisordner] --schreiben
              Sicherung nach Claude/_Archiv/Studiendaten_U15_gesamt_vor_Skriptkorrektur_JJJJ-MM-TT.xlsx
              (Stand unmittelbar vor dem Skriptlauf, hier nach den drei Handänderungen des
              Verfassers), dann Schreiben, dann Zellvergleich. Der Stand vor Phase 1 (15.09.2026,
              22:21 MESZ) liegt gesondert als Studiendaten_U15_gesamt_vor_Phase1_JJJJ-MM-TT.xlsx
              im selben Ordner. Excel muss geschlossen sein.
Regeln:   Keine Namen, kein Geburtsdatum, kein Testdatum je Spieler in der Ausgabe.
"""

import sys, os, io, re, csv, shutil, hashlib, datetime, zipfile, collections
from xml.sax.saxutils import escape

FASSUNG = "2026-09-24, Fassung 2"
HEUTE = datetime.date.today().strftime("%Y-%m-%d")

args = [a for a in sys.argv[1:] if not a.startswith("--")]
BASIS = os.path.abspath(args[0]) if args else os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
MODUS = "schreiben" if "--schreiben" in sys.argv else "pruefen"
P_WB = os.path.join(BASIS, "Statistik", "Studiendaten_U15_gesamt.xlsx")
P_LISTE = os.path.splitext(os.path.abspath(__file__))[0] + "_Liste.csv"
P_SICHER = os.path.join(BASIS, "Claude", "_Archiv", "Studiendaten_U15_gesamt_vor_Skriptkorrektur_%s.xlsx" % HEUTE)
P_OUT = os.path.splitext(os.path.abspath(__file__))[0] + ".txt"

out = io.StringIO()
def w(s=""): out.write(s + "\n")
def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""): h.update(chunk)
    return h.hexdigest()
def ende(code):
    text = out.getvalue()
    with open(P_OUT, "w", encoding="utf-8", newline="\n") as f: f.write(text)
    print(text); sys.exit(code)

w("Datensicherung, Schritt 1.4 Korrekturen mit Eintrag · Modus: " + MODUS)
w("Skriptfassung " + FASSUNG + " · Lauf " + datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
w("Workbook: " + P_WB + "  SHA-256 " + sha256(P_WB) + "  geändert " + datetime.datetime.fromtimestamp(os.path.getmtime(P_WB)).strftime("%Y-%m-%d %H:%M:%S"))
w("Korrekturliste: " + P_LISTE + "  SHA-256 " + sha256(P_LISTE))
w()

# ----------------------------------------------------------------------------
# Korrekturliste lesen
# ----------------------------------------------------------------------------
with open(P_LISTE, encoding="utf-8", newline="") as f:
    LISTE = list(csv.DictReader(f))
SOLL_SPALTEN = ["Nr", "Blatt", "Zelle", "Code", "Feld", "alt", "neu", "Typ", "Grund", "Beleg", "Datum", "Ausfuehrung"]
if list(LISTE[0].keys()) != SOLL_SPALTEN:
    w("ABBRUCH: Spalten der Korrekturliste weichen ab: " + ", ".join(LISTE[0].keys())); ende(2)
w("Korrekturen in der Liste: %d" % len(LISTE))
zellen = collections.Counter((r["Blatt"], r["Zelle"]) for r in LISTE)
if any(n > 1 for n in zellen.values()):
    w("ABBRUCH: eine Zelle steht mehrfach in der Liste: " + str([k for k, n in zellen.items() if n > 1])); ende(2)

import openpyxl
wb = openpyxl.load_workbook(P_WB, data_only=False)
def zellinhalt(ws, ref):
    v = ws[ref].value
    return "" if v is None else (repr(v) if not isinstance(v, str) else v)
def zahl_text(v):
    # Darstellung einer Zahl wie in der Liste (z. B. 1.94 oder 166)
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v)

fehler = 0
for r in LISTE:
    ws = wb[r["Blatt"]]
    ist = ws[r["Zelle"]].value
    ist_txt = "" if ist is None else (zahl_text(ist) if not isinstance(ist, str) else ist)
    if r["Ausfuehrung"] not in ("Skript", "Verfasser"):
        w("   ABBRUCH: Ausfuehrung %r unbekannt" % r["Ausfuehrung"]); fehler += 1
    soll = r["alt"] if r["Ausfuehrung"] == "Skript" else r["neu"]
    ok = (ist_txt == soll)
    if not ok: fehler += 1
    w("%-3s %-11s %-6s %-6s %-22s %-9s alt=%r  neu=%r  ist=%r  %s" % (r["Nr"], r["Blatt"], r["Zelle"], r["Code"], r["Feld"], r["Ausfuehrung"], r["alt"], r["neu"], ist_txt, "ok" if ok else "PASST NICHT"))
    if r["Typ"] not in ("text", "zahl", "leer"):
        w("   ABBRUCH: Typ %r unbekannt" % r["Typ"]); fehler += 1
    if r["Typ"] == "zahl":
        try: float(r["neu"])
        except ValueError: w("   ABBRUCH: neu %r ist keine Zahl" % r["neu"]); fehler += 1
    if r["Typ"] == "leer" and r["neu"] != "":
        w("   ABBRUCH: Typ leer, aber neu nicht leer"); fehler += 1
    if isinstance(ist, str) and ist.startswith("="):
        w("   ABBRUCH: Zelle trägt eine Formel, Formeln werden nicht geändert"); fehler += 1
if fehler:
    w("\nABBRUCH: %d Zellen passen nicht zur Liste. Nichts geschrieben." % fehler); ende(2)
w("\nAlle alten Werte passen zur Liste.")
if MODUS == "pruefen":
    w("Modus --pruefen: nichts geschrieben."); ende(0)

# ----------------------------------------------------------------------------
# Sicherung
# ----------------------------------------------------------------------------
os.makedirs(os.path.dirname(P_SICHER), exist_ok=True)
if os.path.exists(P_SICHER):
    w("ABBRUCH: Sicherung existiert schon: " + P_SICHER); ende(2)
shutil.copy2(P_WB, P_SICHER)
w("Sicherung: " + P_SICHER + "  SHA-256 " + sha256(P_SICHER))

# ----------------------------------------------------------------------------
# Paket patchen
# ----------------------------------------------------------------------------
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
zin = zipfile.ZipFile(P_SICHER, "r")
teile = {n: zin.read(n) for n in zin.namelist()}
infos = {i.filename: i for i in zin.infolist()}
zin.close()

# Blattname -> Teil
wbxml = teile["xl/workbook.xml"].decode("utf-8")
rels = teile["xl/_rels/workbook.xml.rels"].decode("utf-8")
rid_ziel = dict(re.findall(r'<Relationship Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
rid_ziel.update({m[1]: m[0] for m in re.findall(r'<Relationship [^>]*Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)})
blatt_teil = {}
for name, sid, rid in re.findall(r'<sheet name="([^"]+)" sheetId="(\d+)" r:id="(rId\d+)"/>', wbxml):
    blatt_teil[name.replace("&amp;", "&")] = "xl/" + rid_ziel[rid]
w("Blattzuordnung: " + ", ".join("%s -> %s" % kv for kv in blatt_teil.items()))

def spalten_index(ref):
    buchst = re.match(r"([A-Z]+)", ref).group(1)
    n = 0
    for ch in buchst: n = n * 26 + (ord(ch) - 64)
    return n

def patch_zelle(xml, ref, typ, neu):
    m = re.search(r'<c r="%s"(?P<attr>[^>]*?)(?:/>|>(?P<inner>.*?)</c>)' % ref, xml, flags=re.S)
    if not m:
        raise RuntimeError("Zelle %s nicht im XML gefunden" % ref)
    attr = m.group("attr")
    if "<f" in (m.group("inner") or ""):
        raise RuntimeError("Zelle %s trägt eine Formel" % ref)
    stil = re.search(r'\ss="(\d+)"', attr)
    s_attr = ' s="%s"' % stil.group(1) if stil else ""
    if typ == "leer":
        neu_xml = '<c r="%s"%s/>' % (ref, s_attr)
    elif typ == "zahl":
        z = float(neu)
        neu_xml = '<c r="%s"%s><v>%s</v></c>' % (ref, s_attr, (str(int(z)) if z.is_integer() else repr(z)))
    else:
        neu_xml = '<c r="%s"%s t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>' % (ref, s_attr, escape(neu))
    return xml[:m.start()] + neu_xml + xml[m.end():], m.group(0)

alte_xml = {}
SKRIPT = [r for r in LISTE if r["Ausfuehrung"] == "Skript"]
for r in SKRIPT:
    teil = blatt_teil[r["Blatt"]]
    xml = teile[teil].decode("utf-8")
    xml, alt_xml = patch_zelle(xml, r["Zelle"], r["Typ"], r["neu"])
    teile[teil] = xml.encode("utf-8")
    alte_xml[(r["Blatt"], r["Zelle"])] = alt_xml

# Neues Blatt 00_Aenderungen (Inline-Strings, keine Formeln)
def zelle_xml(ref, v):
    if v is None or v == "": return '<c r="%s"/>' % ref
    try:
        z = float(v)
        if isinstance(v, str) and not re.fullmatch(r"-?\d+(\.\d+)?", v.strip()): raise ValueError
        return '<c r="%s"><v>%s</v></c>' % (ref, (str(int(z)) if z.is_integer() else repr(z)))
    except (ValueError, TypeError):
        return '<c r="%s" t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>' % (ref, escape(str(v)))
spalten = "ABCDEFGHIJKL"
kopf = ["Nr", "Blatt", "Zelle", "Code", "Feld", "alt", "neu", "Grund", "Beleg", "Datum", "geschrieben von", "Skript"]
zeilen_xml = ['<row r="1">' + "".join(zelle_xml("%s1" % spalten[j], h) for j, h in enumerate(kopf)) + "</row>"]
for i, r in enumerate(LISTE, 2):
    werte = [r["Nr"], r["Blatt"], r["Zelle"], r["Code"], r["Feld"], r["alt"], r["neu"], r["Grund"], r["Beleg"], r["Datum"],
             ("Claude (Task Datensicherung), freigegeben vom Verfasser" if r["Ausfuehrung"] == "Skript" else "Verfasser von Hand in Excel, vom Skript protokolliert"),
             os.path.basename(__file__)]
    zeilen_xml.append('<row r="%d">' % i + "".join(zelle_xml("%s%d" % (spalten[j], i), v) for j, v in enumerate(werte)) + "</row>")
n_z = len(LISTE) + 1
sheet_neu = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<worksheet xmlns="%s" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
             '<dimension ref="A1:L%d"/><sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>'
             '<sheetFormatPr defaultRowHeight="14.4"/>'
             '<cols><col min="1" max="1" width="5" customWidth="1"/><col min="2" max="2" width="13" customWidth="1"/><col min="3" max="3" width="8" customWidth="1"/>'
             '<col min="4" max="4" width="8" customWidth="1"/><col min="5" max="5" width="18" customWidth="1"/><col min="6" max="7" width="40" customWidth="1"/>'
             '<col min="8" max="8" width="60" customWidth="1"/><col min="9" max="9" width="22" customWidth="1"/><col min="10" max="10" width="11" customWidth="1"/>'
             '<col min="11" max="11" width="42" customWidth="1"/><col min="12" max="12" width="40" customWidth="1"/></cols>'
             '<sheetData>%s</sheetData><pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/></worksheet>'
             % (NS, n_z, "".join(zeilen_xml)))
neu_nr = max(int(m) for m in re.findall(r"worksheets/sheet(\d+)\.xml", "\n".join(teile))) + 1
neu_teil = "xl/worksheets/sheet%d.xml" % neu_nr
neu_rid = "rId%d" % (max(int(m) for m in re.findall(r'Id="rId(\d+)"', rels)) + 1)
neu_sid = max(int(s) for s in re.findall(r'sheetId="(\d+)"', wbxml)) + 1
teile[neu_teil] = sheet_neu.encode("utf-8")
rels = rels.replace("</Relationships>", '<Relationship Id="%s" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet%d.xml"/></Relationships>' % (neu_rid, neu_nr))
teile["xl/_rels/workbook.xml.rels"] = rels.encode("utf-8")
wbxml = wbxml.replace("<sheets>", '<sheets><sheet name="00_Aenderungen" sheetId="%d" r:id="%s"/>' % (neu_sid, neu_rid))
# aktives Blatt bleibt 01_Personen (Index verschiebt sich um eins), volle Neuberechnung beim Öffnen
wbxml = re.sub(r'activeTab="(\d+)"', lambda m: 'activeTab="%d"' % (int(m.group(1)) + 1), wbxml)
if "fullCalcOnLoad" not in wbxml:
    wbxml = wbxml.replace("<calcPr ", '<calcPr fullCalcOnLoad="1" ')
teile["xl/workbook.xml"] = wbxml.encode("utf-8")
ct = teile["[Content_Types].xml"].decode("utf-8")
ct = ct.replace("</Types>", '<Override PartName="/%s" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>' % neu_teil)
teile["[Content_Types].xml"] = ct.encode("utf-8")
app = teile["docProps/app.xml"].decode("utf-8")
app = re.sub(r"<vt:i4>(\d+)</vt:i4>", lambda m: "<vt:i4>%d</vt:i4>" % (int(m.group(1)) + 1), app, count=1)
app = re.sub(r'<vt:vector size="(\d+)" baseType="lpstr">', lambda m: '<vt:vector size="%d" baseType="lpstr"><vt:lpstr>00_Aenderungen</vt:lpstr>' % (int(m.group(1)) + 1), app, count=1)
teile["docProps/app.xml"] = app.encode("utf-8")
# tabSelected nur auf dem bisher aktiven Blatt belassen (neues Blatt ist nicht ausgewählt)

tmp = P_WB + ".neu"
with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for name in list(infos) + [neu_teil]:
        zi = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
        zi.compress_type = zipfile.ZIP_DEFLATED
        zout.writestr(zi, teile[name])
os.replace(tmp, P_WB)
w("Geschrieben: " + P_WB + "  SHA-256 " + sha256(P_WB) + "  Blatt 00_Aenderungen mit %d Zeilen, davon %d vom Skript geschrieben" % (len(LISTE), len(SKRIPT)))

# ----------------------------------------------------------------------------
# Zellvergleich gegen die Sicherung
# ----------------------------------------------------------------------------
w()
w("Zellvergleich neu gegen Sicherung (Formeln als Text, Werte als Werte):")
alt = openpyxl.load_workbook(P_SICHER, data_only=False)
neu = openpyxl.load_workbook(P_WB, data_only=False)
erwartet = {(r["Blatt"], r["Zelle"]) for r in SKRIPT}
gefunden = set()
formel_geaendert = 0
sonst = 0
if set(neu.sheetnames) != set(alt.sheetnames) | {"00_Aenderungen"}:
    w("   ABBRUCH: Blattmenge weicht ab: " + str(neu.sheetnames)); ende(2)
for name in alt.sheetnames:
    wa, wn = alt[name], neu[name]
    mr, mc = max(wa.max_row, wn.max_row), max(wa.max_column, wn.max_column)
    for rr in range(1, mr + 1):
        for cc in range(1, mc + 1):
            va, vn = wa.cell(rr, cc).value, wn.cell(rr, cc).value
            if va != vn:
                ref = wn.cell(rr, cc).coordinate
                if (name, ref) in erwartet:
                    gefunden.add((name, ref))
                    w("   %s!%s: %r -> %r (protokolliert)" % (name, ref, va, vn))
                else:
                    sonst += 1
                    w("   ABWEICHUNG NICHT PROTOKOLLIERT %s!%s: %r -> %r" % (name, ref, va, vn))
                if isinstance(va, str) and va.startswith("="):
                    formel_geaendert += 1
    # Formeln: Zahl je Blatt
    fa = sum(1 for row in wa.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    fn = sum(1 for row in wn.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    w("   %s: Formeln vorher %d, nachher %d" % (name, fa, fn))
nicht_gefunden = erwartet - gefunden
w()
w("Protokollierte Zellen geändert: %d von %d" % (len(gefunden), len(erwartet)))
if nicht_gefunden:
    w("   Zellen aus der Liste ohne sichtbare Änderung (alt = neu?): " + str(sorted(nicht_gefunden)))
w("Nicht protokollierte Abweichungen: %d" % sonst)
w("Geänderte Formelzellen: %d" % formel_geaendert)
ok = (sonst == 0 and formel_geaendert == 0 and not nicht_gefunden)
w("Zellvergleich bestanden: " + ("ja" if ok else "NEIN"))
# Zwischengespeicherte Werte: Hinweis
w()
w("Hinweis: Zwischengespeicherte Formelwerte (z. B. Spalte gilt in 02_Rohdaten) werden beim nächsten Öffnen in Excel neu berechnet (fullCalcOnLoad gesetzt). Der Datenstand verwendet keine Formelspalten.")
ende(0 if ok else 1)
