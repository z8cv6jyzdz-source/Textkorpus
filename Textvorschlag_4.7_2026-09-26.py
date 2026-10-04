# -*- coding: utf-8 -*-
"""
Textvorschlag_4.7_2026-09-26.py — Task 6 des Plans (4.7 auf 550 Wörter)
Bachelorarbeit U15-Plyometrie · DSHS Köln

Zweck: (1) Satzvergleich 4.7 im Master gegen den Textvorschlag mit Wortdifferenz je Stelle,
(2) Messung des Textvorschlags (Wörter als Leerzeichen-Token, Sätze, Median, Maximum, Semikola,
nummerierte Abschnittsverweise, Belegklammern, Objekt- und Anhangsverweise),
(3) Probeexemplar NUR IM ARBEITSORDNER: Kopie des Masters mit 4.7 neu und 4.6 im Wortlaut Rev. 105,
darauf Manuskriptstand_2026-09-25.py und Endabgleich_Manuskript_2026-09-25.py.
Der Master wird nicht verändert (Prozessregel F16 § 1.2, Verfasser 26.09.: keine Änderung am Master durch Claude).

Aufruf:
python Textvorschlag_4.7_2026-09-26.py <Master.docx> <Textvorschlag_4.7_2026-09-26.txt> <Wortlaut_4.6_Rev105_2026-09-26.txt>
       <Ordner 03_Skripte> <Ordner 02_Befunde> <Kennzahlen_2026-09-22.md aus dem Archiv> <Arbeitsordner> <Laufprotokoll.txt>
       [<Satzersetzungen.tsv>]
Das optionale neunte Argument ist eine Tabulator-getrennte Datei mit den Spalten Abschnitt, alter Wortlaut, neuer Wortlaut
(Zeilen mit # sind Kommentare). Die Ersetzungen gelten nur im Probeexemplar und werden je Abschnitt gemessen.
Ohne Semikolon im Skript (chr(59)).
Fassung: 2026-09-26, zweite Fassung (Satzersetzungen für 4.1 und 4.4.2 nach der Rückfrage des Verfassers zu Absatz 1).
"""
import sys
import os
import re
import copy
import difflib
import hashlib
import statistics
import subprocess
from docx import Document
from docx.text.paragraph import Paragraph

MASTER, T47, T46, SKR, BEF, KALT, ARB, LOG = sys.argv[1:9]
ERS = sys.argv[9] if len(sys.argv) > 9 else None
SEMI = chr(59)
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
os.makedirs(ARB, exist_ok=True)
AUS = []


def log(*t):
    AUS.append(' '.join(str(x) for x in t))


def sha(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()


def absaetze_txt(pfad):
    t = open(pfad, encoding='utf-8').read()
    t = '\n'.join(z for z in t.split('\n') if not z.startswith('#'))
    return [a.strip() for a in re.split(r'\n\s*\n', t) if a.strip()]


ABK = ['S.', 'et al.', 'Abs.', 'Nr.', 'Tab.', 'Abb.', 'vgl.', 'bzw.', 'z. B.', 'u. a.', 'ca.', 'Hrsg.', 'Aufl.', 'Std.', 'min.', 'ggf.', 'Abschn.']


def saetze(a):
    """Satzteilung wie Messen_Text_2026-09-25.py"""
    s = a
    for ab in ABK:
        s = s.replace(ab, ab.replace('.', '\u2024'))
    s = re.sub(r'(\d)\.(\d)', '\\1\u2024\\2', s)
    s = re.sub(r'(\d{2})\.(\d{4})', '\\1\u2024\\2', s)
    s = re.sub(r'(\d{2})\.\s', '\\1\u2024 ', s)
    teile = re.split(r'(?<=[.!?])\s+(?=[A-ZÄÖÜ„(0-9])', s)
    return [t.replace('\u2024', '.') for t in teile if t.strip()]


def abschnitt(doc, nr):
    ps = doc.paragraphs
    kopf = [k for k, p in enumerate(ps) if p.style.name.startswith('Heading') and p.text.strip().startswith(nr + ' ')]
    if len(kopf) != 1:
        raise SystemExit('Überschrift ' + nr + ' nicht genau einmal gefunden')
    j = kopf[0] + 1
    body = []
    while j < len(ps) and not ps[j].style.name.startswith('Heading'):
        if ps[j].text.strip():
            body.append(ps[j])
        j += 1
    return body


def messen(name, absaetze):
    alle = []
    log('Messung', name)
    for i, a in enumerate(absaetze, 1):
        ss = saetze(a)
        alle.extend(ss)
        log('  Absatz %d: %d Wörter, %d Sätze, längster %d' % (i, len(a.split()), len(ss), max(len(x.split()) for x in ss)))
    lang = [len(x.split()) for x in alle]
    text = ' '.join(absaetze)
    ohne = re.sub(r'\([^)]*\)', '', text)
    verw = re.findall(r'Abschn(?:itt)?\.?\s*\d|Kapitel\s*\d|siehe (?:oben|unten)', text)
    belege = re.findall(r'\((?:[A-ZÄÖÜ][^()]*?\d{4}[^()]*)\)', text)
    log('  gesamt %d Wörter · %d Sätze · Median %.1f · längster %d · über 32: %d · über 40: %d' % (
        sum(len(a.split()) for a in absaetze), len(lang), statistics.median(lang), max(lang),
        sum(1 for x in lang if x > 32), sum(1 for x in lang if x > 40)))
    log('  Semikola gesamt %d, außerhalb von Klammern %d · Abschnittsverweise %d · Belegklammern %d · Objektverweise %s · Anhangsverweise %s' % (
        text.count(SEMI), ohne.count(SEMI), len(verw), len(belege),
        re.findall(r'(?:Tab|Abb)\.\s*H?\d+', text), re.findall(r'Anhang [A-H]', text)))
    return alle


# ------------------------------------------------------------------ Eingänge
log('Textvorschlag_4.7_2026-09-26.py, Laufprotokoll')
for p in (MASTER, T47, T46, KALT) + ((ERS,) if ERS else ()):
    log('Eingang', os.path.basename(p), os.path.getsize(p), 'Byte, SHA-256', sha(p))
d = Document(MASTER)
alt47 = [p.text.strip() for p in abschnitt(d, '4.7')]
alt46 = [p.text.strip() for p in abschnitt(d, '4.6')]
neu47 = absaetze_txt(T47)
neu46 = absaetze_txt(T46)
if 'word/comments.xml' in __import__('zipfile').ZipFile(MASTER).namelist():
    log('WARNUNG: Master trägt Kommentare')
else:
    log('Master ohne comments.xml')

# ------------------------------------------------------------------ Messung
log('')
s_alt = messen('4.7 im Master', alt47)
s_neu = messen('4.7 Textvorschlag', neu47)
messen('4.6 im Master', alt46)
messen('4.6 Wortlaut Rev. 105', neu46)

# ------------------------------------------------------------------ Satzvergleich 4.7
log('')
log('Satzvergleich 4.7 (Master gegen Textvorschlag), Wortdifferenz je Stelle')
sm = difflib.SequenceMatcher(None, s_alt, s_neu, autojunk=False)
summe = 0
for op, i1, i2, j1, j2 in sm.get_opcodes():
    if op == 'equal':
        log('  gleich: %d Sätze' % (i2 - i1))
        continue
    wa = sum(len(x.split()) for x in s_alt[i1:i2])
    wn = sum(len(x.split()) for x in s_neu[j1:j2])
    summe += wn - wa
    log('  %s (%+d):' % (op, wn - wa))
    for x in s_alt[i1:i2]:
        log('     alt (%d): %s' % (len(x.split()), x))
    for x in s_neu[j1:j2]:
        log('     neu (%d): %s' % (len(x.split()), x))
log('  Summe der Differenzen %+d' % summe)
log('')
log('Satzvergleich 4.6 (Master gegen Wortlaut Rev. 105)')
for zeile in difflib.unified_diff(alt46, neu46, lineterm='', n=0):
    log('  ' + zeile[:400])

# ------------------------------------------------------------------ Probeexemplar (nur Arbeitsordner)
probe = os.path.join(ARB, 'Probeexemplar_nicht_zurueckschreiben.docx')
d2 = Document(MASTER)
for nr, texte in (('4.7', neu47), ('4.6', neu46)):
    body = abschnitt(d2, nr)
    if not all(p.style.name == 'Normal' for p in body):
        raise SystemExit('Abschnitt ' + nr + ' enthält Absätze außerhalb der Formatvorlage Standard')
    vorlage = body[0]
    for p in body[1:]:
        p._p.getparent().remove(p._p)
    letzte = vorlage._p
    for k, txt in enumerate(texte):
        el = vorlage._p if k == 0 else copy.deepcopy(vorlage._p)
        if k > 0:
            letzte.addnext(el)
        for r in el.findall(W + 'r'):
            el.remove(r)
        Paragraph(el, vorlage._parent).add_run(txt)
        letzte = el
ersetzungen = []
if ERS:
    for z in open(ERS, encoding='utf-8').read().split('\n'):
        if z.strip() and not z.startswith('#'):
            nr, alt, neu = z.split('\t')
            ersetzungen.append((nr.strip(), alt, neu))
for nr, alt, neu in ersetzungen:
    treffer = [p for p in abschnitt(d2, nr) if alt in p.text]
    if len(treffer) != 1:
        raise SystemExit('Satzersetzung in ' + nr + ': alter Wortlaut nicht genau einmal gefunden')
    p = treffer[0]
    vorher = p.text
    for r in p._p.findall(W + 'r'):
        p._p.remove(r)
    p.add_run(vorher.replace(alt, neu, 1))
    log('Satzersetzung', nr, '(%+d Wörter):' % (len(neu.split()) - len(alt.split())))
    log('     alt:', alt)
    log('     neu:', neu)
d2.save(probe)
log('')
log('Probeexemplar', probe, os.path.getsize(probe), 'Byte (nur Arbeitsordner, wird nicht zurückgeschrieben)')

ms_txt = os.path.join(ARB, 'Manuskriptstand_Probe.txt')
subprocess.run([sys.executable, os.path.join(SKR, 'Manuskriptstand_2026-09-25.py'), probe, ms_txt], check=True, capture_output=True)
ms = open(ms_txt, encoding='utf-8').read()
log('Manuskriptstand am Probeexemplar:')
for z in ms.split('\n'):
    if re.match(r'^\s+(4\.\d(\.\d)?)\s', z) or z.startswith('Kapitel 4') or z.startswith('Absatztext') or z.startswith('Semikola') or z.startswith('Nummerierte'):
        log('  ' + z.strip())

ea = os.path.join(ARB, 'Endabgleich_Probe')
subprocess.run([sys.executable, os.path.join(SKR, 'Endabgleich_Manuskript_2026-09-25.py'), probe,
                os.path.join(BEF, 'Kennzahlen_2026-09-25_Werte.csv'), os.path.join(BEF, 'Kennzahlen_2026-09-25.md'),
                KALT, os.path.join(SKR, 'Programmkennzahlen_2026-09-23.txt'), os.path.join(SKR, 'Objekte_2026-09-25'), ea],
               check=True, capture_output=True)
lauf = open(os.path.join(ea, 'Endabgleich_Manuskript_2026-09-25.txt'), encoding='utf-8').read()
log('Endabgleich am Probeexemplar:')
for z in lauf.split('\n'):
    if z.startswith('Zahlen im Master') or z.startswith('Satzprüfungen:') or 'ABWEICHUNG' in z or 'SP9c' in z or z.strip().startswith('Wortprüfung') or z.startswith('Vorschläge'):
        log('  ' + z.strip()[:300])

with open(LOG, 'w', encoding='utf-8', newline='\n') as f:
    f.write('\n'.join(AUS) + '\n')
print('\n'.join(AUS))
