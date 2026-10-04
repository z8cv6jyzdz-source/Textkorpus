# -*- coding: utf-8 -*-
"""
Steuerdokumente_Pruefung_2026-09-28.py — Prüfung der Folgeänderungen im Task Steuerdokumente 28.09. (Schritt 4, Übergabe
`04_Uebergaben\\Uebergabe_Steuerdokumente_2026-09-28.md` § 6: für Gliederung v5, Berichtsraster Rev. 3 und Plan Rev. 5 „(a) nur an
neuen oder geänderten Stellen, (e) und (f) über das ganze Dokument, dazu die Trefferliste des Plans“), erweitert um die Übergabe
Einleitung, README, Aufräumskript und Archiv.

(a) Semikolon an neuen oder geänderten Stellen (zeichengenauer Vergleich mit dem Vorstand), außer zwischen Quellen in einer Klammer
(e) Budgetarithmetik: Gliederung v5 § 3.1 und § 3.4, Kopfblöcke des Berichtsrasters, Plan § 5, jeweils gegen Fassung 17 § 5.2 und
    die Ausgabe des Messskripts
(f) jede Angabe „§ x“ oder „§ x.y“ ohne fremden Dokumentnamen davor zeigt auf einen Paragrafen des Dokuments (ein Verweis auf
    einen hier nicht vorhandenen Paragrafen gilt als fremd, wenn vorher in derselben Zeile ein anderes Dokument genannt ist)
(i) Trefferliste des Plans („Durchgehend“ aus Übergabe § 4, dazu „Gliederung v4“, „Berichtsraster Rev. 2“, „Klickfrage 10“,
    „Klickfrage 11“): jeder Treffer ersetzt, historisch, erledigt oder entfallen
(g) Probe: Messskript Fassung 3 misst die Einleitung nach der Übertragung ohne Anpassung (Probe-Master = Master ohne die
    Überschriften 2 bis 3 samt Text, mit den neun Absätzen des Textvorschlags unter „1 Einleitung“, Formatvorlage Standard)
(h) Aufräumskript und Archiv: BOM und CRLF, jeder neue Papierkorbeintrag hat eine Archivkopie mit dem MD5 des Originals, Muster
    00_Steuerung auf Fassung 17, README auf Fassung 17, Plan Rev. 5, Berichtsraster Rev. 3 und Gliederung v5
(k) Übergabe Einleitung: „F16 §“ nur noch historisch in § 2, keine Seitenzahlen der alten Modellrechnung
(l) Nach der übergreifenden Zweitprüfung und den Klicks vom 28.09., 19:25: Erratum Khamis und Roche nirgends als Inhalt der Arbeit,
    Tab. H6 ohne R14, Verdünnungslogik in 6.1 und G3 (6.2 ohne eigenen Absatz), Objekte erst in Task 18 ohne Platzhalter im Text,
    „vorgemerkt für“ an leeren Zielorten, Vorstände von README, Aufräumskript und Übergabe Einleitung im Archiv (unter (h))
Fassung 2 (28.09., nach der übergreifenden Zweitprüfung): (l) neu, (h) mit den Vorständen.

Aufruf: python Steuerdokumente_Pruefung_2026-09-28.py <Vorstand Claude-Ordner> <Arbeitsordner> <F17.md> <Master.docx>
        <Manuskriptstand.py> <Seitenmodell.csv> <Ausgabe.txt>
Ohne Semikolon im Skript (chr(59)).
"""
import copy
import csv
import difflib
import hashlib
import os
import re
import subprocess
import sys
import tempfile

ALT, NEU, F17, MASTER, MSKRIPT, S_CSV, OUT = sys.argv[1:8]
SEMI = chr(59)
f17 = open(F17, encoding='utf-8').read()
aus = []
fehler = []


def lies(p):
    return open(p, encoding='utf-8').read()


def ergebnis(name, treffer, zusatz=''):
    if treffer:
        fehler.append(name)
        aus.extend('   ' + x for x in treffer)
    aus.append('   Ergebnis: %s%s' % ('NICHT BESTANDEN' if treffer else 'bestanden', zusatz))


def zahl(s):
    return int(s.replace('.', ''))


PAARE = [('Gliederung v5', '01_Verfahren/Gliederung_2026-09-23.md', '01_Verfahren/Gliederung_2026-09-28.md'),
         ('Berichtsraster Rev. 3', '02_Befunde/Berichtsraster_2026-09-23.md', '02_Befunde/Berichtsraster_2026-09-23.md'),
         ('Plan Rev. 5', '04_Uebergaben/Plan_Weitere_Schritte_2026-09-25.md', '04_Uebergaben/Plan_Weitere_Schritte_2026-09-25.md'),
         ('Übergabe Einleitung', '04_Uebergaben/Uebergabe_Einleitung_2026-09-28.md', '04_Uebergaben/Uebergabe_Einleitung_2026-09-28.md'),
         ('README', 'README_Ordnerstruktur.md', 'README_Ordnerstruktur.md')]
DOK = {name: (lies(os.path.join(ALT, a)), lies(os.path.join(NEU, n))) for name, a, n in PAARE}


# ------------------------------------------------------------------ (a) Semikolon an neuen oder geänderten Stellen
def neue_stellen(alt, neu):
    a, b = alt.split('\n'), neu.split('\n')
    res = {}
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes():
        if tag == 'equal' or tag == 'delete':
            continue
        for j in range(j1, j2):
            zb = b[j]
            best, r = '', 0.0
            for za in a[i1:i2]:
                sm = difflib.SequenceMatcher(None, za, zb, autojunk=False)
                if sm.real_quick_ratio() < 0.5 or sm.quick_ratio() < 0.5:
                    continue
                q = sm.ratio()
                if q > r:
                    best, r = za, q
            if r >= 0.5:
                pos = set()
                for t2, x1, x2, y1, y2 in difflib.SequenceMatcher(None, best, zb, autojunk=False).get_opcodes():
                    if t2 in ('replace', 'insert'):
                        pos.update(range(y1, y2))
                res[j] = pos
            else:
                res[j] = set(range(len(zb)))
    return res


def zitiersyntax(z, p):
    offen = z.rfind('(', 0, p)
    zu = z.find(')', p)
    if offen != -1 and zu != -1 and z.rfind(')', 0, p) < offen:
        return bool(re.match(r'\s*[A-ZÄÖÜ][^()]*?\b(19|20)\d\d\b', z[p + 1:zu]))
    return False


aus.append('PRÜFUNG (a) Semikolon an neuen oder geänderten Stellen (Vergleich mit dem Vorstand, Zitiersyntax ausgenommen)')
ta = []
for name, (alt, neu) in DOK.items():
    stellen = neue_stellen(alt, neu)
    zl = neu.split('\n')
    n_neu, n_semi = 0, 0
    for j, pos in stellen.items():
        n_neu += len(pos)
        for m in re.finditer(SEMI, zl[j]):
            if m.start() in pos:
                if zitiersyntax(zl[j], m.start()):
                    continue
                n_semi += 1
                ta.append('%s, Zeile %d: …%s…' % (name, j + 1, zl[j][max(0, m.start() - 50):m.start() + 50]))
    aus.append('   %-22s %4d geänderte Zeilen, %6d neue Zeichen, Semikola dort außerhalb der Zitiersyntax: %d, im ganzen Dokument %d (Altbestand)'
               % (name, len(stellen), n_neu, n_semi, neu.count(SEMI)))
ergebnis('(a)', ta)

# ------------------------------------------------------------------ Werte aus Fassung 17, Messskript und Textvorschlag
ub = re.search(r'\*\*Unterbudgets \(Vorgabe\):\*\*(.*)', f17).group(1)
UB = {m.group(1): zahl(m.group(2)) for m in re.finditer(r'(\d(?:\.\d){1,2}|\b7\b) (\d{1,3}(?:\.\d{3})?)\b', ub)}
KAP = {'1': 1500, '4': 2550, '5': 450, '6': 1600, '7': 250}
mess_csv = os.path.join(os.path.dirname(MSKRIPT), 'Manuskriptstand_2026-09-25.csv')
mess_txt = os.path.join(os.path.dirname(MSKRIPT), 'Manuskriptstand_2026-09-25.txt')
W = {r['arbeitsnummer']: int(r['woerter']) for r in csv.DictReader(open(mess_csv, encoding='utf-8'))}
W['4.4'] = W['4.4'] + W['4.4.1'] + W['4.4.2'] + W['4.4.3']
K4 = sum(W[k] for k in ['4.1', '4.2', '4.3', '4.4', '4.5.1', '4.5.2', '4.6', '4.7'])
ALTB = sum(W[k] for k in ['2', '2.1', '2.2', '2.3', '2.4.1', '2.4.2', '2.4.3', '2.5', '3']) + (
    [int(r['woerter']) for r in csv.DictReader(open(mess_csv, encoding='utf-8')) if r['arbeitsnummer'] == '2.4'][0])
GES = K4 + ALTB + W['1'] + sum(W[k] for k in ['5.1', '5.2', '6.1', '6.2', '6.3', '7'])
tv = lies(os.path.join(ALT, '04_Uebergaben', 'Textvorschlag_Einleitung_2026-09-28.md'))
TVW = zahl(re.findall(r'\| Wörter gesamt \| \*\*(\d{1,2}\.\d{3}|\d{3,4})\*\*', tv)[0])
aus.append('')
aus.append('Werte: Unterbudgets Fassung 17 ' + ' · '.join('%s %d' % kv for kv in UB.items()) + ' · Kapitel 4 gemessen %d · Altbestand %d · '
           'Absatztext %d · Textvorschlag Einleitung %d' % (K4, ALTB, GES, TVW))

# ------------------------------------------------------------------ (e) Budgetarithmetik
aus.append('')
aus.append('PRÜFUNG (e) Budgetarithmetik über das ganze Dokument')
te = []
v5 = DOK['Gliederung v5'][1]
b31 = v5[v5.index('### 3.1 Übersicht'):v5.index('### 3.2 ')]
zeilen = [z.split('|')[1:-1] for z in b31.split('\n') if z.startswith('| ') and re.match(r'^\| \d', z)]
B = {c[0].strip(): zahl(c[3].strip()) for c in zeilen}
IST = {c[0].strip(): c[4].strip() for c in zeilen}
for k, v in UB.items():
    if B.get(k) != v:
        te.append('Gliederung v5 § 3.1: %s mit %s statt %d (Fassung 17 § 5.2)' % (k, B.get(k), v))
for k, v in KAP.items():
    if B.get(k) != v:
        te.append('Gliederung v5 § 3.1: Kapitel %s mit %s statt %d' % (k, B.get(k), v))
if B['4'] != sum(B[k] for k in ['4.1', '4.2', '4.3', '4.4', '4.5.1', '4.5.2', '4.6', '4.7']) or B['4.5'] != B['4.5.1'] + B['4.5.2']:
    te.append('Gliederung v5 § 3.1: Kapitel 4 oder 4.5 rechnet nicht')
if B['5'] != B['5.1'] + B['5.2'] or B['6'] != B['6.1'] + B['6.2'] + B['6.3']:
    te.append('Gliederung v5 § 3.1: Kapitel 5 oder 6 rechnet nicht')
summe = re.search(r'\| \*\*Σ\*\* \| \| \*\*Kapitel 1 bis 5\*\* \| \*\*(\d\.\d{3})\*\* \| \*\*(\d\.\d{3}) mit Altbestand\*\*', b31)
if not summe or zahl(summe.group(1)) != sum(KAP.values()) or zahl(summe.group(2)) != GES:
    te.append('Gliederung v5 § 3.1: Summenzeile (Budget 6.350, Ist %d) stimmt nicht' % GES)
for k in ['4.1', '4.2', '4.3', '4.4', '4.5.1', '4.5.2', '4.6', '4.7']:
    if IST[k] != '{:,}'.format(W[k]).replace(',', '.'):
        te.append('Gliederung v5 § 3.1: Ist %s = %s statt %d (Messskript)' % (k, IST[k], W[k]))
if IST['4'] != '{:,}'.format(K4).replace(',', '.') or not IST['1'].startswith('{:,}'.format(ALTB).replace(',', '.')) or str('{:,}'.format(TVW).replace(',', '.')) not in IST['1']:
    te.append('Gliederung v5 § 3.1: Ist Kapitel 4 oder Einleitung stimmt nicht')
b34 = v5[v5.index('### 3.4 '):v5.index('### 3.5 ')]
kb = [zahl(z.split('|')[3].strip()) for z in b34.split('\n') if re.match(r'^\| \d [A-ZÄÖÜ]', z)]
if kb != [KAP[k] for k in ['1', '4', '5', '6', '7']] or '| **Σ** | **4.352** | **6.350** |' not in b34:
    te.append('Gliederung v5 § 3.4: Kapitelbudgets %s' % kb)
ub5 = re.search(r'\*\*Unterbudgets \(Vorgabe[^)]*\):\*\*(.*)', b34).group(1)
if {m.group(1): zahl(m.group(2)) for m in re.finditer(r'(\d(?:\.\d){1,2}|\b7\b) (\d{1,3}(?:\.\d{3})?)\b', ub5)} != UB:
    te.append('Gliederung v5 § 3.4: Unterbudgets weichen von Fassung 17 § 5.2 ab')
aus.append('   Gliederung v5: § 3.1 Budgets gegen Fassung 17 § 5.2 und Ist gegen das Messskript, § 3.4 Kapitelbudgets und Unterbudgets geprüft')
# Berichtsraster: Kopfblöcke
ra = DOK['Berichtsraster Rev. 3'][1]
KOPF = {'3.1': ('1', 1500), '3.4': ('4.1', UB['4.1']), '3.5': ('4.2', UB['4.2']), '3.6': ('4.3', UB['4.3']), '3.7': ('4.4', UB['4.4']),
        '3.8': ('4.5.1', UB['4.5.1']), '3.9': ('4.5.2', UB['4.5.2']), '3.10': ('4.6', UB['4.6']), '3.11': ('4.7', UB['4.7']),
        '3.12': ('5.1', UB['5.1']), '3.13': ('5.2', UB['5.2']), '3.14': ('6', KAP['6']), '3.15': ('7', UB['7'])}
for par, (ab, soll) in KOPF.items():
    i0 = ra.index('### %s ' % par)
    kb = ra[i0:ra.index('**Kopfblock.**', i0) + 60]
    m = re.search(r'\*\*Kopfblock\.\*\* Budget (\d{1,3}(?:\.\d{3})?)', kb)
    if not m or zahl(m.group(1)) != soll:
        te.append('Berichtsraster § %s (%s): Kopfblock-Budget %s statt %d' % (par, ab, m.group(1) if m else 'fehlt', soll))
k6 = ra[ra.index('### 3.14 '):ra.index('### 3.15 ')]
k6kopf = k6[k6.index('**Kopfblock.**'):k6.index('\n', k6.index('**Kopfblock.**'))]
for ab in ['6.1', '6.2', '6.3']:
    m = re.search(re.escape(ab) + r' [^·]*? (\d{3})\b', k6kopf)
    if not m or int(m.group(1)) != UB[ab]:
        te.append('Berichtsraster § 3.14: %s mit %s statt %d' % (ab, m.group(1) if m else 'fehlt', UB[ab]))
if '≤ 6.350' not in ra:
    te.append('Berichtsraster § 7 Nr. 1: „≤ 6.350“ fehlt')
if re.search(r'\*\*Kopfblock\.\*\* Budget ≈', ra):
    te.append('Berichtsraster: Kopfblock mit Näherungsbudget „≈“')
aus.append('   Berichtsraster Rev. 3: %d Kopfblöcke gegen Fassung 17 § 5.2, 6.1 bis 6.3 im Kopfblock Kapitel 6, § 7 Nr. 1' % len(KOPF))
# Plan § 5
pl = DOK['Plan Rev. 5'][1]
b5 = pl[pl.index('**Stand Rev. 5 (28.09., Messskript Fassung 3'):pl.index('Einen Kürzungsauftrag gibt es nicht mehr.')]
P5 = {z.split('|')[1].strip(): zahl(z.split('|')[2].strip()) for z in b5.split('\n') if z.startswith('| ') and not z.startswith('| Posten') and not z.startswith('|---')}
REST = KAP['5'] + KAP['6'] + KAP['7']
soll5 = {'Absatztext im Master mit Altbestand (Kapitel 2 und 3)': GES, 'davon Altbestand, wird durch die Einleitung ersetzt': ALTB,
         'Kapitel 4 (Arbeitsnummer, 4.1 bis 4.7) gegen 2.550': K4, 'Textvorschlag Einleitung gegen 1.500': TVW,
         'Budget der leeren Abschnitte 5.1 bis 7': REST, 'Summe mit vollem Einleitungsbudget': K4 + KAP['1'] + REST,
         'Summe mit dem Textvorschlag Einleitung': K4 + TVW + REST, 'Hartgrenze (Fassung 17 § 1.1)': sum(KAP.values())}
for k, v in soll5.items():
    if P5.get(k) != v:
        te.append('Plan § 5: „%s“ = %s statt %d' % (k, P5.get(k), v))
if len(P5) != len(soll5):
    te.append('Plan § 5: %d Zeilen statt %d' % (len(P5), len(soll5)))
aus.append('   Plan Rev. 5 § 5: %d Zeilen gegen Messskript, Textvorschlag und Budgets (Summen %d und %d gegen %d)'
           % (len(P5), K4 + KAP['1'] + REST, K4 + TVW + REST, sum(KAP.values())))
ergebnis('(e)', te)

# ------------------------------------------------------------------ (f) Paragrafenverweise
aus.append('')
aus.append('PRÜFUNG (f) interne Paragrafenverweise zeigen auf vorhandene Paragrafen')
FREMD = ['F14', 'F16', 'F17', 'Fassung', 'Projektanweisungen', 'Auswertungsplan', 'Umfangsdokument', 'Voraussetzungsprüfungen',
         'Voraussetzungspruefungen', 'Bauplan', 'Übergabe', 'Uebergabe', 'Textvorschlag', 'Befund', 'Vorarbeit', 'Prüfungsordnung',
         'Spezifikation', 'Kennzahlenblatt', 'Auswertungsverfahren', 'Stilprofil', 'CONSORT_Auswertung', 'Prüfprotokoll', 'Leitfaden',
         'Skill', 'Rahmenplan', 'Messskript', 'Seitenmodell', 'Nachtrag', 'Anlage', 'Analyseprotokoll', 'Fragebogenauswertung',
         'Mindestdosis', 'Lehrgang', 'Belegprotokoll', 'Plausibilität', 'Durchsicht', 'Abgleichprotokoll', 'Versuchszahl', 'Umfang',
         'Rev. 1', 'Raster vom 13.09.', 'Berichtsraster_und_Wortbudget', 'Statistische_Verfahren', 'Testprotokoll', 'Objekte_',
         'Korpus', 'Beschaffungsliste', 'Forschungsstand_Strategie', 'Studienauswahl', 'Tabellen_Spezifikation', 'RCT_', 'Datendurchsicht',
         'DSHS', 'SMK', 'dvs', 'Ethikantrag', 'Antrag', 'Gliederung v4', 'v4 ', 'CONSORT', 'Kennzahlen', 'Dokument', 'Relevanz',
         'Protokoll', 'Rechercheprotokoll', 'Evidenz', 'Auswertungsund', 'Berichtsumfang', 'Endabgleich', 'Handprobe', 'README',
         'Register']
EIGEN = {'Gliederung v5': ['Gliederung', 'v5'], 'Berichtsraster Rev. 3': ['Berichtsraster', 'Raster'],
         'Plan Rev. 5': ['Plan'], 'Übergabe Einleitung': ['Übergabe Einleitung', 'diese Übergabe'], 'README': ['README']}
tf = []
for name, (_, neu) in DOK.items():
    if name == 'README':
        continue
    vorhanden = set()
    for z in neu.split('\n'):
        m = re.match(r'^#+\s+(\d+[a-z]?(?:\.\d+[a-z]?)*)\.?\s', z)
        if m:
            vorhanden.add(m.group(1))
    fremd = FREMD + [x for k, v in EIGEN.items() if k != name for x in v]
    nint = 0
    for nr, z in enumerate(neu.split('\n'), 1):
        for m in re.finditer(r'§\s(\d+[a-z]?(?:\.\d+[a-z]?)?)', z):
            vor = z[max(0, m.start() - 60):m.start()]
            nach = z[m.end():m.end() + 8]
            # Kettenverweis „§ 5.2 und § 6“ oder „§ 3, § 4“: Dokumentname vor dem ersten Paragrafen der Kette
            kette = re.search(r'((?:§\s\d+[a-z]?(?:\.\d+[a-z]?)?(?:\s(?:und|mit|bis|oder)\s|,\s|\s·\s)?)+)$', z[:m.start()])
            vor_kette = z[max(0, (kette.start() if kette else m.start()) - 60):(kette.start() if kette else m.start())]
            if any(f in vor for f in fremd) or any(f in vor_kette for f in fremd) or nach.startswith(' dort') or nach.startswith(' der '):
                continue
            # Verweis auf einen hier nicht vorhandenen Paragrafen: fremd, wenn vorher in derselben Zeile ein fremdes Dokument genannt ist
            if m.group(1) not in vorhanden and (any(f in z[:m.start()] for f in fremd) or re.search(r'\bF1\d\b', z[:m.start()])):
                continue
            if z.startswith('| ') and any(f in z.split('|')[1] for f in fremd + ['_2026-', 'Textvorschlag_', 'Uebergabe_']):
                continue
            nint += 1
            if m.group(1) not in vorhanden:
                tf.append('%s Zeile %d: „§ %s“ ohne Paragraf (…%s…)' % (name, nr, m.group(1), z[max(0, m.start() - 60):m.end() + 10]))
    aus.append('   %-22s %3d interne Verweise geprüft, Paragrafen: %s' % (name, nint, ' '.join(sorted(vorhanden, key=lambda x: [int(y) if y.isdigit() else y for y in re.split(r'(\d+)', x) if y]))))
ergebnis('(f)', tf)

# ------------------------------------------------------------------ (i) Trefferliste des Plans
aus.append('')
aus.append('PRÜFUNG (i) Trefferliste des Plans: jeder Treffer ersetzt, historisch, erledigt oder entfallen')
MUSTER = [r'9\.000', r'Teil A', r'Teil B', r'Kapitel 1 bis 7', r'Kapitel 2\b', r'Kapitel 3\b', r'(?<![\d.§ ])(?<!§ )2\.[1-5](?![\d])',
          r'Tasks 7 bis 10', r'Task 14', r'\bR6\b', r'30,3', r'36 Seiten', r'(?<![\d,.])37(?![\d,.])', r'Wortlaut des Antrags',
          r'Antragswortlaut', r'mehr Kapitel', r'Prüfprotokoll', r'vor Kenntnis der KG-Werte', r'in 4\.7', r'F14 §', r'F16 §',
          r'Gliederung v4', r'Berichtsraster Rev\. 2', r'Klickfrage 10', r'Klickfrage 11']
# Zeilen historischer Blöcke (Revisionsvermerke, datierte Stände) und Regeln für Treffer in laufenden Teilen
HIST_ZEILE = [r'^\*\*Auftrag des Verfassers \(25\.09\.2026', r'^\*\*Rev\. [234] ', r'^\*\*Nachtrag 2[68]\.09\.2026', r'^\d\. \*\*Stand Rev\. 4\.\*\*',
              r'^[1-6]\. \*\*', r'^Gemessen am 25\.09\.2026', r'^\| \*\*Kapitel 1 bis 7\*\* \| \*\*8\.647', r'^\*\*Summen\.\*\*',
              r'^\*\*Stand nach Task 6 \(26\.09\.', r'^\| (Absatztext im Master \(Kapitel 1 bis 7\)|Budget der neun|Summe bei|Hartgrenze \(F14|\*\*Kürzungsauftrag|davon Kapitel|darin 4\.7)',
              r'^\*\*Verfasserentscheidung 25\.09\., abends', r'^\*\*Wie 4\.7 auf 550 kommt', r'^\*\*Stand Rev\. 4:\*\*',
              r'^\*\*Kürzungen je Abschnitt', r'^\| (A2, A3|J3|G16d|G27e)', r'^Gemessen: alle Wort-', r'^\*\*Änderungen in Rev\. 2 gegenüber',
              r'^\*Erstellt am 25\.09\.2026', r'^8\. \*\*Absatz 7:\*\*', r'^\*\*Offen aus Block B', r'^\| \*\*Kapitel 1 bis 7\*\* \| \*\*8\.647\*\* \| \*\*7\.603']
REGELN = [
    (r'Task 14 entfällt|Task 14 · entfällt|Task 14:\*\* entfällt|Einleitungsteil von Task 14|Kürzung von Kapitel 2 und Task 14', 'als entfallen gekennzeichnet'),
    (r'R6 entfällt|R6 nach K2', 'neue Regel: R6 entfällt (K2)'),
    (r'Verweise „F14 §“|„F14 §“ in den laufenden Teilen', 'Regel zu den F14-Verweisen (Rev.-5-Vermerk, § 9)'),
    (r'Klickfrage 10 (ist dort entschieden|erledigt|dort entschieden)|Klickfrage 10 \(Wortlaut des Ankersatzes\) · Stufenwahl', 'als erledigt gekennzeichnet (Task 7 neu, Rev. 114)'),
    (r'Altbestand von Kapitel 2 und 3|Altbestand \(Kapitel 2 und 3\)|Altbestand 4\.920 \(Kapitel 2 und 3\)|alten Kapitel 2 und 3', 'Altbestand, entfällt mit der Einleitung'),
    (r'ersetzt die Kürzung von Kapitel 2|Kapitel 2 wird nicht mehr gekürzt', 'als ersetzt gekennzeichnet'),
    (r'Tasks 7 bis 10 der Rev\. 4', 'als ersetzt gekennzeichnet'),
    (r'ohne Prüfprotokolle|Ohne Prüfprotokolle|Prüfprotokoll der Gliederung', 'neue Regel oder Dateibezug'),
    (r'(nicht mehr|nicht) in 4\.7|in 4\.7 bleibt|in 4\.7 nicht berichtet|in 4\.7 \(4\.7 nennt|Produktname in 4\.7', 'geprüft gegen 4.7 im Master'),
    (r'\*Stand 25\.09\., historisch:\*', 'als historisch gekennzeichnet'),
    (r'4\.7 ist der dichteste Abschnitt der Arbeit \(Gliederung v4 § 3\.2\)', 'Zitat aus Gliederung v4, § 2 ist historisch'),
    (r'F14 § 1\.2\)|F14 § 1\.2 \(verbindliche', 'Verfahrensregel, in Fassung 17 § 1.2 unverändert, historischer Teil (§ 2) oder allgemeine Regel'),
    (r'F14 § 11\.7 und § 11\.9', '§ 2 ist historisch (erledigt 26.09.)'),
]
SEK_HIST = ['## 2 Schritt 0', '### 2.1 ', '### 2.2 ', '### 2.3 ', '## 6 Maßnahmenliste', '## 9 Prüfung dieses Dokuments']
ti = []
zaehler = {}
einordnung = {}
sek = 'Kopf'
for nr, z in enumerate(pl.split('\n'), 1):
    if z.startswith('#'):
        sek = z
    for mu in MUSTER:
        for m in re.finditer(mu, z):
            umg = z[max(0, m.start() - 90):m.end() + 90]
            zaehler[mu] = zaehler.get(mu, 0) + 1
            grund = [g for rx, g in REGELN if re.search(rx, umg)]
            if not grund and any(re.match(rx, z) for rx in HIST_ZEILE):
                grund = ['historischer Block (datierter Stand oder Revisionsvermerk)']
            if not grund and any(sek.startswith(s) for s in SEK_HIST):
                grund = ['historischer Paragraf (%s)' % sek.strip('# ')[:30]]
            if grund:
                einordnung[grund[0]] = einordnung.get(grund[0], 0) + 1
            else:
                ti.append('Z%d %s ohne Einordnung: …%s…' % (nr, mu, umg))
aus.append('   Treffer je Muster: ' + ' · '.join('%s %d' % (k, v) for k, v in zaehler.items()))
aus.append('   Ohne Treffer: ' + ' · '.join(mu for mu in MUSTER if mu not in zaehler))
aus.append('   Einordnung: ' + ' · '.join('%s %d' % (k, v) for k, v in sorted(einordnung.items(), key=lambda x: -x[1])))
ergebnis('(i)', ti, ' (%d Treffer, alle eingeordnet)' % sum(zaehler.values()) if not ti else '')

# ------------------------------------------------------------------ (g) Probe Messskript nach der Übertragung
aus.append('')
aus.append('PRÜFUNG (g) Probe: Messskript Fassung 3 misst die Einleitung nach der Übertragung ohne Anpassung')
tg = []
from docx import Document
W_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
d = Document(MASTER)
ps = d.paragraphs


def ueberschrift(p):
    return p.style.name.lower().startswith(('heading', 'überschrift'))


i_e = [i for i, p in enumerate(ps) if ueberschrift(p) and p.text.strip().endswith('Einleitung')]
i_4 = [i for i, p in enumerate(ps) if ueberschrift(p) and p.text.strip() in ('Methodik', '4 Methodik')]
if len(i_e) != 1 or len(i_4) != 1:
    tg.append('Überschriften „Einleitung“ oder „Methodik“ im Master nicht eindeutig')
else:
    e, b = i_e[0], i_4[0]
    for p in ps[e + 1:b]:
        p._element.getparent().remove(p._element)
    blk = tv.split('## 1 Wortlaut')[1].split('\n## 2 ')[0]
    absaetze = [z.strip() for z in blk.split('\n')[1:] if z.strip() and not z.startswith('#')]
    vorlage = [p for p in Document(MASTER).paragraphs if p.style.name in ('Standard', 'Normal') and len(p.text) > 200][0]
    anker = d.paragraphs[e]._element
    for text in reversed(absaetze):
        neu = copy.deepcopy(vorlage._element)
        for r in neu.findall(W_NS + 'r')[1:]:
            neu.remove(r)
        ts = neu.findall('.//' + W_NS + 't')
        ts[0].text = text
        for x in ts[1:]:
            x.getparent().remove(x)
        anker.addnext(neu)
    tmp = tempfile.mkdtemp()
    probe = os.path.join(tmp, 'Probe_Master.docx')
    d.save(probe)
    subprocess.run([sys.executable, MSKRIPT, probe, os.path.join(tmp, 'p.txt'), os.path.join(tmp, 'p.csv'), S_CSV], check=True, stdout=subprocess.DEVNULL)
    ptxt = lies(os.path.join(tmp, 'p.txt'))
    pw = {r['arbeitsnummer']: r for r in csv.DictReader(open(os.path.join(tmp, 'p.csv'), encoding='utf-8'))}
    e_w = int(pw['1']['woerter'])
    semi = int(re.search(r'Semikola außerhalb von Zitierklammern, Kapitel 1 bis 7: (\d+)', ptxt).group(1))
    verw = int(re.search(r'Nummerierte Abschnittsverweise, Kapitel 1 bis 7: (\d+)', ptxt).group(1))
    progn = re.search(r'Prognose Einleitung bis Ende Literaturverzeichnis: (\d+,\d) Seiten', ptxt).group(1)
    alt_zeilen = [k for k in pw if re.match(r'^(2|3)(\.|$)', k)]
    aus.append('   Probe-Master: %d Absätze unter „1 Einleitung“, Überschriften 2 bis 3 entfernt. Messskript: Einleitung %d Wörter '
               '(Textvorschlag %d), Semikola %d, Abschnittsverweise %d, Zeilen 2.x und 3 in der Ausgabe: %d, Seitenprognose %s Seiten'
               % (len(absaetze), e_w, TVW, semi, verw, len(alt_zeilen), progn))
    if len(absaetze) != 9 or e_w != TVW or semi or verw or alt_zeilen or 'Altbestand' in ptxt.split('Budgetabschnitte')[1].split('Semikola')[0].replace('mit dem Altbestand 2, 2.x und 3', ''):
        tg.append('Probe weicht ab (Absätze %d, Wörter %d gegen %d, Semikola %d, Verweise %d, Altbestandszeilen %d)' % (len(absaetze), e_w, TVW, semi, verw, len(alt_zeilen)))
ergebnis('(g)', tg)

# ------------------------------------------------------------------ (h) Aufräumskript, Archiv, README
aus.append('')
aus.append('PRÜFUNG (h) Aufräumskript, Archivkopien und README')
th = []
roh = open(os.path.join(NEU, 'Ordner_aufraeumen.ps1'), 'rb').read()
if not roh.startswith(b'\xef\xbb\xbf') or roh.count(b'\n') != roh.count(b'\r\n'):
    th.append('Aufräumskript: BOM oder CRLF verloren')
ps1 = roh[3:].decode('utf-8')
ps1_alt = open(os.path.join(ALT, 'Ordner_aufraeumen.ps1'), 'rb').read()[3:].decode('utf-8')
neue_eintraege = [x for x in re.findall(r"^\s*'(Claude\\[^']+)',?\r?$", ps1, re.M) if x not in ps1_alt]
arch = {}
for wurzel, _, dateien in os.walk(os.path.join(NEU, '_Archiv')):
    for f in dateien:
        arch.setdefault(f, []).append(os.path.join(wurzel, f))
for e in neue_eintraege:
    rel = e.replace('Claude\\', '').replace('\\', '/')
    orig = os.path.join(ALT, rel)
    name = os.path.basename(rel)
    if name not in arch:
        th.append('Papierkorbeintrag ohne Archivkopie: %s' % e)
        continue
    if not os.path.exists(orig):
        th.append('Original nicht gestagt: %s' % rel)
        continue
    md5o = hashlib.md5(open(orig, 'rb').read()).hexdigest()
    if not any(hashlib.md5(open(a, 'rb').read()).hexdigest() == md5o for a in arch[name]):
        th.append('Archivkopie weicht vom Original ab: %s' % name)
if "'Projektanweisungen_Fassung17.md')" not in ps1:
    th.append('Aufräumskript: Muster 00_Steuerung nicht auf Fassung 17')
f16_in_liste = "'Claude\\00_Steuerung\\Projektanweisungen_Fassung16.md'" in ps1
if f16_in_liste and 'Projektanweisungen_Fassung16.md' not in arch:
    th.append('Fassung 16 im Papierkorb ohne Archivkopie')
rd = DOK['README'][1]
for muss in ['Fassung 17 vom 28.09.2026', '`Projektanweisungen_Fassung17.md`', '(Rev. 5)', '(Rev. 3, Kopfblock Einleitung', '`Gliederung_2026-09-28`',
             '`Textvorschlag_Einleitung_2026-09-28.md`', '_ersetzt_2026-09-28_Steuerdokumente', '_ersetzt_2026-09-28_Uebergaben']:
    if muss not in rd:
        th.append('README ohne „%s“' % muss)
for darf_nicht in ['Fassung 16 vom', '(Rev. 4', '`Gliederung_2026-09-23`', 'Textvorschlag_4.7_2026-09-25.md` (Eingang Task 6)']:
    if darf_nicht in rd:
        th.append('README trägt noch „%s“' % darf_nicht)
soll_arch = ['Gliederung_2026-09-23.md', 'Gliederung_2026-09-23.docx', 'Gliederung_2026-09-23.pdf', 'Berichtsraster_2026-09-23.md',
             'Berichtsraster_2026-09-23.docx', 'Berichtsraster_2026-09-23.pdf', 'Plan_Weitere_Schritte_2026-09-25.md', 'Manuskriptstand_2026-09-25.py',
             'Uebergabe_Kapitel2_2026-09-26.md', 'Uebergabe_2.4_Fassung6_2026-09-28.md', 'Textvorschlag_2.4_2026-09-26.md', 'Textvorschlag_2.4_2026-09-27.md',
             'Textvorschlag_2.4_2026-09-27_Fassung5.md', 'Uebergabe_Steuerdokumente_2026-09-25.md', 'Textvorschlag_4.7_2026-09-25.md',
             'Uebergabe_Textrevision_2026-09-12.md', 'Uebergabe_Steuerdokumente_2026-09-28.md',
             'README_Ordnerstruktur.md', 'Ordner_aufraeumen.ps1', 'Uebergabe_Einleitung_2026-09-28.md']
for n in soll_arch:
    if n not in arch:
        th.append('Archivkopie fehlt: %s (Übergabe § 5.6)' % n)
vorstand = {'Berichtsraster_2026-09-23.md': 'Rev. 2 vom 24.09.2026', 'Plan_Weitere_Schritte_2026-09-25.md': 'Rev. 4',
            'Manuskriptstand_2026-09-25.py': 'Fassung 2 (25.09., spät'}
for n, merkmal in vorstand.items():
    kopf = open(arch[n][0], encoding='utf-8').read()[:2500]
    if merkmal not in kopf or 'Rev. 5' in kopf.split('\n')[0] or 'Rev. 3' in kopf.split('\n')[2]:
        th.append('Archivkopie %s ist nicht der Vorstand' % n)
for n, rel in [('README_Ordnerstruktur.md', 'README_Ordnerstruktur.md'), ('Ordner_aufraeumen.ps1', 'Ordner_aufraeumen.ps1'),
               ('Uebergabe_Einleitung_2026-09-28.md', '04_Uebergaben/Uebergabe_Einleitung_2026-09-28.md')]:
    md5o = hashlib.md5(open(os.path.join(ALT, rel), 'rb').read()).hexdigest()
    if n in arch and not any(hashlib.md5(open(a, 'rb').read()).hexdigest() == md5o for a in arch[n]):
        th.append('Vorstand %s im Archiv weicht vom gestagten Original ab' % n)
aus.append('   Aufräumskript: %d neue Papierkorbeinträge, Fassung 16 in der Liste: %s · Archivkopien: %d Dateien in %s'
           % (len(neue_eintraege), 'ja' if f16_in_liste else 'nein (erst nach K5)', sum(len(v) for v in arch.values()),
              ', '.join(sorted(os.listdir(os.path.join(NEU, '_Archiv'))))))
ergebnis('(h)', th)

# ------------------------------------------------------------------ (k) Übergabe Einleitung
aus.append('')
aus.append('PRÜFUNG (k) Übergabe Einleitung: Verweise und Seitenzahlen')
tk = []
ue = DOK['Übergabe Einleitung'][1]
for m in re.finditer(r'F16 §', ue):
    z = ue[:m.start()].split('\n')[-1] + ue[m.start():].split('\n')[0]
    if not (re.search(r'bisher nach F16 § 1\.1', z) or re.search(r'Teil B in F16 § 1\.1', z)):
        tk.append('„F16 §“ außerhalb der historischen Stellen in § 2: …%s…' % ue[max(0, m.start() - 60):m.end() + 30])
for alt in ['rund 21', 'rund 25 bis 27', '32 bis 36', 'R6 holt', 'Ob der Ankersatz auch in 4.1', 'Vorher läuft der Task Steuerdokumente',
            'im Antragswortlaut, auf die', 'Berichtsraster § 3.1 bis § 3.3', '` § 3.1 bis § 3.3']:
    if alt in ue:
        tk.append('Übergabe Einleitung trägt noch „%s“' % alt)
aus.append('   „F16 §“: %d Stellen, alle historisch in § 2' % ue.count('F16 §') if not tk else '')
ergebnis('(k)', tk)

# ------------------------------------------------------------------ (l) Klicks 19:25, Objekte, Zielorte
aus.append('')
aus.append('PRÜFUNG (l) Erratum, Tab. H6, Verdünnungslogik, Objekte bis Task 18, „vorgemerkt für“')
tl = []
alle = dict((k, v[1]) for k, v in DOK.items())
alle['Fassung 17'] = f17
for name, text in alle.items():
    for nr, z in enumerate(text.split('\n'), 1):
        if 'Erratum' in z and any(x in z for x in ['Anhang G', 'Tab. H6', 'Fließtext', 'R14']) and not any(
                x in z for x in ['nicht', 'Kein Hinweis', 'kein Hinweis', 'ohne Hinweis', 'kein Erratum', 'ohne R1']):
            tl.append('%s Zeile %d: Erratum mit Ort in der Arbeit: …%s…' % (name, nr, z[max(0, z.index('Erratum') - 80):z.index('Erratum') + 80]))
pl5 = DOK['Plan Rev. 5'][1]
v5t = DOK['Gliederung v5'][1]
muss = [('Plan Rev. 5', pl5, 'ohne R1, R9, R10, R12, R13 und R14'),
        ('Plan Rev. 5', pl5, 'Verdünnungslogik als Einordnung des Hauptbefunds (als Limitation in 6.3 G3'),
        ('Plan Rev. 5', pl5, '*Objekte:* keine Objekte und keine Platzhalter im Master.'),
        ('Plan Rev. 5', pl5, '(b) Objekte setzen: im Textteil Tab. 1 mit Beschriftung und Anmerkung aus Textvorschlag 4.4 § 9'),
        ('Plan Rev. 5', pl5, 'die Freigabe ohne unabhängige Methodenprüfung ist für 6.3 vorgemerkt'),
        ('Plan Rev. 5', pl5, 'der Name ist für die KI-Deklaration und Anhang G vorgemerkt'),
        ('Plan Rev. 5', pl5, 'Einbau per Skript (nur Text, keine Objekte und keine Platzhalter)'),
        ('Gliederung v5', v5t, 'Verdünnungslogik als Einordnung des Hauptbefunds'),
        ('Gliederung v5', v5t, 'ohne eigenen Absatz zur Verdünnung'),
        ('Gliederung v5', v5t, 'auch nicht als Platzhalter'),
        ('Berichtsraster Rev. 3', ra, 'Ort 6.1, als Limitation 6.3 G3, 6.2 ohne eigenen Absatz'),
        ('Fassung 17', f17, '**Verdünnungslogik für 6.1:**'),
        ('Fassung 17', f17, 'bis dahin im Text keine Platzhalter, Platzhalter nur für die Anhangsobjekte in Anhang H'),
        ('Fassung 17', f17, 'R14 steht wie R1, R9, R10, R12 und R13 nicht in Tab. H6')]
for name, text, x in muss:
    if x not in text:
        tl.append('%s ohne „%s“' % (name, x))
for x in ['Word-Tabellen per Skript aus den CSV', 'Abb. 1 und Abb. 2 als Platzhalter mit Beschriftungsfeld', 'Einbau per Skript mit Tab. 2 und Tab. 3',
          'Umsetzung der 48-h-Vorgabe je Testtermin', 'steht genau einmal in 6.3', 'für Kapitel 4 und 2']:
    if x in pl5:
        tl.append('Plan Rev. 5 trägt noch „%s“' % x)
r62 = ra[ra.index('**6.2 Methodendiskussion**'):ra.index('**6.3 Stärken und Limitationen**')]
if 'Verdünn' in r62:
    tl.append('Berichtsraster: Verdünnung in einer 6.2-Zeile')
aus.append('   Erratum-Zeilen in F17, Gliederung v5, Raster, Plan, Übergabe Einleitung und README geprüft · %d Pflichtstellen · Negativliste Plan · 6.2-Zeilen des Rasters' % len(muss))
ergebnis('(l)', tl)

aus.append('')
aus.append('GESAMT: ' + ('alle Prüfungen bestanden' if not fehler else 'NICHT BESTANDEN: ' + ', '.join(fehler)))
kopf = ['Steuerdokumente_Pruefung_2026-09-28.py — Prüfung der Folgeänderungen (Task Steuerdokumente 28.09., Schritt 4)',
        'Vorstand: %s · Neu: %s' % (ALT, NEU), '']
open(OUT, 'w', encoding='utf-8', newline='\n').write('\n'.join(kopf + aus) + '\n')
print('\n'.join(aus))
if fehler:
    raise SystemExit(1)
