# -*- coding: utf-8 -*-
"""
Datensicherung_Belegprotokoll_2026-09-24.py (Hilfsskript, erzeugt das Protokoll, rechnet nichts)

Zweck:    Setzt das Belegprotokoll der Phase 1 (Hausstil, .docx und .pdf) aus den Ausgaben der
          Skripte und der ausgefüllten Arbeitsliste zusammen: Prüfregeln mit Datum, markierte
          Werte mit Urteil, Code-Zuordnung, Zuordnungstabelle Fragebogen, Korrekturen, Quell-
          dateien mit SHA-256, Ergebnis der Rückleseprobe und der Abbruchprobe.
Schritt:  Auswertungsverfahren_2026-09-24, Phase 1 (Nachweis „Belegprotokoll“ zu 1.1 und 1.4)
Eingang:  03_Skripte/Datensicherung_Struktur_2026-09-24.txt (Lauf vor den Korrekturen, als
          _vorher.txt gesichert) und der Lauf danach · Datensicherung_Plausibilitaet_2026-09-24.txt ·
          Datensicherung_Korrektur_2026-09-24.txt und _Liste.csv · Datensicherung_Export_2026-09-24.txt ·
          04_Uebergaben/Belegpruefung_2026-09-24.xlsx (ausgefüllt) · Statistik/Datenstand_JJJJ-MM-TT/Pruefsummen.txt
Fassung:  2026-09-24, Fassung 1
Aufruf:   python Datensicherung_Belegprotokoll_2026-09-24.py [Basisordner] --datum JJJJ-MM-TT
Regeln:   Keine Namen, kein Geburtsdatum, kein Testdatum je Spieler. Keine Semikolons.
"""
import sys, os, re, csv, hashlib, datetime, subprocess
import openpyxl
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

args = [a for a in sys.argv[1:] if not a.startswith("--")]
BASIS = os.path.abspath(args[0]) if args else os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
DATUM = sys.argv[sys.argv.index("--datum") + 1] if "--datum" in sys.argv else datetime.date.today().strftime("%Y-%m-%d")
SK = os.path.join(BASIS, "Claude", "03_Skripte")
P_STRUKT_VOR = os.path.join(SK, "Datensicherung_Struktur_2026-09-24_vorher.txt")
P_STRUKT_NACH = os.path.join(SK, "Datensicherung_Struktur_2026-09-24.txt")
P_PLAUS = os.path.join(SK, "Datensicherung_Plausibilitaet_2026-09-24.txt")
P_KORR = os.path.join(SK, "Datensicherung_Korrektur_2026-09-24.txt")
P_LISTE = os.path.join(SK, "Datensicherung_Korrektur_2026-09-24_Liste.csv")
P_EXPORT = os.path.join(SK, "Datensicherung_Export_2026-09-24.txt")
P_BELEG = os.path.join(BASIS, "Claude", "04_Uebergaben", "Belegpruefung_2026-09-24.xlsx")
P_DS = os.path.join(BASIS, "Statistik", "Datenstand_" + DATUM)
ZIEL = os.path.join(BASIS, "Claude", "02_Befunde", "Belegprotokoll_" + DATUM)
NAVY = RGBColor(0x1F, 0x38, 0x64)

def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""): h.update(chunk)
    return h.hexdigest()
def lese(p):
    with open(p, encoding="utf-8") as f: return f.read()

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(2)
def ueber(text, lvl=1):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = True; r.font.color.rgb = NAVY
    r.font.size = Pt(14 if lvl == 1 else 11.5); p.paragraph_format.space_before = Pt(12 if lvl == 1 else 8); p.paragraph_format.space_after = Pt(4)
    return p
def absatz(text, kursiv=False, size=10):
    p = doc.add_paragraph(); r = p.add_run(text); r.italic = kursiv; r.font.size = Pt(size); p.paragraph_format.space_after = Pt(4); return p
def schattiere(zelle, farbe):
    tcPr = zelle._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), farbe); tcPr.append(shd)
def tabelle(kopf, zeilen, breiten=None, size=8):
    t = doc.add_table(rows=1, cols=len(kopf)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(kopf):
        c = t.rows[0].cells[j]; c.text = ""; r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(size); schattiere(c, "D6E4F0")
    for i, z in enumerate(zeilen):
        row = t.add_row().cells
        for j, v in enumerate(z):
            row[j].text = ""; r = row[j].paragraphs[0].add_run("" if v is None else str(v)); r.font.size = Pt(size)
            if i % 2 == 1: schattiere(row[j], "EEF3F9")
    if breiten:
        for row in t.rows:
            for j, b in enumerate(breiten): row.cells[j].width = Cm(b)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

# ------------------------------------------------------------------ Kopf
p = doc.add_paragraph(); r = p.add_run("Belegprotokoll Phase 1 · Datensicherung"); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = NAVY
absatz("Bachelorarbeit U15-Plyometrie · DSHS Köln · Verfasser: Luca Klier · Stand %s · Auswertungsverfahren_2026-09-24, Phase 1 (Schritte 1.1 bis 1.5) · Maßnahmen L5 und B3" % DATUM, size=9)
absatz("Gehört zur Übergabe Datensicherung vom 24.09.2026 (Rev. 87). Gegenstand: Struktur- und Plausibilitätsprüfung von Workbook und Fragebogen-Rohexport, Belegprüfung markierter Werte am Papierbogen durch den Verfasser, Code-Zuordnung, belegte Korrekturen mit Eintrag und das Einfrieren des Datenstands für die blinde R-Rechnung. Datenschutz nach Weg C (Verfasser, 24.09.2026): pseudonymisierte Studiendaten in Claude, Datensparsamkeit. Dieses Protokoll enthält keine Namen, keine Geburtsdaten und keine Testdaten je Spieler. Es bleibt im Projekt und geht nicht in den Blindordner.")

# ------------------------------------------------------------------ 1 Quelldateien
ueber("1 Quelldateien und Prüfsummen")
strukt_vor = lese(P_STRUKT_VOR) if os.path.exists(P_STRUKT_VOR) else ""
strukt_nach = lese(P_STRUKT_NACH)
korr_txt = lese(P_KORR) if os.path.exists(P_KORR) else ""
export_txt = lese(P_EXPORT) if os.path.exists(P_EXPORT) else ""
zeilen = []
def pruefsummen_aus(text, kopf):
    erg = []
    for z in text.splitlines():
        m = re.match(r"\s+(\S.*?\S)\s+([0-9a-f]{64})\s+(\d+) B", z)
        if m: erg.append([kopf, m.group(1), m.group(2), m.group(3)])
    return erg
zeilen += pruefsummen_aus(strukt_vor, "vor den Korrekturen (1.3, erster Lauf)")
zeilen += pruefsummen_aus(strukt_nach, "nach den Korrekturen (1.3, Wiederholung)")
m = re.search(r"Sicherung: (\S+)\s+SHA-256 ([0-9a-f]{64})", korr_txt)
if m: zeilen.append(["Sicherung vor Phase 1", os.path.relpath(m.group(1), BASIS) if m.group(1).startswith(BASIS) else m.group(1), m.group(2), ""])
tabelle(["Stand", "Datei", "SHA-256", "Bytes"], zeilen, (3.2, 4.6, 8.2, 1.3), size=6.5)
absatz("Die Originalexporte des Fragebogens blieben unverändert. Das Workbook wurde nur durch die freigegebenen Korrekturen (Abschnitt 6) verändert, die Sicherung des Standes vor Phase 1 liegt in Claude\\_Archiv.", size=9)

# ------------------------------------------------------------------ 2 Strukturprüfung
ueber("2 Strukturprüfung (1.3)")
absatz("Skript Datensicherung_Struktur_2026-09-24.py, Regeln aus Spezifikation S01 (Regeln 2 bis 4), S03 mit K5, S06 Regeln 3 und 7, dazu Pflichtfelder, Text in Zahlenspalten, Codemengen, Vokabular und Fremdzeilen. Erster Lauf am Workbook-Stand vom 15.09.2026, 22:21 MESZ, Wiederholung nach den Korrekturen.")
def block(text, start, ende_marker):
    a = text.find(start); b = text.find(ende_marker, a + 1) if ende_marker else -1
    return text[a:b] if a >= 0 else ""
erg_vor = block(strukt_vor, "Verstöße, an denen", "Hinweise und Kontrollfälle")
hin_vor = block(strukt_vor, "Hinweise und Kontrollfälle", "Kontrollfälle laut")
kf_vor = block(strukt_vor, "Kontrollfälle laut Übergabe", "Abbruchprobe bestanden")
erg_nach = block(strukt_nach, "Verstöße, an denen", "Hinweise und Kontrollfälle")
hin_nach = block(strukt_nach, "Hinweise und Kontrollfälle", "Kontrollfälle laut")
ueber("2.1 Erster Lauf (vor den Korrekturen)", 2)
status_codes = re.findall(r"VERSTOSS \[S01.3\] (\S+): Status 'Post-Testung ausstehend", erg_vor)
absatz(erg_vor.splitlines()[0].strip(), size=8.5)
if status_codes:
    absatz("VERSTOSS [S01.3] Status „Post-Testung ausstehend (Termin 15.09.2026)“ bei %d Spielern (%s): weder „ausgewertet“ noch „Post-Testung nicht angetreten“" % (len(status_codes), ", ".join(status_codes)), size=8.5)
for l in erg_vor.splitlines()[1:]:
    if "VERSTOSS" in l and "K5] Zeile" not in l and "Status 'Post-Testung ausstehend" not in l:
        absatz(l.strip(), size=8.5)
n_k5 = sum(1 for l in erg_vor.splitlines() if "K5] Zeile" in l)
if n_k5: absatz("dazu %d Einzelverstöße K5 (ungültige Zeilen ohne Bemerkung), Liste in der Skriptausgabe." % n_k5, size=8.5)
for z in hin_vor.splitlines(): absatz(z.strip(), size=8.5)
for z in kf_vor.splitlines(): absatz(z.strip(), size=8.5)
ueber("2.2 Wiederholung nach den Korrekturen", 2)
for z in erg_nach.splitlines(): absatz(z.strip(), size=8.5)
for z in hin_nach.splitlines(): absatz(z.strip(), size=8.5)
m = re.search(r"Abbruchprobe bestanden: (\w+)", strukt_nach)
absatz("Abbruchprobe der Strukturprüfung nach den Korrekturen bestanden: " + (m.group(1) if m else "?"))

# ------------------------------------------------------------------ 3 Prüfregeln
ueber("3 Prüfregeln der Plausibilitätsprüfung (1.1) mit Datum")
plaus = lese(P_PLAUS)
absatz("Skript Datensicherung_Plausibilitaet_2026-09-24.py, Fassung 2. Regeln und Schwellen am 24.09.2026 vorgeschlagen (zwei Varianten A und B), Variante B vom Verfasser am 24.09.2026 per Klick freigegeben, vor dem Lauf. Die Regeln gelten für beide Gruppen und beide Zeitpunkte gleich. Der Verfasser prüfte nur die markierten Werte am Papierbogen, keine Zweiterfassung (Verfasser 24.09.2026). Ein Wert, der dem Papierbogen entspricht, bleibt, auch wenn er extrem ist. Bereits entschieden und nicht erneut vorgelegt: HL-08 post 10 m Versuch 2 (Verfasser 12.09.2026), 10-m-Post-Werte von Verein B „System falsch aufgenommen“.")
regeln = [
    ["P1", "Grobbereich je Test", "5 m 0,70–1,80 s · 10 m 1,40–3,00 s · 30 m 3,50–7,00 s · 505 1,80–4,00 s · SBJ 100–300 cm"],
    ["P2", "Teilzeitenfolge je Lauf", "5 m < 10 m < 30 m, sonst ganzer Lauf markiert"],
    ["P3", "Abschnitte je Lauf", "Δt ≤ 0 s oder v > 10,0 m/s (Spezifikation S02, Lauf auslösegestört) sowie Abschnittsgeschwindigkeit nicht steigend über 0–5, 5–10 und 10–30 m, ganzer Lauf markiert"],
    ["P4", "Spannweite der Versuche eines Spielers", "(Max − Min)/Min > 15 % bei Zeiten, > 20 % beim Standweitsprung, alle Versuche des Satzes markiert"],
    ["P5", "Ausreißer über alle Spieler", "|x − Median| > 3 · 1,4826 · MAD je Test, Seite und Zeitpunkt über alle gültigen Versuche"],
    ["P6", "Prä-Post-Änderung des Bestwerts", "|Post − Prä|/Prä > 10 % bei Zeiten, > 15 % beim Standweitsprung, beide Bestwerte markiert"],
    ["P7", "Alter und Anthropometrie", "Alter prä 12,0–16,0 Jahre · Körperhöhe 140–200 cm · Körpermasse 35–95 kg · BMI 14–30 kg/m² · Mutter 145–190 cm · Vater 155–205 cm · leere oder textliche Felder"],
    ["P8", "Muster", "Zeiten mit mehr als zwei Nachkommastellen, Standweitsprung mit Nachkommastellen, drei identische Versuche, identische Versuchssätze zweier Spieler, Wert trotz Bemerkung „System nicht/falsch aufgenommen“"],
]
tabelle(["Regel", "Gegenstand", "Schwelle (freigegeben 24.09.2026)"], regeln, (1.2, 4, 12))
m = re.search(r"Markierte Werte: (\d+) Markierungen\nje Regel: (.*)\nverschiedene Zellen: (\d+), Spieler betroffen: (\d+)", plaus)
if m: absatz("Ergebnis des Laufs: %s Markierungen (%s) auf %s Zellen bei %s Spielern." % (m.group(1), m.group(2), m.group(3), m.group(4)))

# ------------------------------------------------------------------ 4 Markierte Werte mit Urteil
ueber("4 Markierte Werte mit Urteil des Verfassers")
xb = openpyxl.load_workbook(P_BELEG, data_only=True)
ws = xb["Markierte_Werte"]
kopf = [c.value for c in ws[1]]
mw = []
urteile = {}
for r in range(2, ws.max_row + 1):
    z = [ws.cell(r, c).value for c in range(1, len(kopf) + 1)]
    if z[0] is None: continue
    mw.append([z[0], z[1], z[2], z[3], z[4], z[5], z[6], z[7], z[8], (z[10] or ""), (z[11] if z[11] is not None else ""), (z[12] or "")])
    urteile[str(z[10] or "").strip().lower()] = urteile.get(str(z[10] or "").strip().lower(), 0) + 1
absatz("Urteile: " + ", ".join("%s %d" % (k if k else "(leer)", n) for k, n in sorted(urteile.items())), size=9)
tabelle(["Nr", "Regel", "Code", "Zeitpunkt", "Test", "Seite", "V", "Wert", "Zelle", "Urteil", "Wert laut Bogen", "Bemerkung"], mw, (0.7, 1.1, 1.2, 1.2, 2.2, 0.8, 0.6, 1.2, 2.4, 1.8, 1.6, 3), size=7)

# ------------------------------------------------------------------ 5 Code-Zuordnung
ueber("5 Code-Zuordnung Papier → Workbook (1.2 a) und Zuordnungstabelle Fragebogen (1.2 b)")
wc = xb["Codezuordnung"]
cz = []; oz = []; zf = []
modus = "cz"
for r in range(5, wc.max_row + 1):
    a = wc.cell(r, 1).value
    if a is None: continue
    if isinstance(a, str) and a.startswith("Vergebene Codes"): modus = "oz"; continue
    if isinstance(a, str) and a.startswith("Zuordnungstabelle"): modus = "zf"; continue
    if isinstance(a, str) and a in ("Code", "Listenlabel (H010)"): continue
    z = [wc.cell(r, c).value for c in range(1, 12)]
    if modus == "cz": cz.append([z[0], z[1], z[2], z[3] or "", z[4] or "", z[5] or "", z[6] or "", z[7] if z[7] is not None else "", z[9] or "", z[10] or ""])
    elif modus == "oz": oz.append([z[0], z[1], z[2] or "", z[3] or "", z[4] or "", z[5] or ""])
    else: zf.append([z[0], z[1], z[2], z[3], z[4]])
absatz("Ausgefüllt vom Verfasser aus Papierbögen, Einwilligungserklärungen und Anwesenheitslisten, ohne Namen. Die Namenszuordnung bleibt lokal beim Verfasser.", size=9)
tabelle(["Code", "Verein", "Gruppe", "Prä-Bogen", "Post-Bogen", "Einwilligung", "Familiarisierung (Liste)", "Familiarisierung (Workbook)", "Pflicht", "Bemerkung"], cz, (1.3, 1, 1, 1.5, 1.5, 1.6, 1.8, 1.8, 2.2, 3.5), size=7)
absatz("Vergebene Codes ohne Studiendaten (Abb. 1 Teilnehmerfluss, nicht im Datenstand):", size=9)
tabelle(["Code", "Verein", "Grund (Verfasser)", "Einwilligung", "Eingangstestung", "Bemerkung"], oz, (1.3, 1, 6, 2, 2, 4), size=7)
absatz("Zuordnungstabelle Fragebogen, vom Verfasser am 24.09.2026 per Klick bestätigt. Sie steht als Zuordnung_Fragebogen.csv und Listenplatz_IG.csv im Datenstand.", size=9)
tabelle(["Listenlabel", "H010", "Analysecode", "kein_Listenplatz", "Herkunft der Regel"], zf, (2, 1, 2, 2, 10), size=7)

# ------------------------------------------------------------------ 6 Korrekturen
ueber("6 Korrekturen mit Eintrag (1.4)")
absatz("Nur belegte Korrekturen, vom Verfasser freigegeben, geschrieben von Datensicherung_Korrektur_2026-09-24.py bei geschlossenem Excel nach einer Sicherung. Jede Änderung steht im neuen Blatt 00_Aenderungen des Workbooks. Werte ohne Papierbeleg tragen den Beleg „Verfasser“. Nach dem Schreiben Zellvergleich gegen die Sicherung.", size=9)
if os.path.exists(P_LISTE):
    with open(P_LISTE, encoding="utf-8", newline="") as f: liste = list(csv.DictReader(f))
    n_v = sum(1 for l in liste if l.get("Ausfuehrung") == "Verfasser")
    absatz("Ausführung: „Skript“ = von Datensicherung_Korrektur_2026-09-24.py geschrieben, „Verfasser“ = vom Verfasser am 24.09.2026 selbst in Excel geändert (vor dem Skriptlauf, %d Zellen: 10-m-Wert VS-06 post Versuch 1 nach Papierbogen, Familiarisierung VS-11, Bemerkung VS-11), vom Skript nur protokolliert. Sicherungen in Claude\\_Archiv: Stand vor Phase 1 (15.09.2026, 22:21 MESZ) als Studiendaten_U15_gesamt_vor_Phase1_2026-09-24.xlsx und Stand nach den Handänderungen, vor dem Skriptlauf, als Studiendaten_U15_gesamt_vor_Skriptkorrektur_2026-09-24.xlsx. Lesart des Verfasserurteils zu A5 („ist richtig“) als Fehlversuch, wie A4 und wie alle übrigen ungültigen 505-Versuche der Post-Testung (Beleg „Verfasser“)." % n_v, size=9)
    tabelle(["Nr", "Blatt", "Zelle", "Code", "Feld", "alt", "neu", "Grund", "Beleg", "Ausführung"],
            [[l["Nr"], l["Blatt"], l["Zelle"], l["Code"], l["Feld"], l["alt"], l["neu"], l["Grund"], l["Beleg"], l.get("Ausfuehrung", "")] for l in liste],
            (0.7, 1.8, 1, 1.1, 2.2, 3, 3, 3.6, 1.6, 1.4), size=7)
m1 = re.search(r"Protokollierte Zellen geändert: (\d+) von (\d+)", korr_txt)
m2 = re.search(r"Nicht protokollierte Abweichungen: (\d+)", korr_txt)
m3 = re.search(r"Geänderte Formelzellen: (\d+)", korr_txt)
m4 = re.search(r"Zellvergleich bestanden: (\w+)", korr_txt)
if m1: absatz("Zellvergleich gegen die Sicherung: %s von %s protokollierten Zellen geändert, %s nicht protokollierte Abweichungen, %s geänderte Formelzellen. Bestanden: %s." % (m1.group(1), m1.group(2), m2.group(1), m3.group(1), m4.group(1)))

# ------------------------------------------------------------------ 7 Datenstand
ueber("7 Datenstand, Rückleseprobe, Abbruchprobe (1.5)")
absatz("Skript Datensicherung_Export_2026-09-24.py. Der Datenstand enthält nur die Analysevariablen nach Spezifikation 0.4, Alter statt Geburtsdatum, keine Testdaten, keine Freitexte, keine Formelspalten. Variablennamen und Kategorien wie in Spezifikation 0.4, ein Nachtrag zur Spezifikation war nicht nötig.", size=9)
for muster in (r"Gegenlesen JSON gegen TXT.*", r"  Rückleseprobe bestanden: .*", r"  Abbruchprobe bestanden: .*", r"  Verstöße: \d+", r"  Abweichungen: \d+"):
    m = re.search(muster, export_txt)
    if m: absatz(m.group(0).strip(), size=9)
pp = os.path.join(P_DS, "Pruefsummen.txt")
if os.path.exists(pp):
    absatz("Ordner Statistik\\Datenstand_%s, Prüfsummen (SHA-256):" % DATUM, size=9)
    tabelle(["SHA-256", "Datei"], [[z.split("  ")[0], z.split("  ")[1]] for z in lese(pp).splitlines() if z.strip()], (10, 5), size=7)
    absatz("Pruefsummen.txt selbst: SHA-256 " + sha256(pp), size=8)
absatz("Der Verfasser setzt den Ordner nach dem Einfrieren schreibgeschützt. Ein späterer Stand bekommt einen neuen Ordner mit neuem Datum.", size=9)

# ------------------------------------------------------------------ 8 Nicht in diesem Task
ueber("8 Nicht in diesem Task, offen")
absatz("Korrektur der Begründung „Alter außerhalb“ für VS-18 in Projektanweisungen und Auswertungsplan (AP 15.09. § 6 Nr. 5, vorgemerkt in G17e) · Umgebungsdatei (L4) · Anhang G · Flowchart-Felder vor der Zuteilung, soweit der Verfasser sie nicht in Abschnitt 5 eingetragen hat · Erratum Khamis & Roche vor Phase 4 (L12). Die zwischengespeicherten Formelwerte des Workbooks (etwa Spalte gilt) werden beim nächsten Öffnen in Excel neu berechnet, der Datenstand verwendet keine Formelspalten.", size=9)

os.makedirs(os.path.dirname(ZIEL), exist_ok=True)
doc.save(ZIEL + ".docx")
print("geschrieben:", ZIEL + ".docx", os.path.getsize(ZIEL + ".docx"), "B")
try:
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(ZIEL), ZIEL + ".docx"], check=True, capture_output=True, timeout=180)
    print("geschrieben:", ZIEL + ".pdf", os.path.getsize(ZIEL + ".pdf"), "B")
except Exception as e:
    print("PDF nicht erzeugt:", e)
