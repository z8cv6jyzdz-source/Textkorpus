# -*- coding: utf-8 -*-
"""
Auswertung_B5_Deskriptiv_2026-09-12.py — Phase B, Schritt B5: Deskriptive Statistik
Bachelorarbeit U15-Plyometrie · DSHS Köln · Analyseprotokoll 2026-09-12
Stand 12.09.2026 (Rev. 2 der Kette vom 11.09.): 505 als MITTELWERT BEIDER BEINSEITEN (Ethikantrag § 3) = Zielgröße 505M ergänzt; 505 L/R nur noch deskriptiv. Übrige Rechenschritte unverändert.

Datenquelle: Statistik\Studiendaten_U15_gesamt.xlsx — 02_Rohdaten (B0.2/B0.3) und 01_Personen. Keine Signifikanztests
auf Baseline-Unterschiede (§ 4): berichtet werden n, M, SD, d und die Überlappung. Adhärenz und unerwünschte Effekte:
siehe Auswertung_B0_Mindestdosis_2026-09-12.txt §§ 6–7 (Fragebogen A).

Aufruf (aus Claude\03_Skripte\ oder Claude\; Workbook wird unter ..\Statistik bzw. ..\..\Statistik gesucht):  python Auswertung_B5_Deskriptiv_2026-09-12.py [Pfad zum Workbook]
Ausgabe: gleichnamige .txt neben dem Skript.

Inhalt:
  1  Stichprobencharakteristika (Alter, Größe, Gewicht, %PAH, Reifeband) je Gruppe und Verein
  2  Baseline-Tabelle je Zielgröße (Bestwert): n, M, SD, Differenz, d (gepoolte SD), Überlappung — zwei Bezugsmengen (§ 3):
     (i) alle Spieler mit Prä-Bestwert, (ii) ANCOVA-Set (prä + %PAH; IG zusätzlich mit Post-Wert)
  3  Dasselbe mit dem Mittelwert der Versuche (Sensitivitätsaggregation)
  4  Prä → Post innerhalb der IG (Bestwert und Mittelwert), je Verein und gesamt — deskriptiv, kein Gruppenvergleich
  5  Fehlende %PAH-Fälle (VS-11, VS-18): Wirkung auf den Baseline-Unterschied (§ 12 Nr. 18)
"""
import sys, os, io, math, datetime, statistics
from collections import Counter, defaultdict, OrderedDict
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
def _finde(*rel):
    # Skript liegt in Claude\03_Skripte\ (Ordnerkonvention § 1.4) oder flach in Claude\ — beide Lagen werden gefunden
    for up in ('..', os.path.join('..', '..'), os.path.join('..', '..', '..')):
        p = os.path.join(HERE, up, *rel)
        if os.path.exists(p): return os.path.abspath(p)
    return os.path.join(HERE, '..', *rel)
WB = sys.argv[1] if len(sys.argv) > 1 and not sys.argv[1].startswith('--') else _finde('Statistik', 'Studiendaten_U15_gesamt.xlsx')
OUT = os.path.splitext(os.path.abspath(__file__))[0] + '.txt'
buf = io.StringIO()
def P(*a):
    s = ' '.join(str(x) for x in a); print(s); buf.write(s + '\n')

ZIEL = OrderedDict([
    ('5m',   ('Sprint_5m', '–', 's', 'min')),
    ('10m',  ('Sprint_10m', '–', 's', 'min')),
    ('30m',  ('Sprint_30m', '–', 's', 'min')),
    ('505L', ('COD_505', 'L', 's', 'min')),
    ('505R', ('COD_505', 'R', 's', 'min')),
    ('505M', ('COD_505', 'M', 's', 'min')),   # abgeleitet: Mittelwert beider Beinseiten (Ethikantrag § 3), keine Rohzeilen
    ('SBJ',  ('Standweitsprung', '–', 'cm', 'max')),
])
GRUPPE = {'Hohenlind': 'IG', 'Blau-Weiß Köln': 'IG', 'Vorwärts Spoho': 'KG'}
VEREIN_KURZ = {'Hohenlind': 'A', 'Blau-Weiß Köln': 'B', 'Vorwärts Spoho': 'C'}
def num(v):
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(',', '.'))
    except ValueError: return None
def lade_rohdaten(pfad):
    wb = openpyxl.load_workbook(pfad, data_only=True, read_only=True)
    ws = wb['02_Rohdaten']; rows = list(ws.iter_rows(values_only=True))
    COL = {str(h): i for i, h in enumerate(rows[0])}
    out = []
    for r in rows[1:]:
        if r[COL['Code']] is None: continue
        d = {c: r[COL[c]] for c in ('Code', 'Verein', 'Zeitpunkt', 'Test', 'Seite', 'Versuch', 'Wert', 'ungültig', 'Bemerkung')}
        d['Wert'] = num(d['Wert'])
        d['ungültig'] = str(d['ungültig']).strip() if d['ungültig'] not in (None, '') else ''
        d['gilt_neu'] = 1 if (d['Wert'] is not None and d['ungültig'] == '') else 0
        d['Gruppe'] = GRUPPE[d['Verein']]; d['VereinK'] = VEREIN_KURZ[d['Verein']]
        d['Z'] = next(k for k, (t, s, u, m) in ZIEL.items() if t == d['Test'] and s == d['Seite'])
        out.append(d)
    return out, wb

def ergaenze_505M(agg, tupel=False):
    """505 Mittelwert beider Beinseiten (Ethikantrag § 3, Auswertungsplan 12.09.): je Spieler und Zeitpunkt
    best = (Best_L + Best_R)/2, mittel = (Mittel_L + Mittel_R)/2 — nur wenn beide Seiten einen gültigen Wert haben.
    'werte' = versuchsweise gepaarte Einzelversuche ((V1_L+V1_R)/2, …), k = min(k_L, k_R): dient nur der Messgüte des
    zusammengesetzten Werts (B3/B8); Bestwert und Mittelwert werden NICHT aus den Paaren gebildet."""
    for (c, zp, z) in [k for k in agg if k[2] == '505L']:
        L, R = agg[(c, zp, '505L')], agg.get((c, zp, '505R'))
        if R is None: continue
        k = min(L['k'], R['k'])
        if 'werte' not in L: paare = []                     # Skripte ohne Einzelversuche (B5–B7): nur k, best, mittel
        elif tupel: paare = [(j + 1, (L['werte'][j][1] + R['werte'][j][1]) / 2) for j in range(k)]
        else: paare = [(L['werte'][j] + R['werte'][j]) / 2 for j in range(k)]
        agg[(c, zp, '505M')] = dict(k=k, werte=paare, best=(L['best'] + R['best']) / 2, mittel=(L['mittel'] + R['mittel']) / 2)
    return agg

def aggregiere(data):
    zellen = defaultdict(list)
    for d in sorted(data, key=lambda d: d['Versuch']):
        if d['gilt_neu'] == 1: zellen[(d['Code'], d['Zeitpunkt'], d['Z'])].append(d['Wert'])
    agg = {}
    for key, w in zellen.items():
        richtung = ZIEL[key[2]][3]
        agg[key] = dict(k=len(w), best=(min(w) if richtung == 'min' else max(w)), mittel=sum(w) / len(w))
    return agg
def lade_personen(wb):
    ws = wb['01_Personen']; rows = list(ws.iter_rows(values_only=True)); h = [str(x) for x in rows[0]]
    pers = OrderedDict()
    for r in rows[1:]:
        d = dict(zip(h, r)); c = d.get('Code')
        if not isinstance(c, str) or c[:3] not in ('HL-', 'BW-', 'VS-'): continue
        pers[c] = dict(Gruppe=GRUPPE[d['Verein']], VereinK=VEREIN_KURZ[d['Verein']], PAH=num(d.get('PAH_Prozent')), Alter=num(d.get('Alter_prä_dez')),
                       Groesse=num(d.get('Größe_prä_cm')), Gewicht=num(d.get('Gewicht_prä_kg')), Reife=d.get('Reifeband'), Geb=d.get('Geburtsdatum'))
    return pers

data, wb = lade_rohdaten(WB); agg = ergaenze_505M(aggregiere(data)); pers = lade_personen(wb)
codes = list(pers)
def M(x): return sum(x) / len(x)
def SD(x): return statistics.stdev(x) if len(x) > 1 else float('nan')
def d_pooled(a, b):
    n1, n2 = len(a), len(b)
    sp = math.sqrt(((n1 - 1) * SD(a) ** 2 + (n2 - 1) * SD(b) ** 2) / (n1 + n2 - 2))
    return (M(a) - M(b)) / sp, sp

P('=' * 100)
P('B5 DESKRIPTIVE STATISTIK — Stichprobe, Baseline (ohne Signifikanztests, § 4), Prä→Post innerhalb der IG')
P('Skript:', os.path.basename(__file__), '· Lauf:', datetime.datetime.now().strftime('%Y-%m-%d %H:%M'), '· Workbook:', os.path.abspath(WB))
P('=' * 100)

# ---------------------------------------------------------------- 1 Stichprobe
P('\n1  STICHPROBENCHARAKTERISTIKA (01_Personen)')
P(f"{'Merkmal':16s}{'Gruppe':8s}{'n':>4s}{'M':>9s}{'SD':>8s}{'Min':>8s}{'Max':>8s}")
for merk, key in (('Alter (J.)', 'Alter'), ('Größe (cm)', 'Groesse'), ('Gewicht (kg)', 'Gewicht'), ('%PAH', 'PAH')):
    for lab, filt in (('IG', lambda c: pers[c]['Gruppe'] == 'IG'), ('  A', lambda c: pers[c]['VereinK'] == 'A'), ('  B', lambda c: pers[c]['VereinK'] == 'B'), ('KG', lambda c: pers[c]['Gruppe'] == 'KG')):
        v = [pers[c][key] for c in codes if filt(c) and pers[c][key] is not None]
        P(f'{merk:16s}{lab:8s}{len(v):4d}{M(v):9.2f}{SD(v):8.2f}{min(v):8.2f}{max(v):8.2f}')
P('  Geburtsjahrgänge:', {g: dict(Counter(pers[c]['Geb'].year for c in codes if pers[c]['Gruppe'] == g and pers[c]['Geb'] is not None)) for g in ('IG', 'KG')})
P('  Reifeband (deskriptiv, § 2):', {g: dict(Counter(str(pers[c]['Reife']) for c in codes if pers[c]['Gruppe'] == g)) for g in ('IG', 'KG')})
ig_p = [pers[c]['PAH'] for c in codes if pers[c]['Gruppe'] == 'IG' and pers[c]['PAH'] is not None]; kg_p = [pers[c]['PAH'] for c in codes if pers[c]['Gruppe'] == 'KG' and pers[c]['PAH'] is not None]
lo, hi = max(min(ig_p), min(kg_p)), min(max(ig_p), max(kg_p))
P(f'  %PAH-Überlappungsbereich {lo:.2f}–{hi:.2f}: IG {sum(1 for x in ig_p if lo <= x <= hi)}/{len(ig_p)} · KG {sum(1 for x in kg_p if lo <= x <= hi)}/{len(kg_p)} Spieler im Bereich; d = {d_pooled(ig_p, kg_p)[0]:+.2f}')
ig_a = [pers[c]['Alter'] for c in codes if pers[c]['Gruppe'] == 'IG' and pers[c]['Alter'] is not None]; kg_a = [pers[c]['Alter'] for c in codes if pers[c]['Gruppe'] == 'KG' and pers[c]['Alter'] is not None]
P(f'  Alter: IG {min(ig_a):.2f}–{max(ig_a):.2f} · KG {min(kg_a):.2f}–{max(kg_a):.2f} → ' + ('überlappt nicht' if max(kg_a) < min(ig_a) else 'überlappt') + ' (als Kovariate ungeeignet, § 3)')
P('  Fassung 9 § 3: %PAH IG (n = 17) 94,30 ± 2,99 · KG (n = 11) 89,73 ± 2,79 · Überlappung 86,65–95,02 (9 IG, 10 KG) — IG-Zahlen überholt (HL-08 jetzt vorhanden).')

# ---------------------------------------------------------------- 2/3 Baseline
def baseline(agg_key, titel, mengen):
    P(f'\n{titel}')
    for mname, filt in mengen:
        P(f'  Bezugsmenge: {mname}')
        P(f"{'Zielgröße':10s}{'n IG':>5s}{'M IG':>9s}{'SD IG':>8s}{'n KG':>5s}{'M KG':>9s}{'SD KG':>8s}{'Diff':>8s}{'d':>7s}{'gem. Bereich':>20s}{'IG im Ber.':>11s}{'KG im Ber.':>11s}")
        for z in ZIEL:
            ig = [agg[(c, 'prä', z)][agg_key] for c in codes if pers[c]['Gruppe'] == 'IG' and (c, 'prä', z) in agg and filt(c, z)]
            kg = [agg[(c, 'prä', z)][agg_key] for c in codes if pers[c]['Gruppe'] == 'KG' and (c, 'prä', z) in agg and filt(c, z)]
            d, sp = d_pooled(ig, kg); lo, hi = max(min(ig), min(kg)), min(max(ig), max(kg))
            dec = 0 if z == 'SBJ' else 3
            f = lambda x: f'{x:.{dec}f}' if z != 'SBJ' else f'{x:.1f}'
            P(f"{z:10s}{len(ig):5d}{f(M(ig)):>9s}{f(SD(ig)):>8s}{len(kg):5d}{f(M(kg)):>9s}{f(SD(kg)):>8s}{f(M(ig) - M(kg)):>8s}{d:7.2f}{(f(lo) + '–' + f(hi)):>20s}{(str(sum(1 for x in ig if lo <= x <= hi)) + '/' + str(len(ig))):>11s}{(str(sum(1 for x in kg if lo <= x <= hi)) + '/' + str(len(kg))):>11s}")
alle = ('alle Spieler mit Prä-Wert (Baseline-Tabelle 5.1)', lambda c, z: True)
anc = ('ANCOVA-Set: prä + %PAH; IG zusätzlich mit Post-Wert (Q-Rechnung der Fallzahlbegründung)', lambda c, z: pers[c]['PAH'] is not None and (pers[c]['Gruppe'] == 'KG' or (c, 'post', z) in agg))
baseline('best', '2  BASELINE — BESTWERT (Hauptaggregation). d = (M_IG − M_KG) / gepoolte SD; negativ bei Zeiten = IG schneller. KEINE Signifikanztests (§ 4).', [alle, anc])
P('  Fassung 9 § 3 (alle Fälle mit Bestwert): 5 m 1,05±0,08 vs 1,14±0,04 d −1,22 · 10 m 1,85±0,10 vs 1,97±0,08 d −1,29 · 30 m 4,53±0,23 vs 4,93±0,25 d −1,68 ·')
P('     SBJ 237±13 vs 224±19 d 0,84 · 505 L 2,47±0,11 vs 2,54±0,11 d −0,70 · 505 R 2,48±0,11 vs 2,50±0,13 d −0,17. Q-Rechnung: −1,070 / −1,169 / −1,512 / +0,603 / −0,471 / −0,062.')
P('  Überlappung § 3 (IG-Spieler im gemeinsamen Bereich): 5 m 6/17 · 10 m 14/17 · 30 m 12/17 · SBJ 17/17 · 505 L 13/16 · 505 R 15/16 (auf n = 17 gerechnet).')
baseline('mittel', '3  BASELINE — MITTELWERT DER VERSUCHE (Sensitivitätsaggregation § 11.8)', [alle, anc])

# ---------------------------------------------------------------- 4 Prä→Post IG
P('\n4  PRÄ → POST INNERHALB DER IG — deskriptiv, gepaart; KEIN Gruppenvergleich (KG-Post steht aus)')
P('  Δ = post − prä je Spieler; Vorzeichen: Zeiten negativ = schneller, Standweitsprung positiv = weiter. 95-%-KI des Mittels der Δ (t-Verteilung).')
from scipy import stats
for agg_key, lab in (('best', 'Bestwert'), ('mittel', 'Mittelwert der Versuche')):
    P(f'\n  Aggregation: {lab}')
    P(f"{'Zielgröße':10s}{'Menge':6s}{'n':>3s}{'M prä':>9s}{'SD prä':>8s}{'M post':>9s}{'SD post':>8s}{'M Δ':>9s}{'SD Δ':>8s}{'95%-KI Δ':>20s}{'Δ/SD_prä':>9s}{'verb./gleich/schl.':>19s}")
    for z in ZIEL:
        for mname, filt in (('A', lambda c: pers[c]['VereinK'] == 'A'), ('B', lambda c: pers[c]['VereinK'] == 'B'), ('IG', lambda c: pers[c]['Gruppe'] == 'IG')):
            pl = [c for c in codes if filt(c) and (c, 'prä', z) in agg and (c, 'post', z) in agg]
            if len(pl) < 2: P(f'{z:10s}{mname:6s}{len(pl):3d}   —'); continue
            pre = [agg[(c, 'prä', z)][agg_key] for c in pl]; po = [agg[(c, 'post', z)][agg_key] for c in pl]; dl = [b - a for a, b in zip(pre, po)]
            t = stats.t.ppf(0.975, len(dl) - 1); se = SD(dl) / math.sqrt(len(dl))
            richtung = ZIEL[z][3]; verb = sum(1 for x in dl if (x < 0 if richtung == 'min' else x > 0)); gl = sum(1 for x in dl if x == 0); schl = len(dl) - verb - gl
            dec = 1 if z == 'SBJ' else 3
            f = lambda x: f'{x:.{dec}f}'
            P(f"{z:10s}{mname:6s}{len(pl):3d}{f(M(pre)):>9s}{f(SD(pre)):>8s}{f(M(po)):>9s}{f(SD(po)):>8s}{f(M(dl)):>9s}{f(SD(dl)):>8s}{(f(M(dl) - t * se) + ' bis ' + f(M(dl) + t * se)):>20s}{M(dl) / SD(pre):9.2f}{(str(verb) + '/' + str(gl) + '/' + str(schl)):>19s}")
P('  Datendurchsicht Tab. 2 (Bestwert): A 5 m +0,080 · 10 m +0,097 · 30 m +0,111 · SBJ −6,86 · 505 L +0,033 (n 6) · 505 R +0,017; B 5 m −0,008 · 30 m +0,012 · SBJ −2,56 · 505 L −0,022 · 505 R +0,006 (n 7).')
P('  ⚠ Diese Zahlen sind keine Interventionseffekte (kein Vergleich, Belastungslage der Post-Termine ungleich, Datendurchsicht § 3).')

# ---------------------------------------------------------------- 5 fehlende PAH-Fälle
P('\n5  FEHLENDE %PAH-FÄLLE (VS-11, VS-18) — Wirkung auf den Baseline-Unterschied (§ 12 Nr. 18)')
P(f"{'Zielgröße':10s}{'KG-Rang VS-11':>14s}{'KG-Rang VS-18':>14s}{'d mit beiden':>13s}{'d ohne beide':>13s}{'Δd':>7s}")
for z in ZIEL:
    kg_all = [(c, agg[(c, 'prä', z)]['best']) for c in codes if pers[c]['Gruppe'] == 'KG' and (c, 'prä', z) in agg]
    ig = [agg[(c, 'prä', z)]['best'] for c in codes if pers[c]['Gruppe'] == 'IG' and (c, 'prä', z) in agg]
    richtung = ZIEL[z][3]
    ranked = sorted(kg_all, key=lambda t: t[1], reverse=(richtung == 'max'))   # Rang 1 = bester
    rang = {c: i + 1 for i, (c, _) in enumerate(ranked)}
    kg_mit = [v for _, v in kg_all]; kg_ohne = [v for c, v in kg_all if c not in ('VS-11', 'VS-18')]
    d1 = d_pooled(ig, kg_mit)[0]; d2 = d_pooled(ig, kg_ohne)[0]
    P(f"{z:10s}{(str(rang['VS-11']) + '/' + str(len(kg_all)) if 'VS-11' in rang else '— (kein Wert)'):>14s}{(str(rang['VS-18']) + '/' + str(len(kg_all)) if 'VS-18' in rang else '— (kein Wert)'):>14s}{d1:13.3f}{d2:13.3f}{d2 - d1:+7.3f}")
P('  Rang 1 = bester Wert der KG. § 12 Nr. 18 [EXTRAPOLATION]: „bei allen sechs Zielgrößen schrumpft der Baseline-Unterschied ohne diese Fälle“ — hier geprüft.')
P('  Der Ausfall ist nicht neutral, wenn beide zur schwächeren KG-Hälfte gehören: Der Vergleich ohne sie stellt der IG eine leistungsstärkere KG gegenüber (konservativ für die IG).')

open(OUT, 'w', encoding='utf-8').write(buf.getvalue())
print('\n→ geschrieben:', OUT)
