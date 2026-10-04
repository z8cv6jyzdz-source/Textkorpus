# -*- coding: utf-8 -*-
"""
Analyse_Boumparis_Hausstil_2026-10-02.py — Analysebefund „Boumparis et al. (2026)“ nach .docx und .pdf
Bachelorarbeit U15-Plyometrie · Arbeitsdokument, kein Manuskripttext · Task „Boumparis“, 02.10.2026.

Ruft Werkzeug_Hausstil_2026-09-25.py auf (pandoc, Hausstil nach Projektanweisungen § 9, soffice) und setzt danach:
- DIN A4 quer für das ganze Dokument (Tabellen mit bis zu fünf Spalten),
- Kopfzeilen jeder Tabelle als Wiederholungszeile, Tabellenzeilen nicht über Seiten getrennt,
- Code-Auszeichnung (Dateinamen) in 8 pt,
- Spaltenbreiten nach der Textmenge je Spalte,
- die PDF neu mit soffice.
Muster: Recherche_Hausstil_Umsetzungsrate_2026-10-02.py (Rev. 147), unverändert bis auf diesen Kopf.
Aufruf: python Analyse_Boumparis_Hausstil_2026-10-02.py <Befund.md> <Werkzeug_Hausstil_2026-09-25.py>
"""
import sys, os, subprocess
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MD, WERKZEUG = sys.argv[1], sys.argv[2]
DOCX = os.path.splitext(os.path.abspath(MD))[0] + '.docx'
subprocess.run([sys.executable, WERKZEUG, MD, DOCX], check=True)

d = Document(DOCX)
for s in d.sections:
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = Cm(29.7), Cm(21.0)
    for rand in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(s, rand, Cm(1.8))

for t in d.tables:
    for i, zeile in enumerate(t.rows):
        trPr = zeile._tr.get_or_add_trPr()
        if i == 0:
            h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
        c = OxmlElement('w:cantSplit'); c.set(qn('w:val'), 'true'); trPr.append(c)

def laeufe():
    for p in d.paragraphs:
        yield from p.runs
    for t in d.tables:
        for zeile in t.rows:
            for zelle in zeile.cells:
                for p in zelle.paragraphs:
                    yield from p.runs
for r in laeufe():
    if r.style is not None and r.style.name == 'Verbatim Char':
        r.font.size = Pt(8)

# Spaltenbreiten nach Textmenge je Spalte (neu gegenüber Rev. 139): Mindestbreite nach dem längsten Wort,
# der Rest nach der mittleren Textmenge verteilt.
BREITE_CM = 29.7 - 2 * 1.8
for t in d.tables:
    spalten = len(t.columns)
    gewichte, minima = [], []
    for j in range(spalten):
        texte = [zeile.cells[j].text for zeile in t.rows]
        mittel = sum(len(x) for x in texte) / max(1, len(texte))
        gewichte.append(min(max(mittel, 8.0), 220.0) ** 0.75)
        wort = max((len(w) for x in texte for w in x.split()), default=4)
        minima.append(min(5.0, 0.45 + 0.17 * wort))
    rest = BREITE_CM - sum(minima)
    if rest < 0:
        minima = [m * BREITE_CM / sum(minima) for m in minima]
        rest = 0.0
    tblPr = t._tbl.tblPr
    layout = OxmlElement('w:tblLayout'); layout.set(qn('w:type'), 'fixed'); tblPr.append(layout)
    for j in range(spalten):
        breite = Cm(minima[j] + rest * gewichte[j] / sum(gewichte))
        t.columns[j].width = breite
        for zeile in t.rows:
            zeile.cells[j].width = breite

d.save(DOCX)
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', os.path.dirname(DOCX), DOCX],
               check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
print('geschrieben:', DOCX, 'und', os.path.splitext(DOCX)[0] + '.pdf', '· Abschnitte:', len(Document(DOCX).sections))
