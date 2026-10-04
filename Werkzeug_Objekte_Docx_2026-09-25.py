# -*- coding: utf-8 -*-
"""
Werkzeug_Objekte_Docx_2026-09-25.py — Objekte aus Objekte_2026-09-25.R als Word-Datei im Manuskriptformat
Bachelorarbeit U15-Plyometrie · DSHS Köln · Auswertungsverfahren 2026-09-24, Schritt 7.2 (Setzen)

Zweck: Liest Objekte_2026-09-25.md (Titel, Tabellen, Anmerkungen, Abbildungsunterschriften) und die
PNG-Dateien aus dem Ausgabeordner des R-Skripts und setzt sie in eine .docx nach F14 § 9 (dvs 2020):
Tabellentitel kursiv oberhalb (10 pt), Kopfzeile 15 % grau und fett, Gitternetz, Arial, Anmerkung
unterhalb, Abbildungen mit Beschriftung unterhalb. Die Zellen werden unverändert übernommen, keine
Zahl wird hier gebildet. Die Datei ist Übergabematerial für den Master, kein Manuskripttext.
Aufruf: python Werkzeug_Objekte_Docx_2026-09-25.py <Ausgabeordner des R-Skripts> <Ziel.docx>
Fassung: 2026-09-25, erste Fassung.
"""
import sys
import os
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ORDNER, ZIEL = sys.argv[1], sys.argv[2]
zeilen = open(os.path.join(ORDNER, 'Objekte_2026-09-25.md'), encoding='utf-8').read().split('\n')

d = Document()
sec = d.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)
stil = d.styles['Normal']
stil.font.name = 'Arial'
stil.font.size = Pt(10)
stil.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')


def absatz(text, kursiv=False, fett=False, groesse=10, nach=4):
    p = d.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(groesse)
    r.italic = kursiv
    r.bold = fett
    p.paragraph_format.space_after = Pt(nach)
    return p


def schattiere(zelle, hexfarbe):
    tcPr = zelle._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexfarbe)
    tcPr.append(shd)


def tabelle(kopf, zeilen_t):
    t = d.add_table(rows=1 + len(zeilen_t), cols=len(kopf))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, k in enumerate(kopf):
        c = t.rows[0].cells[j]
        c.text = ''
        r = c.paragraphs[0].add_run(k)
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.bold = True
        schattiere(c, 'D9D9D9')   # 15 % grau
    for i, z in enumerate(zeilen_t, 1):
        for j, v in enumerate(z):
            c = t.rows[i].cells[j]
            c.text = ''
            p = c.paragraphs[0]
            r = p.add_run(v)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            if j > 0 and re.fullmatch(r'[−+]?\d[\d,]*( %| s| cm)?', v.strip()):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph().paragraph_format.space_after = Pt(2)


absatz('Objekte des Manuskripts aus der Ergebnisdatei (Phase 7.2), Übergabematerial für den Master', fett=True, groesse=12)
absatz(zeilen[2] if len(zeilen) > 2 else '', groesse=8, nach=10)
i = 0
n_tab = n_abb = 0
while i < len(zeilen):
    z = zeilen[i]
    if z.startswith('*Tab. ') and z.endswith('*'):
        absatz(z.strip('*'), kursiv=True, nach=2)
        i += 2
        kopf = [c.strip() for c in zeilen[i].strip('|').split('|')]
        i += 2
        rows = []
        while i < len(zeilen) and zeilen[i].startswith('|'):
            rows.append([c.strip() for c in zeilen[i].strip('|').split('|')])
            i += 1
        tabelle(kopf, rows)
        n_tab += 1
        continue
    if z.startswith('Anmerkung. '):
        absatz(z, groesse=9, nach=14)
    if z.startswith('*Abb. '):
        m = re.search(r'\((Abb_[^)]*?)\.png', z)
        if m and os.path.exists(os.path.join(ORDNER, m.group(1) + '.png')):
            d.add_picture(os.path.join(ORDNER, m.group(1) + '.png'), width=Cm(16))
            absatz(z.replace('*', ''), groesse=9, nach=14)
            n_abb += 1
    i += 1
d.save(ZIEL)
print('geschrieben:', ZIEL, '· Tabellen:', n_tab, '· Abbildungen:', n_abb)
