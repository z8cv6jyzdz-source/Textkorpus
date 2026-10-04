# -*- coding: utf-8 -*-
"""
Master_Vorschlagsliste_2026-09-25.py — Vorschlagsliste des Endabgleichs (Abgleichprotokoll_Manuskript_2026-09-25 § 5)
in den Manuskript-Master übertragen
Bachelorarbeit U15-Plyometrie · DSHS Köln · Auswertungsverfahren 2026-09-24, Phase 7.3

Freigaben des Verfassers per Klick am 25.09.2026: 4.2 (Vorschlag 1, Kennwerte bleiben im Text, Reifestatus nach
R11), 4.3 (Vorschlag 6), 4.4 (Vorschläge 2 bis 5 einschließlich Tab. 1, Titel und Anmerkung), Anhang G (Vorschlag 7,
Gliederung v4 M7). Verfasserfestlegung 25.09.: keine Kennungen im Master, Rückverfolgbarkeit über den Endabgleich.
Jede Textstelle wird genau einmal ersetzt, sonst bricht das Skript ab. Zahlen der Tabelle kommen aus der CSV des
R-Objekts (Tab_1_Messguete.csv), keine Zahl von Hand. Ohne Semikolon (Semikola des Bestands über chr(59)).
Aufruf: python Master_Vorschlagsliste_2026-09-25.py <Master.docx> <Tab_1_Messguete.csv> <Ziel.docx>
Fassung: 2026-09-25, erste Fassung.
"""
import sys
import csv
import copy
from docx import Document
from docx.oxml.ns import qn

SRC, CSVT, DST = sys.argv[1], sys.argv[2], sys.argv[3]
SEMI = chr(59)
d = Document(SRC)
P = d.paragraphs
protokoll = []


def absatz_mit(text_anfang):
    treffer = [p for p in P if p.text.startswith(text_anfang)]
    if len(treffer) != 1:
        raise SystemExit('Absatz nicht genau einmal gefunden: ' + text_anfang[:60])
    return treffer[0]


def ersetze_in_absatz(p, alt, neu):
    """Ersetzt alt durch neu innerhalb genau eines Runs des Absatzes (Formatierung bleibt)."""
    runs = [r for r in p.runs if alt in r.text]
    if len(runs) != 1 or p.text.count(alt) != 1:
        raise SystemExit('Textstelle nicht genau einmal in einem Run: ' + alt[:60])
    runs[0].text = runs[0].text.replace(alt, neu)
    protokoll.append((alt[:70], neu[:70]))


# ---------------------------------------------------------------- 4.2 (Vorschlag 1 ohne Kennungen)
p42 = absatz_mit('Die Zielpopulation bildeten männliche Nachwuchsfußballspieler')
ersetze_in_absatz(p42, '94,65 ± 2,90 %PAH', '94,74 ± 2,77 %PAH')
ersetze_in_absatz(p42, '90,41 ± 2,85 %PAH', '90,50 ± 2,73 %PAH')

# ---------------------------------------------------------------- 4.3 (Vorschlag 6)
p43 = absatz_mit('Der Untersuchungsablauf gliederte sich je Verein in vier Phasen')
ersetze_in_absatz(p43, 'und C (15.09.) nach dem geplanten Zeitfenster', 'und C (14.09.) nach dem geplanten Zeitfenster')

# ---------------------------------------------------------------- 4.4 (Vorschlag 2 und Folge)
p44a = absatz_mit('Vorgesehen waren drei Versuche je Spieler und Zielgröße')
ersetze_in_absatz(p44a, ', innerhalb der Gruppen ohne erkennbaren Zusammenhang mit der Leistung', '')
ersetze_in_absatz(p44a, 'fehlen in der betreffenden Zielgröße (Tab. 1, n).', 'fehlen in der betreffenden Zielgröße.')

# ---------------------------------------------------------------- 4.4 (Vorschlag 4)
p44b = absatz_mit('Je Zielgröße wurde der typische Messfehler (TE) der Prä-Testung bestimmt')
alt4 = ('Er überstieg bei allen Zielgrößen den SESOI (Tab. 1)' + SEMI + ' Aussagen über einzelne Spieler sind nicht belastbar, '
        'während der Standardfehler des Gruppenmittels (TE/√n) mit 0,22 bis 0,72 des SESOI Gruppenvergleiche trägt (Abschn. 4.7).')
neu4 = ('Er überstieg bei allen Zielgrößen den SESOI (Tab. 1). Aussagen über einzelne Spieler sind nicht belastbar. '
        'Die Auflösung des Gruppenvergleichs tragen die kleinste nachweisbare Effektstärke und die Breite der Konfidenzintervalle (Lakens, 2022).')
ersetze_in_absatz(p44b, alt4, neu4)

# ---------------------------------------------------------------- 4.4 Tab. 1: Titel (Vorschlag 3)
cap = [p for p in P if p.style.name == 'Caption' and 'Messgüte der Testverfahren' in p.text]
if len(cap) != 1:
    raise SystemExit('Beschriftung von Tab. 1 nicht genau einmal gefunden')
ersetze_in_absatz(cap[0], '. Messgüte der Testverfahren zur Prä-Testung: typischer Messfehler und kleinster praktisch bedeutsamer Unterschied',
                  '. Messgüte je Zielgröße aus den Wiederholungsversuchen der Eingangstestung (Spieler mit mindestens zwei gültigen Versuchen) und typischer Messfehler der Abschlusstestung')

# ---------------------------------------------------------------- 4.4 Tab. 1: Zellen aus der CSV des R-Objekts (Vorschlag 3)
with open(CSVT, encoding='utf-8-sig', newline='') as f:
    zeilen = list(csv.reader(f))
kopf, daten = zeilen[0], zeilen[1:]
t = d.tables[0]
if [c.text for c in t.rows[0].cells] != ['Zielgröße', 'n', 'SD', 'TE [95-%-KI]', 'CV (%)', 'SESOI', 'TE/SESOI']:
    raise SystemExit('Tab. 1 im Master hat nicht die erwartete Kopfzeile')
if len(t.rows) != len(daten) + 1 or len(t.columns) != len(kopf):
    raise SystemExit('Tab. 1: Zeilen- oder Spaltenzahl passt nicht zur CSV')
BREITEN = ['2090', '400', '1600', '760', '730', '1040', '1600']
RAND = '57'  # Zellrand links und rechts in dxa (0,1 cm), damit die Spalten in 14,5 cm passen
if sum(int(b) for b in BREITEN) != 8220:
    raise SystemExit('Spaltenbreiten ergeben nicht die Tabellenbreite 8220')


def setze_zelle(zelle, text):
    """Schreibt den Zellwert in den ersten Run (Formatierung bleibt). Ein Wert mit Konfidenzintervall
    „x [a bis b]“ wird zweizeilig gesetzt: Wert, Zeilenumbruch, Intervall (Breite der Textspalte)."""
    absatz = zelle.paragraphs[0]
    if not absatz.runs:
        absatz.add_run(text)
        return
    r = absatz.runs[0]
    for weg in absatz.runs[1:]:
        weg._r.getparent().remove(weg._r)
    if ' [' in text and text.endswith(']'):
        wert, ki = text.split(' [', 1)
        r.text = wert
        r.add_break()
        r._r.add_t('[' + ki)
    else:
        r.text = text


for ci, wert in enumerate(kopf):
    setze_zelle(t.rows[0].cells[ci], wert)
for ri, zeile in enumerate(daten, start=1):
    for ci, wert in enumerate(zeile):
        setze_zelle(t.rows[ri].cells[ci], wert)
tblPr = t._tbl.tblPr
mar = tblPr.find(qn('w:tblCellMar'))
if mar is None:
    from docx.oxml import OxmlElement
    mar = OxmlElement('w:tblCellMar')
    tblPr.append(mar)
for seite in ('left', 'right'):
    el = mar.find(qn('w:' + seite))
    if el is None:
        from docx.oxml import OxmlElement
        el = OxmlElement('w:' + seite)
        mar.append(el)
    el.set(qn('w:w'), RAND)
    el.set(qn('w:type'), 'dxa')
grid = t._tbl.find(qn('w:tblGrid'))
for g, b in zip(grid.findall(qn('w:gridCol')), BREITEN):
    g.set(qn('w:w'), b)
for row in t.rows:
    for c, b in zip(row.cells, BREITEN):
        tcW = c._tc.tcPr.find(qn('w:tcW'))
        tcW.set(qn('w:w'), b)
protokoll.append(('Tab. 1 Zellen', '%d Zeilen x %d Spalten aus %s' % (len(daten), len(kopf), CSVT.split('/')[-1])))

# ---------------------------------------------------------------- 4.4 Anmerkung zu Tab. 1 (Vorschlag 5, ohne Kennungen)
anm = [p for p in P if p.text.startswith('Anmerkung. Prä-Testung (N = 31)')]
if len(anm) != 1:
    raise SystemExit('Anmerkung zu Tab. 1 nicht genau einmal gefunden')
anm = anm[0]
if anm.runs[0].text != 'Anmerkung. ' or anm.runs[0].italic is not True:
    raise SystemExit('Anmerkung: erster Run nicht wie erwartet')
NEU_ANM = ('n = Spieler mit mindestens zwei gültigen Versuchen in der Eingangstestung. TE = typischer Messfehler '
           '(gepoolte Intra-Personen-Standardabweichung der gültigen Versuche innerhalb der Sitzung) mit 95-%-Konfidenzintervall '
           'aus der χ²-Verteilung. CV = TE in Prozent des Mittelwerts aller gültigen Versuche. SESOI = kleinster praktisch '
           'bedeutsamer Unterschied (0,2 × Standardabweichung der Bestwerte aller Spieler). TE/SESOI > 1: Der Messfehler '
           'übersteigt den SESOI. 505 links/rechts = Wenderichtung, Seitenmittel aus beidseitig gültigen Versuchsnummern. '
           'TE post = typischer Messfehler der Abschlusstestung, n post: Sprint 5 m 25, Sprint 10 m 13, Sprint 30 m 26, '
           '505 links 21, 505 rechts 14, 505-Seitenmittel 12, Standweitsprung 23.')
anm.runs[1].text = NEU_ANM
for r in anm.runs[2:]:
    r._r.getparent().remove(r._r)
protokoll.append(('Anmerkung Tab. 1', NEU_ANM[:70]))

# ---------------------------------------------------------------- Anhang G (Vorschlag 7)
pG = absatz_mit('Anhang G: Sensitivitäts-Poweranalyse und SPSS-Syntax')
ersetze_in_absatz(pG, 'SPSS-Syntax', 'R-Skripte')

# ---------------------------------------------------------------- Kontrolle: n post der Anmerkung gegen die CSV-Kennzahlen? (nur Vorhandensein der Spalte)
if 'TE post [95-%-KI]' not in kopf:
    raise SystemExit('CSV ohne Spalte TE post')
for a, n in protokoll:
    if SEMI in n:
        raise SystemExit('Semikolon in neuem Text: ' + n)
d.save(DST)
print('Master geschrieben: %s' % DST)
for a, n in protokoll:
    print('  ALT: %s\n  NEU: %s' % (a, n))
