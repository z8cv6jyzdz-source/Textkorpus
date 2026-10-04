# -*- coding: utf-8 -*-
"""
Auswertung_B9_ITT_PerProtokoll_2026-09-12.py — Phase B, Schritt B9: Analysesets ITT und Per-Protokoll — Zugehörigkeit je Spieler und Zielgröße
Bachelorarbeit U15-Plyometrie · DSHS Köln · Analyseprotokoll 2026-09-12
Stand 12.09.2026 (Rev. 2): Zielgröße 505M (Mittelwert beider Beinseiten, Ethikantrag) ergänzt; Spalte „PP9“ = Antragskriterium ≥ 75 % (≥ 9/12) ausgewiesen.

ITT (§ 11.7): alle Zugeteilten; ausgewertet werden die vollständigen Fälle (Prä, Post, %PAH), keine Fortschreibung (kein LOCF).
  Ergebnis = „Wirkung des Programmangebots“. Hauptanalyse.
Per-Protokoll (B0.1): IG-Spieler mit ≥ 6 vollständig absolvierten Einheiten (Fragebogen A, CASE-korrigiert) gegen die unveränderte KG.
  Ergebnis = „Wirkung der Durchführung“; nicht randomisierter, beobachtender Vergleich. Zusätzlich, nie allein.
Vier Ausschlussklassen (CONSORT Box 6): Nichteinhaltung (nur PP) · Instrumentenfehler BW-21 (bleibt überall; für PP nicht klassifizierbar)
  · fehlende Kovariate VS-11, VS-18 · Versuchsausschluss (Nenner je Zielgröße). Nicht angetreten BW-07, BW-21: Ausfall bei der Nacherhebung.

Aufruf (aus Claude\03_Skripte\ oder Claude\; Workbook wird unter ..\Statistik bzw. ..\..\Statistik gesucht):  python Auswertung_B9_ITT_PerProtokoll_2026-09-12.py [Workbook] [Fragebogen-Datei]
Ausgabe: Zugehörigkeitstabelle je Spieler × Zielgröße mit Grund; Nenner je Analyse (Item 16).
"""
import sys, os, io, csv, datetime
from collections import Counter, defaultdict, OrderedDict
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
def _finde(*rel):
    # Skript liegt in Claude\03_Skripte\ (Ordnerkonvention § 1.4) oder flach in Claude\ — beide Lagen werden gefunden
    for up in ('..', os.path.join('..', '..'), os.path.join('..', '..', '..')):
        p = os.path.join(HERE, up, *rel)
        if os.path.exists(p): return os.path.abspath(p)
    return os.path.join(HERE, '..', *rel)
WB = sys.argv[1] if len(sys.argv) > 1 and not sys.argv[1].startswith('--') else _finde('Statistik', 'Studiendaten_U15_gesamt.xlsx')
FB = sys.argv[2] if len(sys.argv) > 2 and not sys.argv[2].startswith('--') else _finde('SoSci_Fragebogen', 'Fragebogen Datensatz.txt')
OUT = os.path.splitext(os.path.abspath(__file__))[0] + '.txt'
buf = io.StringIO()
def P(*a):
    s = ' '.join(str(x) for x in a); print(s); buf.write(s + '\n')
SCHWELLE_PP = 6; N_MIN = 8
ZIEL = OrderedDict([
    ('5m',   ('Sprint_5m', '–', 's', 'min')),
    ('10m',  ('Sprint_10m', '–', 's', 'min')),
    ('30m',  ('Sprint_30m', '–', 's', 'min')),
    ('505L', ('COD_505', 'L', 's', 'min')),
    ('505R', ('COD_505', 'R', 's', 'min')),
    ('505M', ('COD_505', 'M', 's', 'min')),   # abgeleitet: Mittelwert beider Beinseiten (Ethikantrag § 3), keine Rohzeilen
    ('SBJ',  ('Standweitsprung', '–', 'cm', 'max')),
])
GRUPPE = {'Hohenlind': 'IG', 'Blau-Weiß Köln': 'IG', 'Vorwärts Spoho': 'KG'}
VEREIN_KURZ = {'Hohenlind': 'A', 'Blau-Weiß Köln': 'B', 'Vorwärts Spoho': 'C'}
LISTE = {**{i: 'HL-%02d' % i for i in range(1, 9)}, **{i: 'BW-%02d' % (i - 8) for i in range(9, 23)}}
CASE_KORREKTUR = {'59': 'BW-06', '70': 'BW-06', '140': 'BW-04', '158': 'BW-04', '165': 'BW-04', '190': 'BW-04', '206': 'BW-04'}
MAP_WB = {'HL-04': 'HL-05'}
def num(v):
    if v is None or v == '': return None
    if isinstance(v, (int, float)): return float(v)
    try: return float(str(v).replace(',', '.'))
    except ValueError: return None
wb = openpyxl.load_workbook(WB, data_only=True, read_only=True)
ws = wb['02_Rohdaten']; rows = list(ws.iter_rows(values_only=True)); COL = {str(h): i for i, h in enumerate(rows[0])}
hat = defaultdict(set); grp = {}; ver = {}
for r in rows[1:]:
    c = r[COL['Code']]
    if c is None: continue
    grp[c] = GRUPPE[r[COL['Verein']]]; ver[c] = VEREIN_KURZ[r[COL['Verein']]]
    z = next(k for k, (t, s, u, m) in ZIEL.items() if t == r[COL['Test']] and s == r[COL['Seite']])
    ung = str(r[COL['ungültig']]).strip() if r[COL['ungültig']] not in (None, '') else ''
    if num(r[COL['Wert']]) is not None and ung == '': hat[(r[COL['Zeitpunkt']], z)].add(c)
for zp in ('prä', 'post'): hat[(zp, '505M')] = hat[(zp, '505L')] & hat[(zp, '505R')]   # 505M: beide Seiten gültig
ws = wb['01_Personen']; rows = list(ws.iter_rows(values_only=True)); h = [str(x) for x in rows[0]]
pah = {}; status = {}
for r in rows[1:]:
    d = dict(zip(h, r)); c = d.get('Code')
    if isinstance(c, str) and c[:3] in ('HL-', 'BW-', 'VS-'): pah[c] = num(d.get('PAH_Prozent')); status[c] = d.get('Status') or ''
codes = sorted(grp)
ganz = Counter()
if os.path.exists(FB):
    raw = open(FB, encoding='utf-16').read()
    for r in csv.DictReader(io.StringIO(raw), delimiter='\t'):
        c = CASE_KORREKTUR.get(r['CASE'], LISTE.get(int(r['H010'])) if r['H010'].strip().isdigit() else None)
        if c is None: continue
        c = MAP_WB.get(c, c)
        if r['H004'] == '1': ganz[c] += 1

P('=' * 110)
P('B9 ANALYSESETS — ITT und Per-Protokoll je Spieler und Zielgröße')
P('Skript:', os.path.basename(__file__), '· Lauf:', datetime.datetime.now().strftime('%Y-%m-%d %H:%M'), '· Workbook:', os.path.abspath(WB))
P('=' * 110)
P('\nLegende je Zelle: I = im ITT-Analyseset (prä + post + %PAH) · P = zusätzlich im Per-Protokoll-Set (IG ≥ 6 Einheiten) · K = KG-Post steht aus (Set offen)')
P('  −prä = kein gültiger Prä-Wert · −post = kein gültiger Post-Wert · −PAH = %PAH fehlt · n.a. = nicht angetreten · <6 = unter Mindestdosis (nur ITT) · Instr. = Instrumentenfehler (nur ITT)')
P(f"{'Code':7s}{'Gr':3s}{'V':2s}{'Einh.':>6s}" + ''.join(f'{z:>9s}' for z in ZIEL) + '   Status')
for c in codes:
    line = f"{c:7s}{grp[c]:3s}{ver[c]:2s}{(str(ganz.get(c, 0)) if grp[c] == 'IG' else '—'):>6s}"
    for z in ZIEL:
        if grp[c] == 'IG' and c in ('BW-07', 'BW-21') and c not in hat[('post', z)]: cell = 'n.a.'
        elif c not in hat[('prä', z)]: cell = '−prä'
        elif pah.get(c) is None: cell = '−PAH'
        elif grp[c] == 'KG': cell = 'K'
        elif c not in hat[('post', z)]: cell = '−post'
        else:
            cell = 'I'
            if c == 'BW-21': cell += ' Instr.'
            elif ganz.get(c, 0) >= SCHWELLE_PP: cell += '+P'
            else: cell += ' <6'
        line += f'{cell:>9s}'
    P(line + '   ' + status.get(c, '')[:38])
P('\nNENNER JE ANALYSE (CONSORT Item 16) — KG in Klammern: maximal erreichbar, Post steht aus')
P(f"{'Zielgröße':10s}{'ITT IG':>8s}{'ITT KG':>8s}{'PP IG':>7s}{'PP9 IG':>8s}{'§ 11.9 ITT':>12s}{'§ 11.9 PP':>11s}{'§ 11.9 PP9':>12s}")
for z in ZIEL:
    itt = [c for c in codes if grp[c] == 'IG' and c in hat[('prä', z)] and c in hat[('post', z)] and pah.get(c) is not None]
    kg = [c for c in codes if grp[c] == 'KG' and c in hat[('prä', z)] and pah.get(c) is not None]
    pp = [c for c in itt if ganz.get(c, 0) >= SCHWELLE_PP]; pp9 = [c for c in itt if ganz.get(c, 0) >= 9]
    P(f"{z:10s}{len(itt):8d}{('(' + str(len(kg)) + ')'):>8s}{len(pp):7d}{len(pp9):8d}{('Inferenz' if len(itt) >= N_MIN and len(kg) >= N_MIN else 'deskriptiv'):>12s}{('Inferenz' if len(pp) >= N_MIN and len(kg) >= N_MIN else 'deskriptiv'):>11s}{('Inferenz' if len(pp9) >= N_MIN and len(kg) >= N_MIN else 'deskriptiv'):>12s}")
P('PP9 = Ethikantrag-Kriterium „Adherence ≥ 75 %“ (≥ 9 von 12): nach § 11.9 nur deskriptiv → Hauptanalyse ITT (Auswertungsplan 12.09.). 505M = Mittelwert beider Seiten; 505 L/R deskriptiv.')
P('\nSprachregelung (§ 10): ITT-Ergebnis = „Wirkung des Programmangebots“; PP-Ergebnis = beobachtender Vergleich der Durchführenden, kein kausaler Effekt.')
P('Verdünnung (§ 11.7): 92 von 216 angebotenen Einheiten vollständig (42,6 %) — ein ITT-Nullbefund bedeutet zunächst nur, dass das Angebot in dieser Umsetzungsrate nichts bewirkt hat.')
open(OUT, 'w', encoding='utf-8').write(buf.getvalue())
print('\n→ geschrieben:', OUT)
