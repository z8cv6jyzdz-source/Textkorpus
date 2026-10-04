# -*- coding: utf-8 -*-
"""
Master_4.7_2026-09-25.py — Abschnitt 4.7 Statistische Auswertung in den Manuskript-Master einfügen
Bachelorarbeit U15-Plyometrie · DSHS Köln · Skill kapiteltext-bachelorarbeit, Schritt 8 (direkter Einbau auf Anweisung)

Freigabe des Verfassers per Klick am 25.09.2026: Stufe „Empfehlung“ des Textvorschlags 4.7
(Textvorschlag_4.7_2026-09-25_Empf.txt, 04_Uebergaben\\Textvorschlag_4.7_2026-09-25.md § 6).
Die Absätze werden als reine Absätze ohne Direktformatierung (Formatvorlage Standard, wie der Bestand) unmittelbar
nach der Überschrift „4.7 Statistische Auswertung“ eingefügt. Vorbedingungen, sonst Abbruch: genau eine solche
Überschrift, der Abschnitt ist leer (nächster Absatz ist die Überschrift „5 Ergebnisse“), keine Kommentare, kein
Semikolon im neuen Text außerhalb von Zitierklammern, kein Abschnittsverweis. Ohne Semikolon im Skript (chr(59)).
Aufruf: python Master_4.7_2026-09-25.py <Master.docx> <Textstufe.txt> <Ziel.docx>
Fassung: 2026-09-25, erste Fassung.
"""
import sys
import re
import copy
import zipfile
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, TXT, DST = sys.argv[1], sys.argv[2], sys.argv[3]
SEMI = chr(59)
with zipfile.ZipFile(SRC) as z:
    if 'word/comments.xml' in z.namelist():
        raise SystemExit('Master trägt comments.xml, Einbau abgebrochen (Prozessregel F14 § 1.2)')
txt = open(TXT, encoding='utf-8').read()
txt = '\n'.join(z for z in txt.split('\n') if not z.startswith('#'))
absaetze = [re.sub(r'\s+', ' ', a).strip() for a in re.split(r'\n\s*\n', txt) if a.strip()]
if len(absaetze) != 7:
    raise SystemExit('Erwartet sieben Absätze, gefunden %d' % len(absaetze))
for a in absaetze:
    ohne_klammern = re.sub(r'\([^)]*\)', '', a)
    if SEMI in ohne_klammern:
        raise SystemExit('Semikolon außerhalb einer Klammer: ' + a[:60])
    if re.search(r'Abschn(?:itt)?\.?\s*\d|Kapitel\s*\d', a):
        raise SystemExit('Abschnittsverweis im neuen Text: ' + a[:60])
    if '[K-' in a or '⟨' in a:
        raise SystemExit('Kennung oder Platzhalter im neuen Text: ' + a[:60])

d = Document(SRC)
P = d.paragraphs
kopf = [i for i, p in enumerate(P) if p.text.strip() == '4.7 Statistische Auswertung' and p.style.name == 'Heading 2']
if len(kopf) != 1:
    raise SystemExit('Überschrift 4.7 nicht genau einmal gefunden')
i47 = kopf[0]
if not (P[i47 + 1].text.strip() == '5 Ergebnisse' and P[i47 + 1].style.name == 'Heading 1'):
    raise SystemExit('Abschnitt 4.7 ist nicht leer oder die Folgeüberschrift fehlt')
n_vor = len(P)
anker = P[i47]._p
for a in absaetze:
    p = OxmlElement('w:p')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = a
    r.append(t)
    p.append(r)
    anker.addnext(p)
    anker = p
d.save(DST)
d2 = Document(DST)
P2 = d2.paragraphs
if len(P2) != n_vor + 7:
    raise SystemExit('Absatzzahl nach dem Einbau unerwartet: %d statt %d' % (len(P2), n_vor + 7))
neu = [p.text for p in P2[i47 + 1:i47 + 8]]
if neu != absaetze:
    raise SystemExit('Rücklesen: eingefügter Text weicht ab')
print('4.7 eingebaut: %d Absätze, %d Wörter, Absätze im Master %d → %d' % (
    len(absaetze), sum(len(a.split()) for a in absaetze), n_vor, len(P2)))
